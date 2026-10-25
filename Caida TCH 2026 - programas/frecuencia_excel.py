# -*- coding: utf-8 -*-
"""Excel: frecuencia del riego suerte a suerte (Riopaila), a partir de las pasadas de riego_frecuencia.py.

Uso: python frecuencia_excel.py <riego_pasadas.pkl> <rio_factores_sep.pkl> <salida.xlsx>
"""
import sys
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
import presupuesto_cosecha as pc  # noqa: E402
from tchmotor import dataset_gee as dg, produccion as pr  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")
P = pd.read_pickle(sys.argv[1])
C = pd.read_pickle(sys.argv[2])
SAL = sys.argv[3]
REF = P.fin.max().date()  # último registro de riego

# ---------------------------------------------------------------- datos
P = P.sort_values(["ste", "ini"]).reset_index(drop=True)
P["valido"] = np.where(P.dias.isna(), "No: primera pasada", np.where(P.cruza_cosecha, "No: hubo cosecha en medio", "Sí"))
ok = P[P.valido == "Sí"]
y = ok.ini.dt.year
med26 = ok[y == 2026].groupby("ste").dias.median()
med_seca = ok[(y == 2026) & ok.ini.dt.month.between(6, 9)].groupby("ste").dias.median()
med_ref = ok[y.between(2023, 2025)].groupby("ste").dias.median()

ma, _ = pc.leer_maestro("riopaila")
M = pd.DataFrame([dict(ste=s, finca=v["finca"], zona=v["zona"], area=v["area"], var=v["var"], corte=v["ncorte"],
                       ult=v["ult"], siembra=v["siembra"]) for s, v in ma.items()])
M["inicio"] = [max([d for d in (a, b) if d] or [None]) if (a or b) else None for a, b in zip(M.ult, M.siembra)]
extra = sorted(set(P.ste) - set(M.ste))
M = pd.concat([M, pd.DataFrame({"ste": extra, "finca": "(no está en el maestro)"})], ignore_index=True)
M = M.sort_values(["zona", "finca", "ste"], na_position="last").reset_index(drop=True)

c26 = C[C.anio == 2026][["ste", "Mes", "tch_calc"]]
maestro_pr, _ = pr.cargar_maestro()
rec, _ = pr.cosechas_recientes(pr.leer_cosechas_recientes(), maestro_pr, dg.cortes_zqm(dg.leer_zqm()))
oc = rec[(rec.planta == "riopaila") & (rec.anio == 2026) & (rec.mes == 10)]
c26 = pd.concat([c26, pd.DataFrame({"ste": oc.sec_ste, "Mes": 10, "tch_calc": oc.tch})]).sort_values("Mes").drop_duplicates("ste", keep="last")
TCH = c26.set_index("ste")

# ---------------------------------------------------------------- libro
F = "Arial"
FB, FN, FH = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, bold=True, color="FFFFFF")
FIN, FNOTE, FTIT = Font(name=F, size=10, color="0000FF"), Font(name=F, size=9, italic=True, color="595959"), Font(name=F, size=14, bold=True)
HEAD, TOT = PatternFill("solid", fgColor="2F3A33"), PatternFill("solid", fgColor="F4F4F0")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
VERDE, AMBAR, ROJO = (PatternFill("solid", fgColor=c) for c in ("D9EAD3", "FFF2CC", "F4CCCC"))
wb = Workbook()


def encabezado(ws, fila, cols):
    for i, (n, w) in enumerate(cols, 1):
        c = ws.cell(fila, i, n)
        c.font, c.alignment, c.fill = FH, CEN, HEAD
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[fila].height = 46


# ------------------------------------------------ Pasadas
pa = wb.active
pa.title = "Pasadas"
cols_p = [("Suerte", 10), ("Inicio de la pasada", 12), ("Fin de la pasada", 12), ("Días que tomó", 8), ("Área regada (ha)", 9),
          ("Área neta (ha)", 9), ("Cobertura", 9), ("Días desde la pasada anterior", 11), ("¿Cuenta como intervalo?", 22),
          ("Año", 6), ("Mes", 5), ("Estado de la suerte", 26)]
encabezado(pa, 1, cols_p)
for k, r in enumerate(P.itertuples(index=False), 2):
    pa.cell(k, 1, r.ste)
    pa.cell(k, 2, r.ini.date())
    pa.cell(k, 3, r.fin.date())
    pa.cell(k, 4, f"=C{k}-B{k}+1")
    pa.cell(k, 5, round(float(r.area), 2))
    pa.cell(k, 6, round(float(r.an), 2) if pd.notna(r.an) else None)
    pa.cell(k, 7, f'=IF(N(F{k})=0,"",E{k}/F{k})')
    pa.cell(k, 8, int(r.dias) if pd.notna(r.dias) else None)
    pa.cell(k, 9, r.valido)
    pa.cell(k, 10, f"=YEAR(B{k})")
    pa.cell(k, 11, f"=MONTH(B{k})")
    pa.cell(k, 12, r.estado if r.estado not in ("nan", "None", "False") else None)
    for j, fmt in ((2, "dd/mm/yyyy"), (3, "dd/mm/yyyy"), (4, "0"), (5, "0.00"), (6, "0.00"), (7, "0%"), (8, "0")):
        pa.cell(k, j).number_format = fmt
    for j in range(1, 13):
        pa.cell(k, j).font = FIN if j in (1, 2, 3, 5, 6, 8, 9, 12) else FB
NP = len(P) + 1
pa.freeze_panes = "B2"
pa.auto_filter.ref = f"A1:L{NP}"
rng = lambda col: f"Pasadas!${col}$2:${col}${NP}"  # noqa: E731
S, I, D, V, A, MS = rng("A"), rng("B"), rng("H"), rng("I"), rng("J"), rng("K")

# ------------------------------------------------ Por suerte
ps = wb.create_sheet("Por suerte", 0)
cols_s = [("Suerte", 10), ("Hacienda", 18), ("Zona", 6), ("Área neta (ha)", 8), ("Variedad", 12), ("Corte de la caña en pie", 8),
          ("Inicio del ciclo actual", 11), ("Edad al último registro (meses)", 9),
          ("Pasadas 2026", 8), ("Días entre pasadas 2026: mediana", 10), ("Días entre pasadas 2026: promedio", 10),
          ("Días entre pasadas 2026: máximo", 10), ("% de intervalos de más de 30 días (2026)", 10),
          ("Días entre pasadas jun-sep 2026: mediana", 11), ("Días entre pasadas 2023-25: mediana", 10),
          ("Cambio 2026 frente a 2023-25 (días)", 10), ("Última pasada", 11), ("Días sin riego al último registro", 10),
          ("¿Se cortó después del último riego?", 10), ("TCH 2026", 8), ("Mes de cosecha 2026", 8)]
encabezado(ps, 1, cols_s)
FREF = "Resumen!$B$3"
for k, r in enumerate(M.itertuples(index=False), 2):
    s = r.ste
    ps.cell(k, 1, s)
    ps.cell(k, 2, r.finca)
    ps.cell(k, 3, int(r.zona) if str(r.zona).isdigit() else (r.zona if isinstance(r.zona, str) else None))
    ps.cell(k, 4, r.area if pd.notna(r.area) else None)
    ps.cell(k, 5, r.var if isinstance(r.var, str) else None)
    ps.cell(k, 6, int(r.corte) if pd.notna(r.corte) else None)
    ps.cell(k, 7, r.inicio if isinstance(r.inicio, date) else None)
    ps.cell(k, 8, f'=IF(ISNUMBER(G{k}),({FREF}-G{k})/30.44,"")')
    ps.cell(k, 9, f'=COUNTIFS({S},A{k},{A},2026)')
    ps.cell(k, 10, float(med26[s]) if s in med26 else None)
    ps.cell(k, 11, f'=IFERROR(AVERAGEIFS({D},{S},A{k},{A},2026,{V},"Sí"),"")')
    ps.cell(k, 12, f'=IF(COUNTIFS({S},A{k},{A},2026,{V},"Sí")=0,"",_xlfn.MAXIFS({D},{S},A{k},{A},2026,{V},"Sí"))')
    ps.cell(k, 13, f'=IF(COUNTIFS({S},A{k},{A},2026,{V},"Sí")=0,"",COUNTIFS({S},A{k},{A},2026,{V},"Sí",{D},">30")/COUNTIFS({S},A{k},{A},2026,{V},"Sí"))')
    ps.cell(k, 14, float(med_seca[s]) if s in med_seca else None)
    ps.cell(k, 15, float(med_ref[s]) if s in med_ref else None)
    ps.cell(k, 16, f'=IF(AND(ISNUMBER(J{k}),ISNUMBER(O{k})),J{k}-O{k},"")')
    ps.cell(k, 17, f'=IF(COUNTIFS({S},A{k})=0,"",_xlfn.MAXIFS({I},{S},A{k}))')
    ps.cell(k, 18, f'=IF(ISNUMBER(Q{k}),{FREF}-Q{k},"")')
    ps.cell(k, 19, f'=IF(AND(ISNUMBER(Q{k}),ISNUMBER(G{k})),IF(G{k}>Q{k},"Sí","No"),"")')
    if s in TCH.index:
        ps.cell(k, 20, round(float(TCH.loc[s, "tch_calc"]), 1))
        ps.cell(k, 21, int(TCH.loc[s, "Mes"]))
    for j, fmt in ((4, "0.00"), (7, "dd/mm/yyyy"), (8, "0.0"), (9, "0"), (10, "0"), (11, "0"), (12, "0"), (13, "0%"),
                   (14, "0"), (15, "0"), (16, "+0;-0;0"), (17, "dd/mm/yyyy"), (18, "0"), (20, "0.0"), (21, "0")):
        ps.cell(k, j).number_format = fmt
    for j in range(1, 22):
        ps.cell(k, j).font = FIN if j in (1, 2, 3, 4, 5, 6, 7, 10, 14, 15, 20, 21) else FB
NS = len(M) + 1
for col in ("J", "N"):
    ps.conditional_formatting.add(f"{col}2:{col}{NS}", CellIsRule(operator="greaterThan", formula=["45"], fill=ROJO))
    ps.conditional_formatting.add(f"{col}2:{col}{NS}", CellIsRule(operator="between", formula=["30.01", "45"], fill=AMBAR))
    ps.conditional_formatting.add(f"{col}2:{col}{NS}", CellIsRule(operator="between", formula=["0", "30"], fill=VERDE))
ps.conditional_formatting.add(f"P2:P{NS}", CellIsRule(operator="greaterThan", formula=["10"], fill=ROJO))
ps.freeze_panes = "C2"
ps.auto_filter.ref = f"A1:U{NS}"

# ------------------------------------------------ Resumen
rs = wb.create_sheet("Resumen", 0)
rs["A1"] = "Riopaila · frecuencia del riego por suerte"
rs["A1"].font = FTIT
rs["A3"], rs["B3"] = "Fecha del último registro de riego", REF
rs["A3"].font, rs["B3"].font, rs["B3"].number_format = FN, FIN, "dd/mm/yyyy"
rs["C3"] = "Las «días sin riego» y la edad se cuentan a esta fecha (fuente: Consolidado de riegos 2026)."
rs["C3"].font = FNOTE
cols_r = [("Hacienda", 22), ("Suertes", 8), ("Área neta (ha)", 10), ("Pasadas 2026", 9),
          ("Promedio de la mediana 2026 (días)", 12), ("Promedio de la mediana 2023-25 (días)", 12),
          ("Suertes con mediana 2026 > 45 días", 11), ("Suertes con mediana 2026 de 31 a 45 días", 11),
          ("Suertes en pie con más de 45 días sin riego", 12)]
encabezado(rs, 5, cols_r)
R_ = lambda col: f"'Por suerte'!${col}$2:${col}${NS}"  # noqa: E731
haciendas = [h for h in M.finca.dropna().unique()]
k = 6
for hac in haciendas:
    crit = f'{R_("B")},A{k}'
    rs.cell(k, 1, hac).font = FB
    rs.cell(k, 2, f"=COUNTIFS({crit})")
    rs.cell(k, 3, f"=SUMIFS({R_('D')},{crit})")
    rs.cell(k, 4, f"=SUMIFS({R_('I')},{crit})")
    rs.cell(k, 5, f'=IFERROR(AVERAGEIFS({R_("J")},{crit}),"")')
    rs.cell(k, 6, f'=IFERROR(AVERAGEIFS({R_("O")},{crit}),"")')
    rs.cell(k, 7, f'=COUNTIFS({crit},{R_("J")},">45")')
    rs.cell(k, 8, f'=COUNTIFS({crit},{R_("J")},">30",{R_("J")},"<=45")')
    rs.cell(k, 9, f'=COUNTIFS({crit},{R_("R")},">45",{R_("S")},"No")')
    for j, fmt in ((2, "0"), (3, "#,##0.0"), (4, "0"), (5, "0"), (6, "0"), (7, "0"), (8, "0"), (9, "0")):
        rs.cell(k, j).number_format, rs.cell(k, j).font = fmt, FB
    k += 1
rs.cell(k, 1, "Total Riopaila").font = FN
for j, col in ((2, "B"), (3, "C"), (4, "D"), (7, "G"), (8, "H"), (9, "I")):
    rs.cell(k, j, f"=SUM({col}6:{col}{k - 1})")
rs.cell(k, 5, f'=IFERROR(AVERAGE({R_("J")}),"")')
rs.cell(k, 6, f'=IFERROR(AVERAGE({R_("O")}),"")')
for j, fmt in ((2, "0"), (3, "#,##0.0"), (4, "0"), (5, "0"), (6, "0"), (7, "0"), (8, "0"), (9, "0")):
    rs.cell(k, j).number_format, rs.cell(k, j).font, rs.cell(k, j).fill = fmt, FN, TOT
rs.cell(k, 1).fill = TOT

# ------------------------------------------------ Cómo leer
cl = wb.create_sheet("Cómo leer")
cl.column_dimensions["A"].width = 120
txt = [
    ("Qué es una pasada", True),
    ("Una suerte se riega por partes durante uno o varios días. Los registros de la misma suerte separados por 7 días o menos "
     "se juntan en una sola pasada (dura en promedio 2 días y cubre la suerte completa). La hoja «Pasadas» trae cada una.", False),
    ("Frecuencia", True),
    ("«Días entre pasadas» = días entre el inicio de una pasada y el inicio de la siguiente en la misma suerte. No cuentan la "
     "primera pasada de cada suerte ni los intervalos en los que hubo una cosecha de la suerte en medio (columna «¿Cuenta como "
     "intervalo?»).", False),
    ("Las medianas (columnas J, N y O de «Por suerte») se calcularon en Python a partir de la hoja «Pasadas»; el resto son "
     "fórmulas sobre esa hoja. Colores: verde ≤30 días, amarillo 31-45, rojo >45; en el cambio frente a 2023-25, rojo si "
     "aumentó más de 10 días.", False),
    ("«Días sin riego al último registro» cuenta desde la última pasada hasta la fecha de la hoja Resumen. Si la suerte se "
     "cortó después de ese riego («¿Se cortó después del último riego?» = Sí), no es una suerte desatendida: es caña nueva.", False),
    ("Fuentes", True),
    ("Riegos: CONSOLIDADO DE RIEGOS 2026-.xlsx, hoja «Consolidado de riegos » (abr-2020 a oct-2026), sin las 412 filas repetidas "
     "idénticas. Suertes, hacienda, variedad, corte y fechas: maestro de suertes de Riopaila. TCH 2026: ZQM / G&P a "
     "septiembre y app de Cosecha en octubre. Azul = dato de la fuente; negro = fórmula.", False),
    ("Generado con frecuencia_excel.py (carpeta «Caida TCH 2026 - programas»).", False),
]
for i, (t, neg) in enumerate(txt, 1):
    c = cl.cell(i, 1, t)
    c.font, c.alignment = (FN if neg else FB), Alignment(wrap_text=True, vertical="top")
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("Por suerte", "Pasadas")
wb.active = 0
wb.save(SAL)
print(f"guardado: {SAL} · {len(M)} suertes · {len(P)} pasadas · último registro {REF}")
print("suertes con mediana 2026:", int(M.ste.isin(med26.index).sum()), "· sin pasadas en 2026:", int((~M.ste.isin(P[P.ini.dt.year == 2026].ste)).sum()))
