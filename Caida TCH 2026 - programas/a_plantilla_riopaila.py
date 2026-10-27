# -*- coding: utf-8 -*-
r"""Riopaila → plantilla del modelo parametrizado (repositorio modelo-caida-tch), con las mismas fuentes y limpiezas de la versión 3.

Lee lo que ya armó la cadena de la versión 3 (para que la comparación sea limpia) y las fuentes crudas donde hace falta:
  · cosechas ......... rio_factores_sep.pkl (ZQM 2011-25 + G&P 2026, un ciclo por suerte y corte, día 15 del mes).
                       Corte 0 → 1 (plantilla; 7 cosechas). Tipo «Mecanico» → mecanizada; lo demás → manual.
                       Sin variedad en el ZQM («#N/A», 38 cosechas) → «SIN DATO».
  · suertes .......... estación EMA de cada suerte (la del modelo), «TCH PPTO» del Maestro de Suertes 2026, centro de la suerte (Earth Engine).
  · clima_diario ..... EMA La Paila y Bugalagrande del repo Precipitacion-riopaila (lluvia, máxima y evaporación calculada);
                       en 2026 los días sin lluvia o máxima se llenan con el diario de Cenicaña (como factores_sep.py).
  · riegos ........... CONSOLIDADO DE RIEGOS 2026-.xlsx, con la limpieza de riego_2026.py (sin repetidos, área > 0, lámina fuera
                       de 20-300 mm → mediana).
  · labores .......... ciclos_labores.pkl (SimpleAgri): roturación del control oficial si el corte está a menos de 45 días (con su
                       oportunidad); si no, SimpleAgri − 10 días; herbicida SimpleAgri − desfase (malezas_aplicaciones.pkl);
                       fertilización SimpleAgri. Los días se cuentan desde el corte anterior que usa la plantilla
                       (fecha de cosecha − edad), así que cada labor queda a los mismos días del corte que midió la versión 3.
  · freatico_* ....... Registro de NF (freatico.json del portal Comité APS): 202 pozos y sus lecturas.
  · renovacion ....... no_renovadas_cosechas_2026.pkl: suertes del plan 2025 que siguieron como soca.
No lleva nombres de personas.

Uso: <tch-motor>\.venv-gee\Scripts\python.exe a_plantilla_riopaila.py <datos intermedios> <repo Precipitacion-riopaila>
       <Clima_Cenicana_...xlsx> <CONSOLIDADO DE RIEGOS 2026-.xlsx> <CONTROL DE ROTURACION.xlsm> <Maestro Suertes 2026.xlsx> <salida.xlsx>
"""
import json
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, r"C:\Users\Agr349\modelo-caida-tch")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
from modelo_tch.esquema import TABLAS  # noqa: E402
from modelo_tch.plantilla import escribir_libro  # noqa: E402
from tchmotor import gee  # noqa: E402

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
PK, REPO, CEN, RIEGOS, CTRL, MAESTRO, SAL = (Path(a) for a in sys.argv[1:8])
DIA = 30.44

# ---------------------------------------------------------------- cosechas
c = pd.read_pickle(PK / "rio_factores_sep.pkl")
cos = pd.DataFrame({
    "suerte": c.ste.astype(str), "fecha_cosecha": c.cosecha, "numero_corte": c.Corte.clip(lower=1).astype(int),
    "area_ha": c.Area, "toneladas": c["TC-Total"], "edad_meses": c.Edad,
    "tipo_cosecha": np.where(c["Tipo Cosecha"] == "Mecanico", "mecanizada", "manual"),
    "causa_cosecha": np.where(c["Causa Cosecha"] == "Incendio", "incendio", "programa"),
    "variedad": c["Variedad Principal"].astype(str), "sacarosa_pct": np.nan, "tenencia": "propia"})
sin_var = cos.variedad.str.strip().str.lower().isin(["nan", "none", "", "#n/a"])
cos.loc[sin_var, "variedad"] = "SIN DATO"          # en la versión 3 iban como texto «nan» y caían en «otra»
print(f"  cosechas sin variedad en el ZQM: {int(sin_var.sum())} → «SIN DATO»")
print(f"cosechas: {len(cos)} ({cos.fecha_cosecha.dt.year.min()}-{cos.fecha_cosecha.dt.year.max()}) · corte 0 → 1: {(c.Corte == 0).sum()}")

# ---------------------------------------------------------------- suertes
mae = list(openpyxl.load_workbook(MAESTRO, read_only=True, data_only=True)["MAESTRO"].iter_rows(min_row=5, values_only=True))
mae = pd.DataFrame(mae[1:], columns=[str(x).strip() if x is not None else f"c{j}" for j, x in enumerate(mae[0])])
mae = mae.loc[:, ~mae.columns.duplicated()]
mae = mae[mae.SUERTE.notna()].assign(ste=lambda d: d.SUERTE.astype(str).str.strip()).drop_duplicates("ste").set_index("ste")
ppto = pd.to_numeric(mae["TCH PPTO"], errors="coerce").where(lambda v: v.between(20, 300))
cen = gee.centroides_suertes("riopaila")
est = c.groupby("ste").est.first()
sue = pd.DataFrame({"suerte": est.index})
sue["hacienda"] = sue.suerte.map(mae.get("HACIENDA", pd.Series(dtype=object)))
sue["zona"] = sue.suerte.map(mae.get("ZONA", pd.Series(dtype=object))).astype("string")
sue["zona_agroecologica"] = np.nan
sue["area_neta_ha"] = np.nan
sue["tch_ppto"] = sue.suerte.map(ppto)
sue["fecha_siembra"] = pd.NaT
sue["estacion_clima"] = sue.suerte.map(est)
sue["pluviometro"] = np.nan
sue["latitud"] = [cen[s][1] if s in cen else np.nan for s in sue.suerte]
sue["longitud"] = [cen[s][0] if s in cen else np.nan for s in sue.suerte]
print(f"suertes: {len(sue)} · con PPTO {sue.tch_ppto.notna().sum()} · con coordenadas {sue.latitud.notna().sum()}")

# ---------------------------------------------------------------- clima diario (EMA + diario de Cenicaña en 2026)
wb = openpyxl.load_workbook(CEN, read_only=True, data_only=True)
it = wb["Diario 2026"].iter_rows(values_only=True)
M = pd.DataFrame(list(it)[1:], columns=next(wb["Diario 2026"].iter_rows(values_only=True)))
M["Fecha"] = pd.to_datetime(M.Fecha)
Mt = M.pivot_table(index="Fecha", columns="Código", values="Temperatura máxima (°C)")
Ml = M.pivot_table(index="Fecha", columns="Código", values="Lluvia (mm)")
e = pd.read_csv(REPO / "clima_estaciones.csv", sep=";")
e["f"] = pd.to_datetime(e.fc_obsrvcion)
dias = pd.date_range("2009-01-01", "2026-10-18", freq="D")
clima = []
for nom, cod in (("La Paila", "PAI"), ("Bugalagrande", "BLG")):
    d = e[e.nm_ema == nom].set_index("f").sort_index()
    d = d[~d.index.duplicated()].reindex(dias)
    f26 = d.index >= "2026-01-01"
    d.loc[f26 & d.tmp_mxma.isna(), "tmp_mxma"] = Mt[cod].reindex(d.index)[f26 & d.tmp_mxma.isna()]
    d.loc[f26 & d.prcptcion.isna(), "prcptcion"] = Ml[cod].reindex(d.index)[f26 & d.prcptcion.isna()]
    clima.append(pd.DataFrame({"estacion": nom, "fecha": d.index, "lluvia_mm": d.prcptcion.values, "temp_max_c": d.tmp_mxma.values,
                               "evaporacion_mm": d.evp_clclada.values}))
clima = pd.concat(clima, ignore_index=True)
clima = clima[clima[["lluvia_mm", "temp_max_c", "evaporacion_mm"]].notna().any(axis=1)]
for k, (lo, hi) in {"lluvia_mm": (0, 300), "temp_max_c": (10, 45), "evaporacion_mm": (0, 20)}.items():
    malo = ~clima[k].between(lo, hi) & clima[k].notna()
    if malo.any():
        print(f"  clima: {malo.sum()} valores de {k} fuera de {lo}-{hi} → vacíos")
        clima.loc[malo, k] = np.nan
print(f"clima: {len(clima)} días-estación {clima.fecha.min().date()} a {clima.fecha.max().date()}")

# ---------------------------------------------------------------- riegos (limpieza de riego_2026.py)
filas = list(openpyxl.load_workbook(RIEGOS, read_only=True, data_only=True)["Consolidado de riegos "].iter_rows(values_only=True))
h = next(i for i, r in enumerate(filas) if r and str(r[0]).strip() == "Fecha")
enc = [str(x).strip() if x is not None else f"c{j}" for j, x in enumerate(filas[h])]
R = pd.DataFrame(filas[h + 1:], columns=enc)
R = R[R.Fecha.notna()].copy()
R["f"] = pd.to_datetime(R.Fecha, errors="coerce")
R["ste"] = R["SEC-SUERTE"].astype(str).str.strip()
R["a"] = pd.to_numeric(R["ÁREA REGADA"], errors="coerce")
R["mm"] = pd.to_numeric(R["MM"], errors="coerce")
dup = R.duplicated(subset=[k for k in enc if not k.startswith("c")], keep="first")
R = R[~dup & R.f.notna() & (R.a > 0)].copy()
R["mm_ok"] = R.mm.where(R.mm.between(20, 300), R.mm.where(R.mm.between(20, 300)).median())
R = R[R.ste.isin(set(sue.suerte)) & (R.a <= 500)]
rie = pd.DataFrame({"suerte": R.ste, "fecha": R.f.dt.normalize(), "area_regada_ha": R.a.round(3), "lamina_mm": R.mm_ok.round(1)})
print(f"riegos: {len(rie)} ({rie.fecha.min().date()} a {rie.fecha.max().date()})")

# ---------------------------------------------------------------- labores (ciclos_labores.pkl + control oficial de roturación)
cl = pd.read_pickle(PK / "ciclos_labores.pkl").dropna(subset=["h_anio", "h_mes"])
lag = float(pd.read_pickle(PK / "malezas_aplicaciones.pkl")["lag"])
F = list(openpyxl.load_workbook(CTRL, read_only=True, data_only=True)["B.D. ROTURACIÓN"].iter_rows(values_only=True))
hh = next(i for i, x in enumerate(F) if x and str(x[0]).strip() == "SECTOR-SUERTE")
enc = [re.sub(r"\s+", " ", str(x)).strip() if x is not None else f"c{j}" for j, x in enumerate(F[hh])]
CT = pd.DataFrame([x[: len(enc)] for x in F[hh + 1:]], columns=enc).loc[:, lambda d: ~d.columns.duplicated()]
CT = CT[CT["SECTOR-SUERTE"].notna()].copy()
CT["ste"] = CT["SECTOR-SUERTE"].astype(str).str.strip()
CT["corte_f"] = pd.to_datetime(CT["FECHA DE CORTE"], errors="coerce")
txt = lambda s: s.astype(str).str.strip().str.upper().replace({"NONE": "", "NAN": "", "#N/A": ""})  # noqa: E731
CT["op1"], CT["op1p"], CT["est1"] = txt(CT["OPORTUNIDAD 1RA LABOR"]), txt(CT["OPORTUNIDAD PENDIENTES PRIMERA"]), txt(CT["ESTADO PRIMERA LABORE"])
CT["dias1"] = pd.to_numeric(CT["DIAS ENTRE PROGRAMACION Y ROTURACION"], errors="coerce")
CT["rot"] = np.select([CT.op1 == "TARDIO", CT.op1.isin(["IDEAL", "NORMAL"]), (CT.est1 == "REALIZADO") & (CT.dias1 <= 0), CT.op1p == "TARDIO"],
                      ["tardia", "a_tiempo", "a_tiempo", "tardia"], "")
CT = CT[(CT.rot != "") & CT.corte_f.notna()]
CG = {s: g for s, g in CT.groupby("ste")}


def rot_control(s, f):
    g = CG.get(s)
    if g is None or pd.isna(f):
        return ""
    dd = (g.corte_f - f).dt.days.abs()
    return g.rot.loc[dd.idxmin()] if dd.min() <= 45 else ""


cl["h_cos"] = pd.to_datetime(dict(year=cl.h_anio.astype(int), month=cl.h_mes.astype(int), day=15))
cl["ini"] = cl.h_cos - pd.to_timedelta(cl.h_edad * DIA, unit="D")       # el corte anterior como lo cuenta la plantilla
cl["rot_ctrl"] = [rot_control(s, f) for s, f in zip(cl.ste, cl.corte_f)]
LAB = []
dia = lambda ini, d: (ini + pd.Timedelta(days=float(d))).normalize() + pd.Timedelta(days=1 if (ini + pd.Timedelta(days=float(d))).hour else 0)  # noqa: E731
for r in cl.itertuples():
    if r.rot_ctrl:
        d_ = r.d_rot - 10 if pd.notna(r.d_rot) else (40 if r.rot_ctrl == "tardia" else 15)
        LAB.append({"suerte": r.ste, "labor": "roturacion", "fecha": dia(r.ini, max(d_, 1)), "oportunidad": r.rot_ctrl})
    elif pd.notna(r.d_rot):
        LAB.append({"suerte": r.ste, "labor": "roturacion", "fecha": dia(r.ini, max(r.d_rot - 10, 1)), "oportunidad": np.nan})
    if pd.notna(r.d_maleza):
        LAB.append({"suerte": r.ste, "labor": "herbicida", "fecha": dia(r.ini, max(r.d_maleza - lag, 1)), "oportunidad": np.nan})
    if pd.notna(r.d_fert):
        LAB.append({"suerte": r.ste, "labor": "fertilizacion", "fecha": dia(r.ini, max(r.d_fert, 1)), "oportunidad": np.nan})
lab = pd.DataFrame(LAB)
print(f"labores: {len(lab)} ({lab.labor.value_counts().to_dict()}) · roturación del control: {(cl.rot_ctrl != '').sum()} de {len(cl)} ciclos")

# ---------------------------------------------------------------- nivel freático
fr = json.load(open(PK / "freatico_luis" / "freatico.json", encoding="utf-8"))
pz = pd.DataFrame([{"pozo": str(p[0]), "latitud": p[1], "longitud": p[2]} for p in fr["pozos"]])
le = pd.DataFrame([{"pozo": str(k), "fecha": pd.Timestamp(f), "profundidad_m": v} for f, d in fr["lecturas"].items() for k, v in d.items() if v is not None])
le["profundidad_m"] = pd.to_numeric(le.profundidad_m.replace({"S": 2, "s": 2}), errors="coerce").clip(upper=2)   # S = pozo seco (más de 2 m)
le = le[le.pozo.isin(set(pz.pozo)) & le.profundidad_m.between(0, 10)]
print(f"freático: {len(pz)} pozos · {len(le)} lecturas {le.fecha.min().date()} a {le.fecha.max().date()}")

# ---------------------------------------------------------------- renovación
nr = pd.read_pickle(PK / "no_renovadas_cosechas_2026.pkl")
ren = pd.DataFrame({"suerte": sorted(nr[nr.no_renov].ste.astype(str).unique()), "anio_plan": 2025, "renovada": "no"})
print(f"renovación: {len(ren)} suertes del plan 2025 sin renovar")

datos = {"cosechas": cos, "suertes": sue, "clima_diario": clima, "riegos": rie, "labores": lab, "freatico_pozos": pz,
         "freatico_lecturas": le, "renovacion": ren}
for t in TABLAS:
    d = datos.get(t.nombre, pd.DataFrame())
    for col in t.columnas:
        if col.nombre not in d:
            d[col.nombre] = np.nan
    datos[t.nombre] = d[[col.nombre for col in t.columnas]].reset_index(drop=True)
SAL.parent.mkdir(parents=True, exist_ok=True)
escribir_libro(SAL, datos)
print("escrito:", SAL)
