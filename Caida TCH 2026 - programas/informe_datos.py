# -*- coding: utf-8 -*-
"""Datos del informe navegable «Caída del TCH Riopaila 2026» (tablero_web): cada causa con sus suertes y cada suerte con
su ficha, para no tener que abrir el Excel a buscar la justificación.

Entradas (todas ya calculadas por la cadena):
  tablero_v3/suertes_caida_tch.csv (tablero_datos.py) · rio_integrado_dataset.pkl y caida_por_suerte.pkl
  (caida_integrada_rio.py) · freatico_suertes_lecturas.csv (freatico_suertes.py) · validacion_*.csv (validacion_modelo.py)
  · coeficientes del soporte v3 · Maestro Suertes 2026 (fecha exacta del corte) · clima diario de las EMA (repo
  Precipitacion-riopaila + Meteoportal 2026) · NDVI/NDMI mensual por suerte (tch-motor/data/gee/mensual_riopaila.csv.gz).
Lluvia de la ficha: la EMA que usa el modelo y, al lado, el pluviómetro más cercano (precipitacion.csv.csv del mismo repo).
Salidas en <carpeta web>/data: resumen.json, causas.json, suertes.json, fichas.json, validacion.json.

Uso: <tch-motor>\\.venv-gee\\Scripts\\python.exe informe_datos.py <datos intermedios> <soporte_v3.xlsx>
       <Maestro Suertes 2026.xlsx> <repo Precipitacion-riopaila> <Clima_Cenicana_...xlsx> <mensual_riopaila.csv.gz> <carpeta web>
"""
import json
import shutil
import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
from tchmotor import gee  # noqa: E402  (solo para los centroides de las suertes)

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
PK, SOP, MAESTRO, REPO, CEN, SAT, WEB = (Path(a) for a in sys.argv[1:8])
DATA = WEB / "data"
DATA.mkdir(parents=True, exist_ok=True)
wm = lambda g, k: float(np.average(g[k], weights=g.Area)) if len(g) and g.Area.sum() > 0 else None  # noqa: E731
es = lambda v, d=1: f"{v:+.{d}f}".replace("-", "\u2212").replace(".", ",")  # noqa: E731
ep = lambda p: "p < 0,001" if p < 0.001 else "p = " + f"{p:.2f}".replace(".", ",")  # noqa: E731
r1 = lambda v, n=1: None if v is None or (isinstance(v, float) and not np.isfinite(v)) or pd.isna(v) else round(float(v), n)  # noqa: E731


def limpio(o):
    """NaN e infinitos → null (el JSON del navegador no los acepta)."""
    if isinstance(o, dict):
        return {k: limpio(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [limpio(v) for v in o]
    if isinstance(o, float) and not np.isfinite(o):
        return None
    if isinstance(o, (np.floating,)):
        return None if not np.isfinite(o) else float(o)
    if isinstance(o, (np.integer,)):
        return int(o)
    return o


def guardar(obj, nombre, **kw):
    json.dump(limpio(obj), open(DATA / nombre, "w", encoding="utf-8"), ensure_ascii=False, allow_nan=False,
              **(kw or {"separators": (",", ":")}))

ds = pd.read_pickle(PK / "rio_integrado_dataset.pkl")
x = pd.read_pickle(PK / "caida_por_suerte.pkl")
S = pd.read_csv(PK / "tablero_v3" / "suertes_caida_tch.csv")
coef = pd.read_excel(SOP, sheet_name="Coeficientes").set_index("variable")
PROB = pd.read_excel(SOP, sheet_name="Probadas sin efecto").set_index("variable")  # lluvia en roturación y fertilización
x["n"] = x.groupby(["ste", "Mes"]).cumcount()
x["clave"] = x.ste + " · mes " + x.Mes.astype(int).astype(str) + np.where(x.n > 0, " (" + (x.n + 1).astype(str) + ")", "")
q = ds.loc[x.index].copy()
q["clave"] = x.clave
S = S.set_index("clave").loc[q.clave].reset_index()
S.index = q.index
ref = ds[ds.anio.between(2020, 2025) & ds.Mes.between(1, 9)]
ds["prev_tch"] = ds.sort_values(["ste", "cosecha"]).groupby("ste").tch_calc.shift()
q["prev_tch"] = ds.loc[q.index, "prev_tch"]

# ---------------------------------------------------------------- fecha exacta del corte (Maestro: FECHA DE ULTIMO CORTE)
M = list(openpyxl.load_workbook(MAESTRO, read_only=True, data_only=True)["MAESTRO"].iter_rows(min_row=5, values_only=True))
pm = pd.DataFrame(M[1:], columns=[str(v).strip() if v is not None else f"c{j}" for j, v in enumerate(M[0])])
pm = pm.loc[:, ~pm.columns.duplicated()]
pm = pm[pm.SUERTE.notna()].assign(ste=lambda t: t.SUERTE.astype(str).str.strip()).drop_duplicates("ste").set_index("ste")
fuc = pd.to_datetime(pm["FECHA DE ULTIMO CORTE"], errors="coerce")
q["fecha_corte"] = [f.date().isoformat() if pd.notna(f := fuc.get(s)) and f.year == 2026 and f.month == m else None
                    for s, m in zip(q.ste, q.Mes)]
print(f"fecha exacta del corte en {q.fecha_corte.notna().sum()} de {len(q)} cosechas")

# ---------------------------------------------------------------- normales por mes de cosecha (para leer cada suerte)
def normal_mes(col, d):
    return d.dropna(subset=[col]).groupby("Mes").apply(lambda g: wm(g, col)).to_dict()


ref_agua = ds[ds.anio.between(2022, 2025) & ds.Mes.between(1, 9) & (ds.def_na == 0)]
ref_nf = ds[ds.anio.between(2022, 2025) & ds.Mes.between(1, 9) & (ds.nf_na == 0)]
N_D32, N_LLU = normal_mes("d32", ref), normal_mes("llu03", ref)
N_DEFSR, N_MM, N_DEF = normal_mes("defsr_fin", ref_agua), normal_mes("mm_fin", ref_agua), normal_mes("def_fin", ref_agua)
N_NF = normal_mes("nf_ciclo", ref_nf)

# ---------------------------------------------------------------- causas: nombre, columna de aporte, grupo, texto
CAUSAS = [
    {"id": "mecanica", "nombre": "Cosecha mecanizada", "col": "Cosecha mecanizada", "grupo": "Agronómico y operativo",
     "mide": "Si la cosecha de 2026 fue mecanizada: la cosechadora daña la cepa y genera despoblación.",
     "efecto": f"En 2026 la cosecha mecanizada rindió {es(coef.loc['mec_2026', 'efecto %'])} % de TCH frente a la manual en la misma "
               f"suerte (en 2020-25 la diferencia fue de −2 a −4 %)."},
    {"id": "corte-anterior", "nombre": "Corte anterior mecanizado", "col": "Corte anterior mecanizado", "grupo": "Agronómico y operativo",
     "mide": "Si el corte que dejó esta soca fue mecanizado: la cepa queda aporreada y rebrota menos.",
     "efecto": f"Soca que sale de un corte mecanizado después de manual: {es(coef.loc['pm_new', 'efecto %'])} % de TCH; de dos "
               f"cortes mecanizados seguidos: {es(coef.loc['pm_seg', 'efecto %'])} % (2011-2026)."},
    {"id": "corte-humedo", "nombre": "Cosecha con suelo húmedo", "col": "Corte anterior en suelo húmedo",
     "grupo": "No controlable",
     "mide": "Lluvia en el mes del corte anterior en la estación de la suerte. Si el corte se hizo con el suelo húmedo, la "
             "cepa queda más golpeada y la soca rebrota menos; se toma como húmedo un mes con 100 mm o más.",
     "efecto": f"Por cada 100 mm de lluvia en el mes del corte anterior: {es(coef.loc['ca_mm100', 'efecto %'])} % de TCH "
               f"({ep(coef.loc['ca_mm100', 'p'])}). La diferencia entre corte manual y mecanizado no fue confiable."},
    {"id": "calor", "nombre": "Temperatura mayor a 32 °C 30 días antes de cosecha", "col": "Días Tmax >32 °C", "grupo": "No controlable",
     "mide": "Días con temperatura máxima por encima de 32 °C en los 30 días antes del corte (EMA La Paila o Bugalagrande).",
     "efecto": f"Cada día por encima de 32 °C: {es(coef.loc['d32', 'efecto %'], 2)} % de TCH."},
    {"id": "sequia", "nombre": "Sequía al final del ciclo", "col": "Balance hídrico: sequía al final", "grupo": "No controlable",
     "mide": "Balance hídrico de los últimos meses del ciclo: evaporación menos lluvia (déficit sin riego, mm).",
     "efecto": f"Cada 100 mm de déficit que queda después del riego: {es(coef.loc['def100', 'efecto %'])} % de TCH."},
    {"id": "riego", "nombre": "Riego aplicado", "col": "Balance hídrico: riego aplicado", "grupo": "Flujo de caja",
     "mide": "Lámina de riego aplicada al final del ciclo (mm), que reduce el déficit.",
     "efecto": f"Cada 100 mm de riego evitan {es(-coef.loc['def100', 'efecto %'])[1:]} % de pérdida de TCH (mismo coeficiente del déficit)."},
    {"id": "freatico", "nombre": "Nivel freático", "col": "Nivel freático", "grupo": "No controlable",
     "mide": "Nivel promedio del agua freática en la suerte entre el corte anterior y la cosecha (pozos del Registro de NF, "
             "interpolados a la suerte).",
     "efecto": f"Por cada 10 cm que el nivel promedio estuvo por encima de 1,5 m: {es(coef.loc['nf_exc15', 'efecto %'] / 10, 2)} % de TCH."},
    {"id": "lluvia", "nombre": "Lluvia del arranque", "col": "Precipitación 0 a 3 meses", "grupo": "No controlable",
     "mide": "Lluvia de los primeros 3 meses del ciclo (EMA), cuando rebrota la soca.",
     "efecto": f"Cada 100 mm de más en los meses 0-3: {es(coef.loc['llu03_100', 'efecto %'])} % de TCH."},
    {"id": "variedades", "nombre": "Variedades", "col": "Variedades", "grupo": "Agronómico y operativo",
     "mide": "La variedad sembrada, frente a la mezcla de variedades de 2020-25.",
     "efecto": "Cada variedad tiene su propio efecto, medido en la misma suerte a lo largo de los años."},
    {"id": "edad", "nombre": "Edad de cosecha", "col": "Edad de cosecha", "grupo": "Agronómico y operativo",
     "mide": "Meses de edad al cortar, frente a la edad normal de 2020-25.",
     "efecto": "Cortar más viejo da más toneladas: 2026 se cortó con más edad y eso sumó."},
    {"id": "renovacion", "nombre": "Renovación no hecha", "col": "No renovación", "grupo": "Flujo de caja",
     "mide": "Suertes del plan de renovación 2025 que siguieron como soca.",
     "efecto": f"Suerte que debía renovarse y no se renovó: {es(coef.loc['no_renov', 'efecto %'])} % de TCH (poco seguro, {ep(coef.loc['no_renov', 'p'])})."},
    {"id": "roturacion", "nombre": "Oportunidad de roturación", "col": "Oportunidad roturación", "grupo": "Flujo de caja",
     "mide": "Roturación fuera de oportunidad (control de roturación; si no está, SimpleAgri).",
     "efecto": f"Roturación tardía: {es(coef.loc['rot_tard', 'efecto %'])} % de TCH ({ep(coef.loc['rot_tard', 'p'])}). "
               f"Roturar con el suelo húmedo (lluvia en los 5 días antes) no mostró efecto confiable: "
               f"{es(PROB.loc['Roturación en suelo húmedo', 'efecto %'])} % por cada 10 mm ({ep(PROB.loc['Roturación en suelo húmedo', 'p'])})."},
    {"id": "malezas", "nombre": "Control de malezas", "col": "Control de malezas", "grupo": "Flujo de caja",
     "mide": "Días del corte al primer herbicida; cuenta el atraso después de 90 días.",
     "efecto": f"Por mes de atraso después de 90 días: {es(coef.loc['mal_extra30', 'efecto %'])} % de TCH ({ep(coef.loc['mal_extra30', 'p'])})."},
    {"id": "fertilizacion", "nombre": "Oportunidad de fertilización", "col": "Oportunidad fertilización", "grupo": "Flujo de caja",
     "mide": "Días del corte a la fertilización; cuenta si pasó de 120 días.",
     "efecto": f"Fertilización después de 120 días: {es(coef.loc['fert_120', 'efecto %'])} % de TCH ({ep(coef.loc['fert_120', 'p'])}). "
               f"Fertilizar con el suelo húmedo no mostró efecto confiable: "
               f"{es(PROB.loc['Fertilización en suelo húmedo', 'efecto %'])} % por cada 10 mm ({ep(PROB.loc['Fertilización en suelo húmedo', 'p'])})."},
    {"id": "corte", "nombre": "Número de corte (socas más jóvenes en 2026)", "col": "N° de corte", "grupo": "Agronómico y operativo",
     "mide": "Cuántos cortes lleva la soca, frente a la historia 2020-25 de la misma suerte. Cada corte le quita productividad a la "
             "soca; en 2026 se cosecharon socas más jóvenes que lo normal (más segundos cortes y menos de 6 o más), por eso este "
             "factor sumó.",
     "efecto": "Frente a la plantilla, en la misma suerte: " + ", ".join(
         f"{r.corte.replace(' o más', '')} corte{' o más' if 'más' in r.corte else ''} {es(r.efecto_pct)} %" for r in pd.read_pickle(PK / "efecto_corte.pkl")["efecto"].itertuples()) + "."},
    {"id": "incendio", "nombre": "Incendio", "col": "Incendio", "grupo": "Agronómico y operativo",
     "mide": "Cosecha forzada por incendio.", "efecto": f"Cosecha por incendio: {es(coef.loc['inc', 'efecto %'])} % de TCH."},
]
ok = S.tch_hist.notna()
w = S.area[ok]
for k in CAUSAS:
    k["t_ha"] = r1(np.average(S.loc[ok, k["col"]], weights=w), 2)
sin_exp = r1(np.average(S.loc[ok, "Sin explicar"], weights=w), 2)
HA, TR, TH = float(w.sum()), float(np.average(S.tch_real[ok], weights=w)), float(np.average(S.tch_hist[ok], weights=w))

# ---------------------------------------------------------------- una fila por cosecha 2026 con sus datos de soporte
# hacienda de respaldo (cartografía de tablones) para las suertes que no están en el Maestro 2026
HAC_GEO = {f["properties"]["sec_ste"]: str(f["properties"].get("hacienda") or "").strip().title()
           for f in gee._features("riopaila") if f["properties"].get("sec_ste")}
fil = []
for i in q.index:
    r, s = q.loc[i], S.loc[i]
    m = int(r.Mes)
    reg_lab = r.Corte >= 2 and isinstance(r.rot, str)
    fil.append({
        "k": s.clave, "s": r.ste, "h": s.hacienda if isinstance(s.hacienda, str) and s.hacienda not in ("", "None", "Nan") else HAC_GEO.get(r.ste, ""),
        "z": None if pd.isna(s.zona) else int(s.zona),
        "za": s.zona_agroecologica if isinstance(s.zona_agroecologica, str) else "", "m": m, "f": r.fecha_corte,
        "c": int(r.Corte), "v": s.variedad, "a": r1(r.Area, 2), "ed": r1(r.Edad), "tipo": s.cosecha,
        "ant": r.ant, "ant_f": None if pd.isna(r.prev_anio) else f"{int(r.prev_anio)}-{int(r.prev_mes):02d}", "ant_tch": r1(r.prev_tch),
        "real": r1(s.tch_real), "hist": r1(s.tch_hist), "ppto": r1(s.ppto), "caida": r1(s.caida, 2),
        "inc": int(r.inc), "nr": int(r.no_renov), "ema": r.est,
        "d32": r1(r.d32, 0), "d32n": r1(N_D32.get(m), 1),
        "llu": r1(r.llu03, 0), "llun": r1(N_LLU.get(m), 0),
        "dsr": None if r.def_na else r1(r.defsr_fin, 0), "dsrn": r1(N_DEFSR.get(m), 0),
        "rie": None if r.def_na else r1(r.mm_fin, 0), "rien": r1(N_MM.get(m), 0),
        "def": None if r.def_na else r1(r.def_fin, 0), "defn": r1(N_DEF.get(m), 0),
        "nf": r1(r.nf_ciclo, 2), "nf45": r1(r.nf_45, 2), "nfn": r1(N_NF.get(m), 2),
        "mmca": r1(r.mm_ca, 0), "mmrot": r1(r.mm_rot, 0) if reg_lab else None, "mmfert": r1(r.mm_fert, 0) if reg_lab else None,
        "rot": r.rot if reg_lab else None, "rot_de": r.rot_de if reg_lab else None,
        "dmal": None if not reg_lab or pd.isna(r.d_mal) else int(r.d_mal),
        "dfert": None if not reg_lab or pd.isna(r.d_fert) else int(r.d_fert),
        "ap": {k["id"]: r1(s[k["col"]], 2) for k in CAUSAS}, "se": r1(s["Sin explicar"], 2),
        "fp": s.factor_principal if isinstance(s.factor_principal, str) else "",
    })
# las que no están en el Maestro ni en la cartografía (re-suerteo): la hacienda de su sector
SEC_H = pd.Series({f["s"]: f["h"] for f in fil if f["h"]}).groupby(lambda s: s[:4]).agg(lambda v: v.mode().iloc[0]).to_dict()
for f in fil:
    f["h"] = f["h"] or SEC_H.get(f["s"][:4], "")
# rebrote visto por el satélite antes y después del primer corte mecanizado (mecanizada_satelite.py)
SATD = PK / "satelite_mecanizada"
SAT_IM, SAT_RS, SAT_PS = None, None, None
if (SATD / "satelite_resumen.csv").exists():
    SAT_RS = pd.read_csv(SATD / "satelite_resumen.csv")
    SAT_PS = pd.read_csv(SATD / "satelite_por_suerte.csv")
    SAT_IM = pd.read_csv(SATD / "satelite_imagenes.csv") if (SATD / "satelite_imagenes.csv").exists() else pd.DataFrame()
    shutil.rmtree(WEB / "img" / "sat", ignore_errors=True)
    (WEB / "img" / "sat").mkdir(parents=True)
    for x in SAT_IM.itertuples():
        shutil.copy(SATD / "imagenes" / x.imagen, WEB / "img" / "sat" / x.imagen)
    imgs = {(x.grupo, x.ste): x for x in SAT_IM.itertuples()}
    for f in fil:
        for gr in ("A", "B"):
            x = imgs.get((gr, f["s"]))
            f["sat" + gr] = None if x is None else {"img": f"img/sat/{x.imagen}", "edad": int(x.edad), "f1": str(x.fecha1),
                                                   "f2": str(x.fecha2), "ndvi1": r1(x.ndvi1, 2), "ndvi2": r1(x.ndvi2, 2),
                                                   "rala1": r1(100 * x.rala1, 0), "rala2": r1(100 * x.rala2, 0)}
guardar(fil, "suertes.json")
print(f"suertes.json: {len(fil)} cosechas")

# ---------------------------------------------------------------- series de clima diario (las mismas del modelo)
dias = pd.date_range("2009-01-01", "2026-10-18", freq="D")
wb = openpyxl.load_workbook(CEN, read_only=True, data_only=True)
it = wb["Diario 2026"].iter_rows(values_only=True)
enc = next(it)
Mt = pd.DataFrame(list(it), columns=enc)
Mt["Fecha"] = pd.to_datetime(Mt.Fecha)
Ml = Mt.pivot_table(index="Fecha", columns="Código", values="Lluvia (mm)")
Mx = Mt.pivot_table(index="Fecha", columns="Código", values="Temperatura máxima (°C)")
e = pd.read_csv(REPO / "clima_estaciones.csv", sep=";")
e["f"] = pd.to_datetime(e.fc_obsrvcion)
EMA = {}
for est, cod in (("La Paila", "PAI"), ("Bugalagrande", "BLG")):
    d = e[e.nm_ema == est].set_index("f").sort_index()
    d = d[~d.index.duplicated()].reindex(dias)
    falta = d.index >= "2026-01-01"
    d.loc[falta & d.prcptcion.isna(), "prcptcion"] = Ml[cod].reindex(d.index)[falta & d.prcptcion.isna()]
    d.loc[falta & d.tmp_mxma.isna(), "tmp_mxma"] = Mx[cod].reindex(d.index)[falta & d.tmp_mxma.isna()]
    EMA[est] = d
LLM = {k: v.prcptcion.resample("MS").sum(min_count=20) for k, v in EMA.items()}  # mm por mes
CAL = {k: (v.tmp_mxma > 32).resample("MS").sum().where(v.tmp_mxma.resample("MS").count() >= 20) for k, v in EMA.items()}
NLL = {k: v[(v.index.year >= 2011) & (v.index.year <= 2025)].groupby(v.index[(v.index.year >= 2011) & (v.index.year <= 2025)].month).mean()
       for k, v in LLM.items()}
NCA = {k: v[(v.index.year >= 2011) & (v.index.year <= 2025)].groupby(v.index[(v.index.year >= 2011) & (v.index.year <= 2025)].month).mean()
       for k, v in CAL.items()}

# pluviómetros (repo Precipitacion-riopaila): lluvia diaria de cada PLUV; a cada suerte, el más cercano a menos de 3 km
pl = pd.read_csv(REPO / "precipitacion.csv.csv", sep=";", encoding="latin-1", usecols=[1, 2, 3])
pl.columns = ["PLUV", "Fecha", "mm"]
pl["Fecha"] = pd.to_datetime(pl.Fecha, format="%d/%m/%Y", errors="coerce")
pl["mm"] = pd.to_numeric(pl.mm.astype(str).str.replace(",", "."), errors="coerce")
PMm = pl.dropna().pivot_table(index="Fecha", columns="PLUV", values="mm", aggfunc="mean").resample("MS").sum(min_count=25)
PN = PMm[PMm.index.year <= 2025].groupby(PMm.index[PMm.index.year <= 2025].month).mean()
gp = json.load(open(r"C:\Users\Agr349\agrocontrol-campo\public\data\contexto_pluviometros.geojson", encoding="utf-8"))["features"]
GP = pd.DataFrame([(f["properties"]["Pluviometr"], *f["geometry"]["coordinates"]) for f in gp], columns=["PLUV", "lon", "lat"])
GP = GP[GP.PLUV.isin(PMm.columns)].set_index("PLUV")
CEN_S = gee.centroides_suertes("riopaila")


def pluv_de(ste):
    c_ = CEN_S.get(ste)
    if c_ is None or GP.empty:
        return None, None
    d = np.hypot((GP.lon - c_[0]) * 111320 * np.cos(np.radians(c_[1])), (GP.lat - c_[1]) * 110570)
    return (int(d.idxmin()), float(d.min())) if d.min() <= 3000 else (None, None)


# ---------------------------------------------------------------- fichas
NF = pd.read_csv(PK / "freatico_suertes_lecturas.csv", parse_dates=["fecha"])
NFG = {s: g.sort_values("fecha") for s, g in NF.groupby("ste")}
sat = pd.read_csv(SAT).rename(columns={"sec_ste": "ste"})
sat["f"] = pd.to_datetime(dict(year=sat.anio, month=sat.mes, day=15))
sat = sat[sat.despejado >= 0.6]
SG = {s: g.set_index("f") for s, g in sat.groupby("ste")}


def curva(ste, ini, fin):
    g = SG.get(ste)
    if g is None:
        return None
    out = []
    for e_ in range(0, 15):
        f = (ini + pd.DateOffset(months=e_)).to_period("M").to_timestamp() + pd.Timedelta(days=14)
        if f > fin + pd.Timedelta(days=15):
            break
        if f in g.index:
            out.append([e_, r1(g.loc[f, "NDVI"], 3), r1(g.loc[f, "NDMI"], 3)])
    return out or None


fichas = {}
for i in q.index:
    r = q.loc[i]
    h = ds[ds.ste == r.ste].sort_values("cosecha")
    hist = [[int(a.anio), int(a.Mes), int(a.Corte), "Mec" if a.mec == 1 else "Man", r1(a.tch_calc), r1(a.Edad)]
            for a in h.itertuples()]
    ini, fin = r.ini, r.cosecha
    meses = pd.date_range(ini.to_period("M").to_timestamp(), fin.to_period("M").to_timestamp(), freq="MS")
    est = r.est
    pv, pdist = pluv_de(r.ste)
    lluvia = [[f.strftime("%Y-%m"), r1(LLM[est].get(f), 0), r1(NLL[est].get(f.month), 0), r1(CAL[est].get(f), 0),
               r1(NCA[est].get(f.month), 1), r1(PMm[pv].get(f), 0) if pv else None, r1(PN[pv].get(f.month), 0) if pv else None]
              for f in meses]
    g = NFG.get(r.ste)
    fre = None
    if g is not None:
        gg = g[(g.fecha >= ini - pd.Timedelta(days=120)) & (g.fecha <= fin + pd.Timedelta(days=20))]
        fre = {"pozo": int(gg.pozo_cercano.mode().iloc[0]) if len(gg) else None,
               "dist": r1(gg.dist_m.median(), 0) if len(gg) else None,
               "l": [[f.date().isoformat(), r1(v, 2)] for f, v in zip(gg.fecha, gg.nivel)]}
    prev = h[h.cosecha < r.cosecha].tail(2)
    curvas = {"2026": curva(r.ste, ini, fin)}
    for p in prev.itertuples():
        curvas[f"{p.anio}"] = curva(r.ste, p.ini, p.cosecha)
    fichas[q.loc[i, "clave"]] = {"hist": hist, "ini": ini.date().isoformat(), "lluvia": lluvia, "fre": fre,
                                 "pluv": {"id": pv, "dist": r1(pdist, 0)} if pv else None,
                                 "sat": {k: v for k, v in curvas.items() if v},
                                 "sat_tipo": {f"{p.anio}": ("Mec" if p.mec == 1 else "Man") for p in prev.itertuples()}}
guardar(fichas, "fichas.json")
print(f"fichas.json: {len(fichas)} · {(DATA / 'fichas.json').stat().st_size / 1e6:.2f} MB")

# ---------------------------------------------------------------- soporte de cada causa (gráficas del informe)
sop = {}
a = ds.groupby("anio").apply(lambda g: pd.Series({"pct": 100 * (g.Area * g.mec).sum() / g.Area.sum()}))
sop["mecanica"] = {"anios": [[int(y), r1(v.pct)] for y, v in a.iterrows()],
                   "brecha": [[y, r1(coef.loc[f"mec_{y}", "efecto %"])] for y in range(2020, 2027)]}
hist_s = ds[ds.anio.between(2015, 2025)].groupby("ste").apply(lambda g: wm(g, "tch_calc"))
qq = q.assign(h=q.ste.map(hist_s)).dropna(subset=["h"])
ORD = ["manual", "mecánica tras manual", "mecánica seguida", "plantilla o sin dato"]
sop["corte-anterior"] = {"tabla": [[k, int(len(g)), r1(g.Area.sum(), 0), r1(100 * (wm(g, "tch_calc") / wm(g, "h") - 1))]
                                   for k in ORD if len(g := qq[qq.ant == k])]}
# satélite: diferencia de NDVI y NDMI por mes de edad (efectos fijos de suerte y del mes observado)
cur = []
for r in ds[ds.cosecha >= "2019-06-01"].itertuples():
    g = SG.get(r.ste)
    if g is None:
        continue
    for e_ in range(1, 13):
        f = (r.ini + pd.DateOffset(months=e_)).to_period("M").to_timestamp() + pd.Timedelta(days=14)
        if f in g.index and f <= r.cosecha:
            cur.append((r.ste, e_, f, g.loc[f, "NDVI"], g.loc[f, "NDMI"], r.ant))
cur = pd.DataFrame(cur, columns=["ste", "edad", "f", "NDVI", "NDMI", "ant"]).dropna()
cur["ym"] = cur.f.dt.to_period("M").astype(str)
sat_ed = []
for e_ in range(1, 10):
    d = cur[cur.edad == e_]
    row = [e_]
    for ind in ("NDVI", "NDMI"):
        mo = smf.ols(f"{ind} ~ C(ste) + C(ym) + C(ant, Treatment('manual'))", data=d).fit()
        row += [r1(mo.params.get(f"C(ant, Treatment('manual'))[T.{k}]"), 3) for k in ("mecánica tras manual", "mecánica seguida")]
    sat_ed.append(row)
sop["corte-anterior"]["sat"] = sat_ed  # [edad, NDVI tras manual, NDVI seguida, NDMI tras manual, NDMI seguida]
if SAT_RS is not None:
    fila = lambda d: [int(d.edad), d.indice, int(d.n_grupo), int(d.n_control), r1(d.cambio_grupo, 3), r1(d.cambio_control, 3),  # noqa: E731
                      r1(d.efecto, 3), r1(d.t, 1), r1(d.peor_grupo_pct, 0), r1(d.peor_control_pct, 0)]
    sop["corte-anterior"]["cosecha2026"] = [fila(d) for d in SAT_RS[SAT_RS.comparacion.str.startswith("A ")].itertuples()]
    b = SAT_PS[SAT_PS.grupo == "B"].copy()
    b["edad"] = np.where(b.ndvi_100_2.notna(), 100, 60)
    b["dndvi"] = np.where(b.edad == 100, b.ndvi_100_2 - b.ndvi_100_1, b.ndvi_60_2 - b.ndvi_60_1)
    b["drala"] = np.where(b.edad == 100, b.rala_100_2 - b.rala_100_1, b.rala_60_2 - b.rala_60_1)
    b = b.dropna(subset=["dndvi"]).sort_values("dndvi")
    sop["mecanica"]["alerta2027"] = {"n": int((SAT_PS.grupo == "B").sum()),
        "resumen": [fila(d) for d in SAT_RS[SAT_RS.comparacion.str.startswith("B ")].itertuples()],
        "suertes": [[x.ste, str(x.corte2), int(x.edad), r1(x.dndvi, 2), r1(100 * x.drala, 0)] for x in b.itertuples()]}
sop["corte-anterior"]["ejemplo"] = {
    "titulo": "Hacienda Riopaila, suertes 3501-550, 3501-551, 3501-560 y 3501-561 (32 ha)",
    "suertes": ["3501-550", "3501-551", "3501-560", "3501-561"]}
for cid, col, refd in (("calor", "d32", ref), ("lluvia", "llu03", ref), ("sequia", "defsr_fin", ref_agua),
                       ("riego", "mm_fin", ref_agua), ("freatico", "nf_ciclo", ref_nf)):
    sop[cid] = {"meses": [[m, r1(wm(g, col), 1), r1(wm(gr, col), 1)]
                          for m in range(1, 10) if len(g := q[(q.Mes == m)].dropna(subset=[col])) and len(gr := refd[refd.Mes == m].dropna(subset=[col]))]}
# freático en el tiempo: % del área de Riopaila con nivel a menos de 1 m, por lectura
area_s = ds.drop_duplicates("ste", keep="last").set_index("ste").Area
NF["ha"] = NF.ste.map(area_s)
tl = NF.dropna(subset=["ha"]).groupby("fecha").apply(lambda g: pd.Series({
    "lt1": 100 * g.ha[g.nivel < 1].sum() / g.ha.sum(), "prom": np.average(g.nivel, weights=g.ha)}))
sop["freatico"]["tiempo"] = [[f.date().isoformat(), r1(v.lt1), r1(v.prom, 2)] for f, v in tl.iterrows()]
# lluvia mensual de la EMA La Paila frente a lo normal (2025-2026)
sop["lluvia"]["ema"] = [[f.strftime("%Y-%m"), r1(LLM["La Paila"].get(f), 0), r1(NLL["La Paila"].get(f.month), 0)]
                        for f in pd.date_range("2024-07-01", "2026-09-01", freq="MS")]
sop["calor"]["ema"] = [[f.strftime("%Y-%m"), r1(CAL["La Paila"].get(f), 0), r1(NCA["La Paila"].get(f.month), 1)]
                       for f in pd.date_range("2024-07-01", "2026-09-01", freq="MS")]
s26 = q.groupby("var").Area.sum() / q.Area.sum()
sR = ref.groupby("var").Area.sum() / ref.Area.sum()
sop["variedades"] = {"v": [[v, r1(100 * s26.get(v, 0)), r1(100 * sR.get(v, 0))] for v in (s26 + sR).fillna(0).sort_values(ascending=False).index[:9]]}
EDS = ["<9", "9-10", "10-11", "11-12", "12-13", "13-14", "14-15", ">15"]
e26 = q.groupby("edad_b").Area.sum() / q.Area.sum()
eR = ref.groupby("edad_b").Area.sum() / ref.Area.sum()
sop["edad"] = {"e": [[k, r1(100 * e26.get(k, 0)), r1(100 * eR.get(k, 0))] for k in EDS],
               "prom": [r1(wm(q, "Edad")), r1(wm(ref, "Edad"))]}
okh = ds.mm_ca.notna() & (ds.p_na == 0) & ds.anio.between(2012, 2026)
hh = ds[okh].assign(h=(ds.mm_ca >= 100).astype(float)).groupby("anio").apply(lambda g: 100 * wm(g, "h"))
sop["corte-humedo"] = {"anios": [[int(y), r1(v)] for y, v in hh.items()],
                       "mm": [r1(wm(q[q.mm_ca.notna() & (q.p_na == 0)], "mm_ca"), 0),
                              r1(wm(ref[ref.mm_ca.notna() & (ref.p_na == 0)], "mm_ca"), 0)]}
EC = pd.read_pickle(PK / "efecto_corte.pkl")  # de caida_integrada_rio.py
sop["corte"] = {"efecto": [[r.corte, r1(r.efecto_pct)] for r in EC["efecto"].itertuples()],
                "mezcla": [[k, r1(100 * v["2026"]), r1(100 * v["normal"])] for k, v in EC["mezcla"].iterrows()],
                "prom": [r1(EC["corte_2026"]), r1(EC["corte_normal"])]}
guardar({"causas": CAUSAS, "soporte": sop}, "causas.json")

# ---------------------------------------------------------------- resumen y validación
pp = S[S.ppto.notna() & ok]
wp = pp.area
resumen = {
    "periodo": "enero a septiembre de 2026", "cosechas": int(ok.sum()), "cosechas_total": len(S), "ha": r1(HA, 0),
    "real": r1(TR, 1), "hist": r1(TH, 1), "caida": r1(TH - TR, 2), "caida_pct": r1(100 * (TR / TH - 1), 1),
    "t_caida": r1(float(((S.tch_hist - S.tch_real) * S.area)[ok].sum()), 0), "sin_explicar": sin_exp,
    "puente": {"cosechas": int(len(pp)), "ha": r1(wp.sum(), 0), "ppto": r1(np.average(pp.ppto, weights=wp), 1),
               "hist": r1(np.average(pp.tch_hist, weights=wp), 1), "real": r1(np.average(pp.tch_real, weights=wp), 1)},
}
guardar(resumen, "resumen.json", indent=1)
V = PK / "tablero_v3"
va = pd.read_csv(V / "validacion_anual.csv")
v2 = pd.read_csv(V / "validacion_2026.csv")
vs = pd.read_csv(V / "validacion_segmentos.csv")
vm = pd.read_csv(V / "validacion_metodos.csv")
pred = -float(v2.t_ha.sum() - v2.t_ha.iloc[-1])
validacion = {
    "anual": [[int(a.anio), r1(a.real), r1(a.predicho)] for a in va.itertuples()],
    "corr": r1(np.corrcoef(va.real, va.predicho)[0, 1], 2),
    "v2": [[a.factor, r1(a.t_ha, 2)] for a in v2.itertuples()],
    "segmentos": [[a.dimension, str(a.segmento), int(a.cosechas), r1(a.real), r1(a.predicho)] for a in vs.itertuples()],
    "r2": [[a.metodo, r1(a.r2, 2), a.nota] for a in vm.itertuples()],
}
guardar(validacion, "validacion.json")
print(json.dumps(resumen, ensure_ascii=False))
print("causas:", [(k["nombre"], k["t_ha"]) for k in CAUSAS], "sin explicar", sin_exp)
