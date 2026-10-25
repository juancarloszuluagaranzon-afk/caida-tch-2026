# -*- coding: utf-8 -*-
"""Base de maduración Riopaila: «Calificacion Rot» y «Calificacion Fert» a partir de oportunidad y calidad (BI-BL).

Regla de la base (reproduce el 100 % de lo ya calificado): IDEAL-NORMAL + FRIABLE = 1; IDEAL-NORMAL + HUMEDO = 3;
TARDIO + FRIABLE = 3; TARDIO + HUMEDO = 5 (1 = mejor condición, 5 = peor). Oport. Fert de septiembre no existe (no hay
control de fertilización): donde hay Calidad Fert se califica provisionalmente con IDEAL-NORMAL, que es el 98 % de la base.
Salida: la base con las dos columnas llenas y un Excel en el orden de la base para pegar, con la escala invertida aparte.

Uso: python calificacion_rot_fert.py <base completada.csv> <salida.csv> <para pegar.xlsx>
"""
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
BASE, SAL, XL = sys.argv[1:4]
d = pd.read_csv(BASE, sep=";", encoding="latin-1", dtype=str, keep_default_na=False)
t = {c: d[c].str.strip() for c in ["Oport. Rot", "Calidad Rot.", "Oport. Fert", "Calidad Fert", "Calificacion Rot", "Calificacion Fert"]}
PUNTOS = {("IDEAL-NORMAL", "FRIABLE"): "1", ("IDEAL-NORMAL", "HUMEDO"): "3", ("TARDIO", "FRIABLE"): "3", ("TARDIO", "HUMEDO"): "5"}


def calificar(op, cal):
    return np.array([PUNTOS.get((o, c), "") for o, c in zip(op, cal)])


rot = calificar(t["Oport. Rot"], t["Calidad Rot."])
op_fert = np.where((t["Oport. Fert"] == "") & (t["Calidad Fert"] != ""), "IDEAL-NORMAL", t["Oport. Fert"])
fert = calificar(op_fert, t["Calidad Fert"])
prov = (t["Oport. Fert"] == "") & (fert != "")
for col, v in (("Calificacion Rot", rot), ("Calificacion Fert", fert)):
    ya = t[col] != ""
    print(f"{col}: la regla reproduce {100 * (v[ya] == t[col][ya]).mean():.0f} % de {ya.sum()} filas ya calificadas")
    m = (t[col] == "") & (v != "")
    d.loc[m, col] = v[m]
    print(f"   llenadas {m.sum()} {pd.Series(v[m]).value_counts().sort_index().to_dict()} · siguen vacías {(d[col].str.strip() == '').sum()}")
print(f"Calificacion Fert provisional (Oport. Fert supuesta IDEAL-NORMAL): {prov.sum()} filas")
d.to_csv(SAL, sep=";", encoding="latin-1", index=False)

# Excel para pegar, en el orden de la base
nuevo_rot = (t["Calificacion Rot"] == "") & (d["Calificacion Rot"].str.strip() != "")
nuevo_fert = (t["Calificacion Fert"] == "") & (d["Calificacion Fert"].str.strip() != "")
inv = lambda s: s.str.strip().map({"1": 5, "3": 3, "5": 1})  # noqa: E731
p = pd.DataFrame({
    "Suerte": d.suerte, "Mes": pd.to_numeric(d.MES), "Corte": pd.to_numeric(d.Corte),
    "Oport. Rot": t["Oport. Rot"], "Calidad Rot.": t["Calidad Rot."], "Calificacion Rot": pd.to_numeric(d["Calificacion Rot"]),
    "Oport. Fert": t["Oport. Fert"], "Calidad Fert": t["Calidad Fert"], "Calificacion Fert": pd.to_numeric(d["Calificacion Fert"]),
    "Rot (1 malo, 5 bueno)": inv(d["Calificacion Rot"]), "Fert (1 malo, 5 bueno)": inv(d["Calificacion Fert"]),
    "Nota": np.select([prov & nuevo_fert, d["Calificacion Rot"].str.strip() == "", d["Calificacion Fert"].str.strip() == ""],
                      ["Fert provisional: sin Oport. Fert, se supuso IDEAL-NORMAL", "Rot sin calificar: falta oportunidad o calidad",
                       "Fert sin calificar: sin fertilización al suelo en SimpleAgri"], "")})
p.loc[(p.Nota == "") & (p["Calificacion Rot"].isna()), "Nota"] = "Rot sin calificar: falta oportunidad o calidad"
with pd.ExcelWriter(XL, engine="openpyxl") as w:
    p.to_excel(w, sheet_name="Calificaciones", index=False)
    pd.DataFrame({"Escala de la base (1 = mejor, 5 = peor)": [
        "IDEAL-NORMAL + FRIABLE = 1", "IDEAL-NORMAL + HUMEDO = 3", "TARDIO + FRIABLE = 3", "TARDIO + HUMEDO = 5",
        "Columnas «1 malo, 5 bueno»: la misma calificación invertida (1 ↔ 5).",
        "Verde: calificación nueva (septiembre). Amarillo: Fert provisional."]}).to_excel(w, sheet_name="Regla", index=False)
wb = load_workbook(XL)
ws = wb["Calificaciones"]
verde, amarillo = PatternFill("solid", fgColor="D9EAD3"), PatternFill("solid", fgColor="FFF2CC")
cols = {c.value: c.column for c in ws[1]}
for i in range(len(p)):
    if nuevo_rot.iloc[i]:
        for c in ("Calificacion Rot", "Rot (1 malo, 5 bueno)"):
            ws.cell(i + 2, cols[c]).fill = verde
    if nuevo_fert.iloc[i]:
        for c in ("Calificacion Fert", "Fert (1 malo, 5 bueno)"):
            ws.cell(i + 2, cols[c]).fill = amarillo if prov.iloc[i] else verde
for c in ws[1]:
    c.font = Font(bold=True)
    ws.column_dimensions[c.column_letter].width = max(10, min(55, len(str(c.value)) + 3))
ws.column_dimensions[ws.cell(1, cols["Nota"]).column_letter].width = 55
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
wb["Regla"].column_dimensions["A"].width = 75
wb.save(XL)
print(f"guardado: {SAL}\npara pegar: {XL}")
