# -*- coding: utf-8 -*-
"""Excel de soporte de la diapositiva «Impacto de la restricción presupuestal en el TCH» de Riopaila (2026 y 2027).

Lee las tablas que dejaron los programas de esta carpeta (en la carpeta de trabajo caida_oct) y arma un libro con una
hoja por factor (suerte a suerte, con fórmulas), la hoja de efectos y supuestos, y el resumen que alimenta la diapositiva.

Uso (en la carpeta caida_oct): python impacto_excel.py <salida.xlsx>
"""
import json
import sys
import warnings
from datetime import date

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
import presupuesto_cosecha as pc  # noqa: E402
from tchmotor import rutas  # noqa: E402

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
SAL = sys.argv[1]
HOY = date.today()
wm = lambda g, k, w="Area": np.average(g[k].dropna(), weights=g[w][g[k].notna()])  # noqa: E731

# ================================================================ datos y efectos exactos
c = pd.read_pickle("rio_factores_sep.pkl")
est = {e["sec_ste"]: e["tch_estimado"] for e in json.load(open(rutas.TABLERO_DATA / "estimados.json", encoding="utf-8")) if e["planta"] == "riopaila"}

# curva por corte en el mismo lote (como no_renovadas_2026.py)
TOPE = 13
cc = c.copy()
cc["k"] = cc.Corte.clip(upper=TOPE).astype(int)
cc["e2"] = cc.Edad ** 2
mo = smf.wls("y ~ C(k, Treatment(1)) + C(anio) + C(Mes) + Edad + e2 + C(ste)", cc, weights=cc.Area).fit()
b = pd.Series({k: (0.0 if k == 1 else mo.params[f"C(k, Treatment(1))[T.{k}]"]) for k in range(1, TOPE + 1)})
ren = float(b[1:6].mean())
curva = pd.DataFrame({"corte": b.index, "efecto": np.exp(b.values) - 1, "factor": np.exp(ren - b.values)})

# renovación
cs = pd.read_pickle("no_renovadas_cosechas_2026.pkl")
r26 = cs[cs.no_renov].copy()
ma, _ = pc.leer_maestro("riopaila")
M = pd.DataFrame([dict(ste=s, ncorte=v["ncorte"], var=v["var"], area=v["area"], prox=v["prox"]) for s, v in ma.items()])
import openpyxl  # noqa: E402
F = [list(r) for r in openpyxl.load_workbook(r"C:\Users\Agr349\AppData\Local\Temp\Areas no renovadas (1).xlsx", read_only=True,
                                              data_only=True).worksheets[0].iter_rows(values_only=True)]
txt = lambda v: str(v).strip() if v is not None else ""  # noqa: E731
i0 = next(i for i, r in enumerate(F) if txt(r[0]) == "Zona" and txt(r[2]) == "Sect-Sue")
plan = pd.DataFrame([dict(ste=txt(r[2]), area=float(r[3]), renovada=txt(r[4]).upper() == "SI") for r in F[i0 + 1:] if txt(r[2])])
proc = {txt(r[13]) for r in F if len(r) > 19 and txt(r[13]) and "-" in txt(r[13])}
NO = set(plan.ste[~plan.renovada]) - proc
mn = plan[plan.ste.isin(NO)].merge(M.rename(columns={"area": "area_m"}), on="ste", how="left")
mn["area_m"] = mn.area_m.fillna(mn.area)
pend_ren = mn[mn.ste.isin(est) & ~mn.ste.isin(r26.ste) & (mn.ncorte > 1) & (mn["var"].str.upper() != "RENOVACION")].copy()
pend_ren["anio"] = pd.to_datetime(pend_ren.prox, errors="coerce").dt.year
pend_ren = pend_ren[pend_ren.anio == 2026]
pend_ren["tch"] = pend_ren.ste.map(est)
siguen = mn[(mn.ncorte > 1) & (mn["var"].str.upper() != "RENOVACION")]

# roturación, malezas, fertilización
rt = pd.read_pickle("roturacion.pkl")
j, pendr = rt["j"], rt["pend"]
rot26 = j[j.c_anio == 2026][["ste", "corte_f", "rot1", "Area", "c_tch_calc", "c_Mes"]].rename(columns={"Area": "ha", "c_tch_calc": "tch", "c_Mes": "mes"})
rot26["fuente"] = "control + cosecha"
p26 = pendr[pendr.corte_f < "2025-11-01"][["ste", "corte_f", "rot1", "area"]].rename(columns={"area": "ha"})
p26["tch"] = p26.ste.map(est).fillna(105.0)
p26["mes"] = None
p26["fuente"] = "control, por cosechar nov-dic (TCH motor)"
rot26 = pd.concat([rot26, p26], ignore_index=True)
co = pd.read_pickle("cana2027_labores.pkl")
cy = pd.read_pickle("ciclos_labores.pkl")
lag = pd.read_pickle("labores_perdidas.pkl")["lag"]
co["tch"] = co.ste.map(est).fillna(105.0)
p27 = pendr[pendr.corte_f >= "2025-11-01"]
co["rot_oficial"] = co.ste.map(dict(zip(p27.ste, p27.rot1)))
dc = co.d_rot - lag
co["rot_sa"] = np.where(co.d_rot.notna(), np.where(dc > 30, "TARDIO", np.where(dc > 15, "NORMAL", "IDEAL")),
                        np.where(co.dias_hoy > 30 + lag, "PENDIENTE (TARDÍO)", "en plazo"))
co["rot_f"] = co.rot_oficial.fillna(co.rot_sa)
co["rot_fuente"] = np.where(co.rot_oficial.notna(), "control oficial", "SimpleAgri (−%d días)" % lag)
mal26 = cy[cy.h_anio == 2026].copy()
base = pd.read_pickle("base_rio_labores.pkl")

# riego y lluvia 2026 (modelo C de riego_2026.py) y arranque (factores_sep.py)
rr = pd.read_pickle("rio_riego.pkl")
d = rr.dropna(subset=["llu03", "d32", "llu_e0_4", "llu_e4_8", "llu_fin"]).copy()
for k in ("e0_4", "e4_8", "fin"):
    d[f"def100_{k}"] = d[f"def_{k}"] / 100
ED = "C(edad_b, Treatment('12-13'))"
mec = " + ".join(f"mec_{y}" for y in range(2022, 2027))
moC = smf.wls(f"y ~ C(ste) + C(anio) + C(Mes) + {ED} + C(soca_b) + C(var) + inc + Corte + {mec} + llu03_100 + d32 + def100_e0_4 + def100_e4_8 + def100_fin",
              d, weights=d.Area).fit()
q, r = d[(d.anio == 2026) & (d.Mes <= 9)], d[d.anio.between(2022, 2025) & (d.Mes <= 9)]
riego26 = [(k, moC.params[f"def100_{k}"], wm(q, f"def_{k}"), wm(r, f"def_{k}")) for k in ("e0_4", "e4_8", "fin")]
prod_es = c[(c.anio == 2026) & (c.Mes <= 9)]
ha_es, ton_es = float(prod_es.Area.sum()), float(prod_es["TC-Total"].sum())
tab = pd.read_pickle("tabla_riopaila_sep.pkl")
llu26, lluref = float(tab.loc["ene-sep", "lluvia 0-3 m 2026"]), float(tab.loc["ene-sep", "lluvia 0-3 m ref"])
c_arr = float(tab.loc["ene-sep", "arranque (lluvia 0-3 m)"] / 100) / ((llu26 - lluref) / 100) if llu26 != lluref else -0.012
co27 = pd.read_pickle("cohorte_2027_riego.pkl")
r27 = pd.read_pickle("riego_2027.pkl")
sl = pd.read_pickle("siembra_lucro.pkl")
s26, ren27, pl, fut = sl["s26"], sl["ren"], sl["pl"], sl["fut"]
TCHM = sl["tchm"]
HA26_COS = float(cs.Area.sum())
pend_all = [(s, v["area"]) for s, v in ma.items() if v["prox"] and v["prox"].year == 2026 and s not in set(cs.ste) and v["var"].upper() != "RENOVACION"]
HA26_PEND = float(sum(a or 0 for _, a in pend_all))

# ================================================================ libro
FN_ = "Arial"
FB, FNB, FH = Font(name=FN_, size=10), Font(name=FN_, size=10, bold=True), Font(name=FN_, size=10, bold=True, color="FFFFFF")
FIN, FNOTE, FTIT = Font(name=FN_, size=10, color="0000FF"), Font(name=FN_, size=9, italic=True, color="595959"), Font(name=FN_, size=14, bold=True)
FLINK = Font(name=FN_, size=10, color="008000")
HEAD, TOT, YEL = PatternFill("solid", fgColor="1E3328"), PatternFill("solid", fgColor="E3E8DB"), PatternFill("solid", fgColor="FFFF00")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
WR = Alignment(wrap_text=True, vertical="top")
wb = Workbook()
L = get_column_letter


def hoja(nombre, titulo, nota):
    ws = wb.create_sheet(nombre)
    ws["A1"], ws["A1"].font = titulo, FTIT
    ws["A2"], ws["A2"].font = nota, FNOTE
    return ws


def encabezado(ws, fila, cols):
    for i, (n, w) in enumerate(cols, 1):
        x = ws.cell(fila, i, n)
        x.font, x.alignment, x.fill = FH, CEN, HEAD
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[fila].height = 40


def celda(ws, f, col, v, fmt=None, inp=False):
    x = ws.cell(f, col, None if (isinstance(v, float) and np.isnan(v)) else v)
    x.font = FIN if inp else FB
    if fmt:
        x.number_format = fmt
    return x


def total(ws, f, etiqueta, formulas):
    ws.cell(f, 1, etiqueta).font = FNB
    for col, (form, fmt) in formulas.items():
        x = ws.cell(f, col, form)
        x.font, x.number_format = FNB, fmt
    for col in range(1, max(formulas) + 1):
        ws.cell(f, col).fill = TOT


# ------------------------------------------------ Efectos y supuestos
ef = hoja("Efectos y supuestos", "Efectos medidos y supuestos (celdas azules = dato; se pueden cambiar y el libro se recalcula)",
          "Efectos medidos con cosechas de Riopaila, cada suerte frente a su propio promedio 2020-25, con efectos de año, mes, edad y corte.")
filas_ef = [
    ("HA26C", "Ha cosechadas 2026 (ene-oct)", HA26_COS, "#,##0", "ZQM / G&P a septiembre + app de Cosecha octubre"),
    ("HA26P", "Ha por cosechar 2026 (nov-dic)", HA26_PEND, "#,##0", "Maestro: próximo corte en 2026 sin cosecha registrada"),
    ("HA26", "Ha a cosechar 2026", "=B4+B5", "#,##0", "Denominador del impacto en TCH 2026"),
    ("HA27", "Ha a cosechar 2027", 5393, "#,##0", "Presupuesto de cosecha 2027 (proyección)"),
    ("ROT", "Roturación 1.ª labor TARDÍA (>30 días)", -0.063, "0.0%", "Control oficial de roturación; cosechas 2024-26, n=803; p<0,01; controla suelo húmedo (sin control −10,9 %)"),
    ("MAL30", "Malezas: por cada 30 días de atraso del 1.er control químico", -0.0146, "0.00%", "SimpleAgri; socas cosechadas 2024-26, n=1.011; p=0,007. Con la fecha real de aplicaciones (cosechas 2026, n=259): −2,3 % (p=0,11); se deja el menor"),
    ("MALP", "Malezas: plazo del protocolo (días)", 90, "0", "Pedido del usuario (22-oct): el primer control hasta 90 días cuenta como oportuno; se cuenta solo el atraso más allá"),
    ("MAL0", "Malezas: ciclo sin control químico", -0.067 + 0.0146, "0.0%", "−6,7 % frente a ≤60 días (p=0,001), menos lo que ya restaría un control a los 90 días (−1,46 %)"),
    ("DEF04", "Riego: déficit sin cubrir meses 0-4 (por 100 mm)", float(moC.params["def100_e0_4"]), "0.00%", "Modelo de riego 2022-26 (arranque seco ayuda)"),
    ("DEF48", "Riego: déficit sin cubrir meses 4-8 (por 100 mm)", float(moC.params["def100_e4_8"]), "0.00%", "Modelo de riego 2022-26"),
    ("DEFF", "Riego: déficit sin cubrir del mes 8 a la cosecha (por 100 mm)", float(moC.params["def100_fin"]), "0.00%", "Modelo de riego 2022-26; p<0,001"),
    ("ARR", "Lluvia del arranque, meses 0-3 (por 100 mm)", c_arr, "0.00%", "Modelo de factores 2011-26 (más lluvia al arrancar = menos TCH)"),
    ("TCHM", "TCHM Riopaila 2024-25 (t/ha por mes)", float(TCHM), "0.00", "ZQM: toneladas ÷ (área × edad)"),
    ("LUCP", "Lucro: días normales entre corte y siembra", 60, "0", "Supuesto: lo que pasa de 60 días es pérdida"),
    ("NINO", "Escenario El Niño: déficit final sin cubrir extra (mm)", 117, "0", "Lo que pasó en 2026 (149 contra 32 mm normales)"),
    ("TCH27", "TCH esperado 2027 para el escenario de riego", 103.6, "0.0", "Presupuesto 2027, escenario «como 2026»"),
    ("CSI", "Corte a siembra (por 100 días)", 0, "0.0%", "Medido −0,8 % (p=0,22): no significativo → 0"),
    ("FER", "Fertilización tardía", 0, "0.0%", "Base de maduración 2026: sin efecto (TARDÍO +8 %, n=3; suelo húmedo +0,9 %, p=0,72)"),
]
encabezado(ef, 3, [("Concepto", 58), ("Valor", 12), ("Fuente / cómo se midió", 100)])
REF = {}
for i, (k, t, v, fmt, nota) in enumerate(filas_ef, 4):
    ef.cell(i, 1, t).font = FB
    x = celda(ef, i, 2, v, fmt, inp=not (isinstance(v, str) and v.startswith("=")))
    if not (isinstance(v, str) and v.startswith("=")):
        x.fill = YEL
    ef.cell(i, 3, nota).font = FNOTE
    REF[k] = f"'Efectos y supuestos'!$B${i}"
r0 = 4 + len(filas_ef) + 1
ef.cell(r0, 1, "Curva por número de corte en el mismo lote (para «Menor renovación»)").font = FNB
encabezado(ef, r0 + 1, [("Corte", 58), ("Frente a la plantilla", 12), ("Factor para llevar al lote renovado (promedio de cortes 1-5)", 100)])
for i, rw in enumerate(curva.itertuples(), r0 + 2):
    celda(ef, i, 1, int(rw.corte) if rw.corte < TOPE else TOPE, "0", inp=True)
    celda(ef, i, 2, float(rw.efecto), "+0.0%;-0.0%", inp=True)
    celda(ef, i, 3, float(rw.factor), "0.000", inp=True)
CURV_K = f"'Efectos y supuestos'!$A${r0 + 2}:$A${r0 + 1 + len(curva)}"
CURV_F = f"'Efectos y supuestos'!$C${r0 + 2}:$C${r0 + 1 + len(curva)}"
ef.cell(r0 + 2 + len(curva), 1, "El corte 13 agrupa 13 o más.").font = FNOTE


def factor_corte(celda_corte):
    return f"INDEX({CURV_F},MATCH(MIN({celda_corte},13),{CURV_K},0))"


# ------------------------------------------------ Renovación
rv = hoja("Renovación", "Menor área de renovación: suertes del plan 2025 que no se renovaron",
          f"85 suertes no renovadas (93 del plan menos 8 que se renovaron tarde). TCH renovado = TCH × factor de la curva por corte. Pérdida = (TCH renovado − TCH) × ha.")
cols = [("Suerte", 10), ("Origen", 30), ("Corte", 7), ("Área (ha)", 9), ("TCH (real o motor)", 10), ("Factor curva", 9), ("TCH renovado", 10), ("Toneladas perdidas", 11)]
encabezado(rv, 4, cols)
f = 5
f_ini = f
for rw in r26.sort_values("ste").itertuples():
    celda(rv, f, 1, rw.ste, inp=True); celda(rv, f, 2, "Cosechada 2026 (ene-oct)", inp=True); celda(rv, f, 3, int(rw.Corte), "0", True)
    celda(rv, f, 4, float(rw.Area), "0.00", True); celda(rv, f, 5, float(rw.tch), "0.0", True)
    celda(rv, f, 6, f"={factor_corte(f'C{f}')}", "0.000"); celda(rv, f, 7, f"=E{f}*F{f}", "0.0"); celda(rv, f, 8, f"=(G{f}-E{f})*D{f}", "#,##0")
    f += 1
f_cos_fin = f - 1
for rw in pend_ren.sort_values("ste").itertuples():
    celda(rv, f, 1, rw.ste, inp=True); celda(rv, f, 2, "Por cosechar 2026 (TCH del motor)", inp=True); celda(rv, f, 3, int(rw.ncorte), "0", True)
    celda(rv, f, 4, float(rw.area_m), "0.00", True); celda(rv, f, 5, float(rw.tch), "0.0", True)
    celda(rv, f, 6, f"={factor_corte(f'C{f}')}", "0.000"); celda(rv, f, 7, f"=E{f}*F{f}", "0.0"); celda(rv, f, 8, f"=(G{f}-E{f})*D{f}", "#,##0")
    f += 1
f_fin = f - 1
total(rv, f, "Total 2026", {4: (f"=SUM(D{f_ini}:D{f_fin})", "#,##0"), 8: (f"=SUM(H{f_ini}:H{f_fin})", "#,##0")})
RV26 = (f"Renovación!$D${f}", f"Renovación!$H${f}")
f += 2
rv.cell(f, 1, "2027").font = FNB
rv.cell(f + 1, 1, "Ha de las no renovadas que siguen como soca (maestro)").font = FB
celda(rv, f + 1, 4, float(siguen.area_m.sum()), "#,##0", True)
rv.cell(f + 2, 1, "Pérdida media por ha de las cosechadas en 2026 (t/ha)").font = FB
celda(rv, f + 2, 4, f"=SUM(H{f_ini}:H{f_cos_fin})/SUM(D{f_ini}:D{f_cos_fin})", "0.0")
rv.cell(f + 3, 1, "Toneladas 2027 (si no se renuevan)").font = FNB
celda(rv, f + 3, 4, f"=D{f + 1}*D{f + 2}", "#,##0").font = FNB
RV27 = (f"Renovación!$D${f + 1}", f"Renovación!$D${f + 3}")
rv.freeze_panes = "B5"

# ------------------------------------------------ Roturación
ro = hoja("Roturación", "Oportunidad de la roturación (primera labor)",
          "Pérdida = −efecto × TCH × ha cuando la 1.ª labor fue TARDÍA (>30 días) o sigue pendiente fuera de plazo. NORMAL (16-30 días) no resta.")
cols = [("Suerte", 10), ("Fecha de corte", 11), ("Fuente", 30), ("Oportunidad 1.ª labor", 18), ("Área (ha)", 9), ("TCH", 8), ("Efecto", 8), ("Toneladas perdidas", 11)]
encabezado(ro, 4, cols)
f = 5
ro.cell(f, 1, "2026: cosechas de 2026 con su roturación (control oficial)").font = FNB
f += 1
a26 = f
for rw in rot26.sort_values(["rot1", "ste"]).itertuples():
    celda(ro, f, 1, rw.ste, inp=True); celda(ro, f, 2, rw.corte_f.date() if pd.notna(rw.corte_f) else None, "dd/mm/yyyy", True)
    celda(ro, f, 3, rw.fuente, inp=True); celda(ro, f, 4, rw.rot1, inp=True); celda(ro, f, 5, float(rw.ha), "0.00", True)
    celda(ro, f, 6, float(rw.tch), "0.0", True)
    celda(ro, f, 7, f'=IF(OR(D{f}="TARDIO",D{f}="PENDIENTE (TARDÍO)"),{REF["ROT"]},0)', "0.0%")
    celda(ro, f, 8, f"=-G{f}*F{f}*E{f}", "#,##0")
    f += 1
b26 = f - 1
total(ro, f, "Total 2026 (tardías)", {5: (f'=SUMIFS(E{a26}:E{b26},G{a26}:G{b26},"<0")', "#,##0"), 8: (f"=SUM(H{a26}:H{b26})", "#,##0")})
RO26 = (f"Roturación!$E${f}", f"Roturación!$H${f}")
f += 2
ro.cell(f, 1, "2027: socas cortadas desde nov-2025 (estado a hoy)").font = FNB
f += 1
a27 = f
for rw in co.sort_values(["rot_f", "ste"]).itertuples():
    celda(ro, f, 1, rw.ste, inp=True); celda(ro, f, 2, rw.corte_f.date(), "dd/mm/yyyy", True); celda(ro, f, 3, rw.rot_fuente, inp=True)
    celda(ro, f, 4, rw.rot_f, inp=True); celda(ro, f, 5, float(rw.area), "0.00", True); celda(ro, f, 6, float(rw.tch), "0.0", True)
    celda(ro, f, 7, f'=IF(OR(D{f}="TARDIO",D{f}="PENDIENTE (TARDÍO)"),{REF["ROT"]},0)', "0.0%")
    celda(ro, f, 8, f"=-G{f}*F{f}*E{f}", "#,##0")
    f += 1
b27 = f - 1
total(ro, f, "Total 2027 (tardías o pendientes fuera de plazo)", {5: (f'=SUMIFS(E{a27}:E{b27},G{a27}:G{b27},"<0")', "#,##0"), 8: (f"=SUM(H{a27}:H{b27})", "#,##0")})
RO27 = (f"Roturación!$E${f}", f"Roturación!$H${f}")
ro.cell(f + 1, 1, "TCH 2027: el del motor (si no hay, 105). «PENDIENTE (TARDÍO)» en SimpleAgri puede ser un registro faltante.").font = FNOTE
ro.freeze_panes = "B5"

# ------------------------------------------------ Malezas
MA = pd.read_pickle("malezas_aplicaciones.pkl")
ml = hoja("Malezas", "Oportunidad del control químico de malezas (primer control del ciclo)",
          f"Fecha real de la base de aplicaciones de los contratistas (2025 y ago-oct 2026); donde no hay, SimpleAgri corrida {MA['lag']:.0f} días "
          "(desfase medido). Pérdida % = efecto por 30 días × (días − plazo)/30; sin control químico = efecto fijo. Toneladas = −% × TCH × ha.")
cols = [("Suerte", 10), ("Fecha de corte", 11), ("Fuente de la fecha", 20), ("Días al 1.er control (o días ya pasados)", 13), ("Área (ha)", 9), ("TCH", 8), ("Pérdida %", 9), ("Toneladas perdidas", 11)]
encabezado(ml, 4, cols)
f = 5
ml.cell(f, 1, "2026: socas cosechadas en 2026 (ene-oct)").font = FNB
f += 1
a26 = f
m26 = pd.concat([MA["cyc"].assign(fuente=MA["cyc"].fuente), MA["rest"].assign(fuente=np.where(MA["rest"].d_uso.notna(), "SimpleAgri corregido", "sin registro"))])
for rw in m26.sort_values("ste").itertuples():
    hecho = pd.notna(rw.d_uso)
    celda(ml, f, 1, rw.ste, inp=True); celda(ml, f, 2, rw.corte_f.date(), "dd/mm/yyyy", True)
    celda(ml, f, 3, rw.fuente if hecho else "sin control", inp=True); celda(ml, f, 4, float(rw.d_uso) if hecho else None, "0", True)
    celda(ml, f, 5, float(rw.h_area), "0.00", True); celda(ml, f, 6, float(rw.h_tch), "0.0", True)
    celda(ml, f, 7, f'=IF(C{f}="sin control",{REF["MAL0"]},{REF["MAL30"]}*MAX(0,D{f}-{REF["MALP"]})/30)', "0.0%")
    celda(ml, f, 8, f"=-G{f}*F{f}*E{f}", "#,##0")
    f += 1
b26 = f - 1
total(ml, f, "Total 2026", {5: (f'=SUMIFS(E{a26}:E{b26},G{a26}:G{b26},"<0")', "#,##0"), 8: (f"=SUM(H{a26}:H{b26})", "#,##0")})
ML26 = (f"Malezas!$E${f}", f"Malezas!$H${f}")
f += 2
ml.cell(f, 1, "2027: socas cortadas desde nov-2025 (estado a hoy; lo pendiente cuenta con los días que ya pasaron)").font = FNB
f += 1
a27 = f
for rw in MA["co"].sort_values("ste").itertuples():
    celda(ml, f, 1, rw.ste, inp=True); celda(ml, f, 2, rw.corte_f.date(), "dd/mm/yyyy", True)
    celda(ml, f, 3, rw.fuente, inp=True); celda(ml, f, 4, float(rw.dias), "0", True)
    celda(ml, f, 5, float(rw.area), "0.00", True); celda(ml, f, 6, float(rw.tch), "0.0", True)
    celda(ml, f, 7, f'={REF["MAL30"]}*MAX(0,D{f}-{REF["MALP"]})/30', "0.0%")
    celda(ml, f, 8, f"=-G{f}*F{f}*E{f}", "#,##0")
    f += 1
b27 = f - 1
total(ml, f, "Total 2027 (mínimo)", {5: (f'=SUMIFS(E{a27}:E{b27},G{a27}:G{b27},"<0")', "#,##0"), 8: (f"=SUM(H{a27}:H{b27})", "#,##0")})
ML27 = (f"Malezas!$E${f}", f"Malezas!$H${f}")
ml.freeze_panes = "B5"

# ------------------------------------------------ Riego y lluvia
rl = hoja("Riego y lluvia", "Falta de riego y precipitaciones",
          "Falta de riego = déficit de agua (evaporación − lluvia − riego) que quedó sin cubrir, frente a lo normal 2022-25. Precipitaciones = lluvia de los primeros 3 meses del ciclo.")
encabezado(rl, 4, [("2026 (cosechas ene-sep)", 44), ("2026 (mm)", 11), ("Normal 2022-25 (mm)", 12), ("Diferencia (mm)", 11), ("Efecto por 100 mm", 11), ("Aporte al TCH", 11)])
for i, (k, coef, v26, vref) in enumerate(riego26, 5):
    nom = {"e0_4": "Déficit sin cubrir meses 0-4", "e4_8": "Déficit sin cubrir meses 4-8", "fin": "Déficit sin cubrir mes 8 a cosecha"}[k]
    rl.cell(i, 1, nom).font = FB
    celda(rl, i, 2, float(v26), "0", True); celda(rl, i, 3, float(vref), "0", True); celda(rl, i, 4, f"=B{i}-C{i}", "+0;-0")
    celda(rl, i, 5, f"={REF[{'e0_4': 'DEF04', 'e4_8': 'DEF48', 'fin': 'DEFF'}[k]]}", "0.00%"); celda(rl, i, 6, f"=E{i}*D{i}/100", "+0.0%;-0.0%")
rl.cell(8, 1, "Producción ene-sep 2026 (t) y ha").font = FB
celda(rl, 8, 2, ton_es, "#,##0", True); celda(rl, 8, 3, ha_es, "#,##0", True)
q_all = rr[(rr.anio == 2026) & (rr.Mes <= 9)].dropna(subset=["def_fin"])
r_all = rr[rr.anio.between(2022, 2025) & (rr.Mes <= 9)].dropna(subset=["def_fin"])
area_def = float(q_all.Area[q_all.def_fin > wm(r_all, "def_fin")].sum())
total(rl, 9, "Falta de riego 2026: toneladas (columna C = ha con déficit final mayor a lo normal)", {6: ("=-SUM(F5:F7)*B8", "#,##0"), 3: (area_def, "#,##0")})
RG26 = ("'Riego y lluvia'!$C$9", "'Riego y lluvia'!$F$9")
rl.cell(11, 1, "Lluvia del arranque (meses 0-3), cosechas ene-sep").font = FB
celda(rl, 11, 2, llu26, "0", True); celda(rl, 11, 3, lluref, "0", True); celda(rl, 11, 4, "=B11-C11", "+0;-0")
celda(rl, 11, 5, f"={REF['ARR']}", "0.00%"); celda(rl, 11, 6, "=E11*D11/100", "+0.0%;-0.0%")
c26a = c[(c.anio == 2026) & (c.Mes <= 9)].dropna(subset=["llu03"])
area_llu = float(c26a.Area[c26a.llu03 > lluref].sum())
total(rl, 12, "Precipitaciones 2026: toneladas (columna C = ha con lluvia de arranque mayor a lo normal)", {6: ("=-F11*B8", "#,##0"), 3: (area_llu, "#,##0")})
PR26 = ("'Riego y lluvia'!$C$12", "'Riego y lluvia'!$F$12")
encabezado(rl, 14, [("2027 (lo que la caña ya vivió)", 44), ("Cohorte (mm)", 11), ("Normal (mm)", 12), ("Diferencia (mm)", 11), ("Efecto por 100 mm", 11), ("Ha de la etapa", 11)])
co27["area"] = co27.area.astype(float)
filas27 = []
for k, nom, ref in (("def_e0_4", "Déficit sin cubrir meses 0-4", "DEF04"), ("def_e4_8", "Déficit sin cubrir meses 4-8", "DEF48")):
    x = co27.dropna(subset=[k])
    filas27.append((nom, wm(x, k, "area"), wm(x, "n_" + k, "area"), ref, float(x.area.sum())))
x = co27.dropna(subset=["llu03"])
filas27.append(("Lluvia del arranque (meses 0-3)", wm(x, "llu03", "area"), wm(x, "n_llu03", "area"), "ARR", float(x.area.sum())))
for i, (nom, v, n, ref, ha) in enumerate(filas27, 15):
    rl.cell(i, 1, nom).font = FB
    celda(rl, i, 2, float(v), "0", True); celda(rl, i, 3, float(n), "0", True); celda(rl, i, 4, f"=B{i}-C{i}", "+0;-0")
    celda(rl, i, 5, f"={REF[ref]}", "0.00%"); celda(rl, i, 6, ha, "#,##0", True)
rl.cell(19, 1, "Toneladas (− = a favor) = −efecto × diferencia/100 × ha × TCH esperado 2027").font = FNOTE
total(rl, 20, "Riego medido hasta hoy 2027", {6: (f"=-(E15*D15/100*F15+E16*D16/100*F16)*{REF['TCH27']}", "#,##0")})
total(rl, 21, "Escenario El Niño 2027 (final del ciclo como 2026)", {3: (f"={REF['HA27']}", "#,##0"), 6: (f"=-{REF['DEFF']}*{REF['NINO']}/100*{REF['HA27']}*{REF['TCH27']}", "#,##0")})
x27 = co27.dropna(subset=["llu03"])
area_seco = float(x27.area[x27.llu03 < x27.n_llu03].sum())
total(rl, 22, "Precipitaciones 2027 (arranque; columna C = ha con arranque más seco que lo normal, a favor)",
      {3: (area_seco, "#,##0"), 6: (f"=-E17*D17/100*F17*{REF['TCH27']}", "#,##0")})
RG27 = ("'Riego y lluvia'!$C$21", "'Riego y lluvia'!$F$21")
PR27 = ("'Riego y lluvia'!$C$22", "'Riego y lluvia'!$F$22")
rl.cell(23, 1, "En la diapositiva, la falta de riego 2027 es el escenario El Niño; lo medido hasta hoy (fila 20) casi no resta.").font = FNOTE

# ------------------------------------------------ Lucro cesante
lu = hoja("Lucro cesante", "Lucro cesante: días que el lote quedó cortado y sin sembrar",
          "Toneladas = ha × (días − días normales) / 30,44 × TCHM. Con todos los días, columna aparte.")
cols = [("Suerte", 10), ("Área (ha)", 9), ("Corte anterior", 11), ("Siembra (o hoy)", 11), ("Días sin caña", 9), ("Días por encima de lo normal", 11),
        ("Toneladas perdidas", 11), ("Toneladas con todos los días", 12)]
encabezado(lu, 4, cols)
f = 5
lu.cell(f, 1, "2026: renovaciones sembradas en 2026").font = FNB
f += 1
a = f
for rw in s26.sort_values("dias_sin_cana", ascending=False).itertuples():
    celda(lu, f, 1, rw.ste, inp=True); celda(lu, f, 2, float(rw.area), "0.00", True)
    celda(lu, f, 3, pd.Timestamp(rw.corte_prev).date(), "dd/mm/yyyy", True); celda(lu, f, 4, pd.Timestamp(rw.siembra).date(), "dd/mm/yyyy", True)
    celda(lu, f, 5, f"=D{f}-C{f}", "0"); celda(lu, f, 6, f"=MAX(0,E{f}-{REF['LUCP']})", "0")
    celda(lu, f, 7, f"=B{f}*F{f}/30.44*{REF['TCHM']}", "#,##0"); celda(lu, f, 8, f"=B{f}*E{f}/30.44*{REF['TCHM']}", "#,##0")
    f += 1
bb = f - 1
total(lu, f, "Total 2026", {2: (f"=SUM(B{a}:B{bb})", "#,##0.0"), 7: (f"=SUM(G{a}:G{bb})", "#,##0"), 8: (f"=SUM(H{a}:H{bb})", "#,##0")})
LU26 = (f"'Lucro cesante'!$B${f}", f"'Lucro cesante'!$G${f}")
f += 2
lu.cell(f, 1, f"2027: cortadas y todavía sin sembrar (días hasta {HOY.strftime('%d/%m/%Y')})").font = FNB
f += 1
a = f
for rw in ren27.itertuples():
    celda(lu, f, 1, rw.ste, inp=True); celda(lu, f, 2, float(rw.area), "0.00", True)
    celda(lu, f, 3, pd.Timestamp(rw.corte_prev).date(), "dd/mm/yyyy", True); celda(lu, f, 4, HOY, "dd/mm/yyyy", True)
    celda(lu, f, 5, f"=D{f}-C{f}", "0"); celda(lu, f, 6, f"=MAX(0,E{f}-{REF['LUCP']})", "0")
    celda(lu, f, 7, f"=B{f}*F{f}/30.44*{REF['TCHM']}", "#,##0"); celda(lu, f, 8, f"=B{f}*E{f}/30.44*{REF['TCHM']}", "#,##0")
    f += 1
bb = f - 1
total(lu, f, "Total 2027 (hasta hoy)", {2: (f"=SUM(B{a}:B{bb})", "#,##0.0"), 7: (f"=SUM(G{a}:G{bb})", "#,##0"), 8: (f"=SUM(H{a}:H{bb})", "#,##0")})
LU27 = (f"'Lucro cesante'!$B${f}", f"'Lucro cesante'!$G${f}")
lu.cell(f + 1, 1, "Falta el plan de renovación 2026-27: las renovaciones que se corten de aquí en adelante sumarán lucro en 2027.").font = FNOTE

# ------------------------------------------------ Corte a siembra
cs_ = hoja("Corte a siembra", "Oportunidad corte a siembra: efecto en el TCH de la plantilla",
           "419 plantillas cosechadas 2019-2026 con siembra (maestro) y corte anterior conocidos; TCH con efectos de año, mes, edad y variedad.")
encabezado(cs_, 4, [("Días entre corte y siembra", 26), ("Plantillas", 10), ("TCH medio", 10), ("Frente a ≤60 días", 14)])
grupos = pl.groupby("sin_g").apply(lambda g: (len(g), np.average(g.tch_calc, weights=g.Area)))
efg = {"≤60": "base", "61-120": "+4,3 % (p=0,07)", "121-180": "+5,2 % (p=0,05)", ">180": "+2,0 % (p=0,44)"}
for i, kk in enumerate(["≤60", "61-120", "121-180", ">180"], 5):
    n, t = grupos.get(kk, (0, np.nan))
    cs_.cell(i, 1, kk).font = FB
    celda(cs_, i, 2, int(n), "0", True); celda(cs_, i, 3, float(t), "0.0", True); celda(cs_, i, 4, efg[kk], inp=True)
cs_.cell(10, 1, "Lineal: −0,8 % por cada 100 días (p=0,22, no significativo). En la diapositiva = 0 t.").font = FNB
encabezado(cs_, 12, [("Plantillas de 2026 (cosechadas)", 26), ("Área (ha)", 10), ("Días sin caña", 10), ("TCH", 14)])
f = 13
for rw in pl[pl.anio == 2026].sort_values("dias_sin_cana", ascending=False).itertuples():
    celda(cs_, f, 1, rw.ste, inp=True); celda(cs_, f, 2, float(rw.Area), "0.00", True); celda(cs_, f, 3, int(rw.dias_sin_cana), "0", True)
    celda(cs_, f, 4, float(rw.tch_calc), "0.0", True)
    f += 1
total(cs_, f, "Plantillas 2026 con más de 60 días entre corte y siembra (ha)", {2: (f'=SUMIFS(B13:B{f - 1},C13:C{f - 1},">60")', "#,##0")})
CS26 = f"'Corte a siembra'!$B${f}"
f += 2
encabezado(cs_, f, [("Plantillas sembradas desde oct-2025 (cosecha 2027)", 26), ("Área (ha)", 10), ("Días sin caña", 10), ("Siembra", 14)])
f += 1
a = f
for rw in fut.sort_values("dias_sin_cana", ascending=False).itertuples():
    celda(cs_, f, 1, rw.ste, inp=True); celda(cs_, f, 2, float(rw.area), "0.00", True); celda(cs_, f, 3, int(rw.dias_sin_cana), "0", True)
    celda(cs_, f, 4, pd.Timestamp(rw.siembra).date(), "dd/mm/yyyy", True)
    f += 1
total(cs_, f, "Plantillas 2027 con más de 60 días entre corte y siembra (ha)", {2: (f'=SUMIFS(B{a}:B{f - 1},C{a}:C{f - 1},">60")', "#,##0")})
CS27 = f"'Corte a siembra'!$B${f}"

# ------------------------------------------------ Fertilización
fe = hoja("Fertilización", "Oportunidad y calidad de la fertilización (base de validación de maduración Riopaila 2026)",
          "227 cortes de ene-ago 2026 cruzados con el ZQM. TCH frente al promedio 2020-25 de la misma suerte. Ninguna calificación bajó el TCH: efecto = 0.")
encabezado(fe, 4, [("Suerte", 10), ("Mes", 6), ("Área (ha)", 9), ("Oportunidad fertilización", 16), ("Suelo al fertilizar", 14), ("Calificación (1-5)", 10),
                   ("Oportunidad roturación", 16), ("Suelo al roturar", 14), ("TCH", 8), ("Frente a su promedio 2020-25", 12)])
f = 5
for rw in base.sort_values("ste").itertuples():
    for j_, v, fmt in ((1, rw.ste, None), (2, int(rw.Mes), "0"), (3, float(rw.Area), "0.00"), (4, rw.of, None), (5, rw.cf, None),
                       (6, float(rw.kf) if pd.notna(rw.kf) else None, "0"), (7, rw.orr, None), (8, rw.cr, None), (9, float(rw.tch_calc), "0.0"),
                       (10, float(np.exp(rw.rel) - 1), "+0.0%;-0.0%")):
        celda(fe, f, j_, v, fmt, True)
    f += 1
bb = f - 1
total(fe, f, "Fertilización TARDÍA 2026 (ha)", {3: (f'=SUMIFS(C5:C{bb},D5:D{bb},"TARDIO")', "#,##0")})
FE26 = f"Fertilización!$C${f}"
f += 2
fe.cell(f, 1, "Efectos (modelo con mes, edad, corte y tipo de cosecha): oportunidad TARDÍO +8,2 % (n=3); suelo húmedo +0,9 % (p=0,72); "
              "calificación 1-5 +0,6 % por punto (p=0,61). Roturación TARDÍO −6,2 % (n=6), confirma el control oficial.").font = FNOTE
fe.cell(f + 1, 1, "2027 (SimpleAgri, socas cortadas desde nov-2025): primera fertilización después de 120 días").font = FB
celda(fe, f + 1, 3, float(co.area[co.fert_c == ">120"].sum()), "#,##0", True)
FE27 = f"Fertilización!$C${f + 1}"

# ------------------------------------------------ Resumen
rs = wb.active
rs.title = "Resumen"
rs["A1"], rs["A1"].font = "Riopaila · impacto de la restricción presupuestal en el TCH", FTIT
rs["A2"], rs["A2"].font = ("Todas las cifras salen de las otras hojas (fórmulas). TCH = toneladas ÷ ha a cosechar del año. "
                           "2027: estado de labores a hoy; riego = escenario El Niño."), FNOTE
heads = [("Factor", 26), ("2026 · Ha afectadas", 12), ("2026 · Toneladas", 12), ("2026 · TCH", 10), ("2027 · Ha afectadas", 12), ("2027 · Toneladas", 12), ("2027 · TCH", 10), ("Hoja", 18)]
encabezado(rs, 4, heads)
lineas = [("Menor renovación", RV26, RV27, "Renovación"),
          ("Corte a siembra", (CS26, None), (CS27, None), "Corte a siembra"),
          ("Fertilización", (FE26, None), (FE27, None), "Fertilización"),
          ("Roturación tardía", RO26, RO27, "Roturación"),
          ("Control de malezas", ML26, ML27, "Malezas"),
          ("Falta de riego", RG26, RG27, "Riego y lluvia"),
          ("Precipitaciones", PR26, PR27, "Riego y lluvia"),
          ("Lucro cesante", LU26, LU27, "Lucro cesante")]
for i, (nom, a26_, a27_, hj) in enumerate(lineas, 5):
    rs.cell(i, 1, nom).font = FB
    for col, (area, ton) in ((2, a26_), (5, a27_)):
        x = rs.cell(i, col, f"={area}")
        x.font, x.number_format = FLINK, "#,##0"
        y = rs.cell(i, col + 1, f"={ton}" if ton else 0)
        y.font, y.number_format = (FLINK if ton else FIN), "#,##0"
        z = rs.cell(i, col + 2, f"={L(col + 1)}{i}/{REF['HA26'] if col == 2 else REF['HA27']}")
        z.font, z.number_format = FB, "0.0"
    rs.cell(i, 8, hj).font = FNOTE
ft = 5 + len(lineas)
total(rs, ft, "Total", {3: (f"=SUM(C5:C{ft - 1})", "#,##0"), 4: (f"=C{ft}/{REF['HA26']}", "0.0"),
                        6: (f"=SUM(F5:F{ft - 1})", "#,##0"), 7: (f"=F{ft}/{REF['HA27']}", "0.0")})
rs.cell(ft + 2, 1, "Ha a cosechar").font = FB
rs.cell(ft + 2, 2, f"={REF['HA26']}").number_format = "#,##0"
rs.cell(ft + 2, 5, f"={REF['HA27']}").number_format = "#,##0"
rs.cell(ft + 3, 1, "Negativo = a favor. Corte a siembra y fertilización: sin efecto medible (0 t). Los factores se midieron por separado y pueden traslaparse.").font = FNOTE

# ------------------------------------------------ Cómo se hizo
ch = wb.create_sheet("Cómo se hizo")
ch.column_dimensions["A"].width = 130
txt_ = [
    ("Fuentes", True),
    ("Cosechas: ZQM histórico 2011-2025, G&P Suertes_26 (hoja ZQM101, ene-sep 2026) y app de Cosecha (octubre). Maestro de suertes de Riopaila. "
     "Estimados del motor (corrida 16-oct-2026).", False),
    ("Renovación: «Areas no renovadas.xlsx» (plan 2025: 124 suertes; 93 marcadas NO, de las que 8 se renovaron tarde).", False),
    ("Roturación: «CONTROL DE ROTURACION. est 2025 con formatos.xlsm» (cortes oct-2023 a ene-2026, criterio de Riopaila: IDEAL ≤15 días, "
     "NORMAL 16-30, TARDÍO >30) y SimpleAgri (planillas de roturación 540-544) para los cortes posteriores, corridas 10 días porque la planilla "
     "se confirma después de la labor.", False),
    ("Malezas: fecha real de aplicación de las bases de aplicaciones de los contratistas («LABOR APLICACIONES 31-12-2025» y «BASE DE DATOS "
     "APLICACIONES 2026», 14-ago a 10-oct-2026); donde no hay (feb-ago 2026 o antes de 2025), SimpleAgri corrida 7 días, el desfase medido entre "
     "la aplicación y la confirmación de la planilla. Fertilización: base de validación de maduración Riopaila 2026 y SimpleAgri.", False),
    ("Riego: «CONSOLIDADO DE RIEGOS 2026-.xlsx» (área regada y lámina por riego). Clima: EMA La Paila y Bugalagrande (repositorio de "
     "precipitación de Riopaila) completadas con el Meteoportal de Cenicaña.", False),
    ("Lucro cesante y corte a siembra: fechas de siembra del maestro y fecha del corte anterior (ZQM, archivo de áreas no renovadas).", False),
    ("", False),
    ("Método", True),
    ("Cada efecto se midió con cosechas reales de Riopaila (2024-2026 para las labores), comparando cada suerte con su propio promedio 2020-25 y "
     "descontando año, mes, edad y número de corte. Las toneladas de 2026 y 2027 aplican ese efecto a las suertes de cada año (hojas por factor).", False),
    ("2026 = cosechas del año (ene-oct reales y nov-dic con el TCH del motor donde aplica). 2027 = caña que se está levantando, con el estado de "
     "las labores a hoy: lo pendiente cuenta con los días que ya pasaron, así que es un mínimo.", False),
    ("Programas: carpeta «01 Estimados y TCH/Caida TCH 2026 - programas» (no_renovadas_2026, roturacion, labores_sa, labores_perdidas, "
     "riego_2026, riego_2027, siembra_lucro, impacto_excel).", False),
]
for i, (t, neg) in enumerate(txt_, 1):
    x = ch.cell(i, 1, t)
    x.font, x.alignment = (FNB if neg else FB), WR

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title not in ("Resumen", "Cómo se hizo", "Efectos y supuestos")
wb.active = 0
wb.save(SAL)
print("guardado:", SAL)
print("riego 2026 por etapa:", [(k, round(100 * cf, 2), round(a), round(b_)) for k, cf, a, b_ in riego26], "· ene-sep", round(ton_es), "t", round(ha_es), "ha")
print("arranque coef", round(100 * c_arr, 2), "%/100 mm · lluvia", round(llu26), "vs", round(lluref))
