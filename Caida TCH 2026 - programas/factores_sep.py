# -*- coding: utf-8 -*-
"""Caída del TCH 2026: cuánto pesa cada factor, por planta, con las cosechas hasta SEPTIEMBRE (G&P Suertes_26).

Reconstruye el modelo del 29-sep-2026 (m2_modelos.py / m3_castilla_cen.py, recuperados en recuperados_29sep/):
log TCH con efectos fijos de suerte, año y mes + edad, soca, variedad, incendio, cosecha mecánica (en Riopaila
con una brecha por año), lluvia de los meses 0-3 del ciclo (arranque) y días con Tmax > 32 °C en los 30 días
antes de cosechar (final). Ponderado por área, errores agrupados por suerte. Aportes de 2026 frente a 2020-25 en
los mismos meses; resultado = cada suerte frente a su propio promedio 2020-25. Diatraea aparte (solo Castilla).

Fuentes (solo lectura):
  Riopaila: ZQM histórico 2011-2025 + hoja ZQM101 del G&P 2026 (ene-sep); EMA La Paila / Bugalagrande (repo
    mar96ospina-a/Precipitacion-riopaila, hasta 8-sep-2026) + Meteoportal de Cenicaña (diario 2026).
  Castilla: Historico por suerte (2020-2026) Castilla.xlsx (todas las tenencias, a agosto) + septiembre de tierras
    propias del G&P; pluviómetros y estaciones de Recursos Hidricos; CSV diario de Cenicaña + Meteoportal; barrenador.

Uso: <tch-motor>\\.venv-gee\\Scripts\\python.exe factores_sep.py <carpeta del repo Precipitacion-riopaila> [<salida>]
"""
import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
from tchmotor import dataset_gee as dg  # noqa: E402
from tchmotor.agronomia import cargar_var_map, map_variedad  # noqa: E402

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
pd.set_option("display.width", 250)
pd.set_option("display.max_columns", 40)

REPO = Path(sys.argv[1])
SAL = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(".")
DOC = Path(r"C:\Users\Agr349\Documents")
GP = DOC / "G&P Suertes_26.xlsx"
CAS = DOC / r"RIOPAILA AGRICOLA\04 Presupuesto, inversion y junta\JUNTA DIRECTIVA\CASTILLA"
CEN = DOC / r"RIOPAILA AGRICOLA\Clima Cenicaña\Clima_Cenicana_Riopaila_Castilla_1993-2026.xlsx"
EDADES = ["<9", "9-10", "10-11", "11-12", "12-13", "13-14", "14-15", ">15"]
ULTIMO_MES = 9
PERIODOS = [("ene-ago", range(1, 9)), ("ene-sep", range(1, 10)), ("ene-feb", [1, 2]), ("mar-abr", [3, 4]),
            ("may-jun", [5, 6]), ("jul-ago", [7, 8]), ("sep", [9])]
dias = pd.date_range("2009-01-01", "2026-10-18", freq="D")
wm = lambda g, k: np.average(g[k].dropna(), weights=g.Area[g[k].notna()]) if g[k].notna().any() else np.nan  # noqa: E731


def fitc(f, d, grp):
    mo = smf.wls(f, data=d, weights=d.Area)
    rows = mo.data.row_labels
    return mo.fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(d.loc[rows, grp])[0]})


def hot(serie, f):
    s = serie.loc[f - pd.Timedelta(days=30): f - pd.Timedelta(days=1)]
    return (s > 32).sum() * 30 / s.notna().sum() if s.notna().sum() >= 20 else np.nan


def rain(serie, a):
    s = serie.loc[a: a + pd.Timedelta(days=90)]
    return s.sum() * 91 / s.notna().sum() if s.notna().sum() >= 70 else np.nan


def relativo(d, ut, y, anio):
    """y menos el promedio de la suerte en 2020-25 (sin el propio corte si es de 2020-25)."""
    eb = d[anio].between(2020, 2025)
    b = d[eb].groupby(ut)[y].agg(["sum", "count"])
    d = d.join(b, on=ut)
    rel = np.where(eb, d[y] - (d["sum"] - d[y]) / (d["count"] - 1), d[y] - d["sum"] / d["count"])
    rel = pd.Series(rel, index=d.index)
    rel[(eb & (d["count"] <= 1)) | d["count"].isna()] = np.nan
    return rel


# ---------------------------------------------------------------- Meteoportal 2026 (diario)
wb = openpyxl.load_workbook(CEN, read_only=True, data_only=True)
it = wb["Diario 2026"].iter_rows(values_only=True)
enc = next(it)
M = pd.DataFrame(list(it), columns=enc)
M["Fecha"] = pd.to_datetime(M.Fecha)
Mt = M.pivot_table(index="Fecha", columns="Código", values="Temperatura máxima (°C)")
Ml = M.pivot_table(index="Fecha", columns="Código", values="Lluvia (mm)")


# ---------------------------------------------------------------- G&P 2026, hoja ZQM101
def leer_gp():
    w = openpyxl.load_workbook(GP, read_only=True, data_only=True)
    filas, hdr, vacias = [], None, 0
    for r in w["ZQM101"].iter_rows(values_only=True):
        if hdr is None:
            if any(str(c).strip() == "Ubic Técnica" for c in r if c is not None):
                hdr = dg._dedup([str(x).strip() if x is not None else f"c{j}" for j, x in enumerate(r)])
            continue
        if all(v is None for v in r[:6]):
            vacias += 1
            if vacias > 50:
                break
            continue
        vacias = 0
        filas.append(r)
    w.close()
    g = pd.DataFrame(filas, columns=hdr)
    g["AÑO"] = 2026
    return g


def ciclos(df):
    """Como dg.cortes_zqm, conservando tipo y causa de cosecha: un ciclo por suerte + soca (meses consecutivos ≤2),
    con la fila de mayor TC-Total (el acumulado)."""
    d = df.copy()
    for c in ["AÑO", "Mes", "Corte", "Area", "Edad", "TC-Total"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    for c in ["Ubic Técnica", "Empresa", "Variedad Principal", "Tipo Cosecha", "Causa Cosecha"]:
        d[c] = d[c].astype(str).str.strip()
    d = d[d["Empresa"].str.startswith(("02", "03")) & d.Mes.isin(range(1, 13))]
    d["planta"] = np.where(d["Empresa"].str.startswith("03"), "riopaila", "castilla")
    d["t"] = d["AÑO"] * 12 + d["Mes"] - 1
    d = d.dropna(subset=["t", "Corte"])
    d = d[~d["Ubic Técnica"].isin(["", "None", "nan"])].sort_values(["Ubic Técnica", "Corte", "t"])
    gap = d.groupby(["Ubic Técnica", "Corte"])["t"].diff()
    d["ciclo"] = (gap.isna() | (gap > 2)).astype(int).groupby([d["Ubic Técnica"], d["Corte"]]).cumsum()
    idx = d.assign(_tc=d["TC-Total"].fillna(-1)).groupby(["Ubic Técnica", "Corte", "ciclo"])["_tc"].idxmax()
    c = d.loc[idx].copy()
    c["tch_calc"] = c["TC-Total"] / c["Area"]
    return c[c.Area.gt(0) & c.tch_calc.between(20, 300) & c.Edad.between(8, 24)]


print("leyendo ZQM y G&P…", flush=True)
hist = dg._hoja(dg.ZQM_HIST, 0)
gp = leer_gp()
cols = ["AÑO", "Mes", "Ubic Técnica", "Empresa", "Corte", "Variedad Principal", "Area", "Edad", "TC-Total",
        "Tipo Cosecha", "Causa Cosecha"]
for d_ in (hist, gp):
    for k in cols:
        if k not in d_.columns:
            d_[k] = np.nan
Z = ciclos(pd.concat([hist[cols], gp[cols]], ignore_index=True))
print("cortes del ZQM por planta y año (2024-26):",
      Z[Z["AÑO"] >= 2024].groupby(["planta", "AÑO"]).size().to_dict(), flush=True)


# ---------------------------------------------------------------- modelo común
def aportes(mo, q, r, ed, mec, d):
    for k in EDADES:
        q["e_" + k] = (q.edad_b == k).astype(float)
        r["e_" + k] = (r.edad_b == k).astype(float)
    a = {"arranque (lluvia 0-3 m)": mo.params["llu03_100"] * (wm(q, "llu03_100") - wm(r, "llu03_100")),
         "calor (días >32 °C)": mo.params["d32"] * (wm(q, "d32") - wm(r, "d32")),
         "edad": sum(ed[k] * (wm(q, "e_" + k) - wm(r, "e_" + k)) for k in EDADES),
         "incendio": mo.params["inc"] * (wm(q, "inc") - wm(r, "inc")),
         "soca": mo.params["Corte"] * (wm(q, "Corte") - wm(r, "Corte"))}
    if mec == "por año":
        a["cosecha mecánica"] = mo.params["mec_2026"] * wm(q, "mec_2026") - sum(
            mo.params[f"mec_{y}"] * wm(r, f"mec_{y}") for y in range(2020, 2026))
    else:
        a["cosecha mecánica"] = mo.params["mec"] * (wm(q, "mec") - wm(r, "mec"))
    bv = {k.split("[T.")[1][:-1]: v for k, v in mo.params.items() if k.startswith("C(var)")}
    s26, sR = q.groupby("var").Area.sum() / q.Area.sum(), r.groupby("var").Area.sum() / r.Area.sum()
    a["variedad"] = sum((s26.get(v, 0) - sR.get(v, 0)) * bv.get(v, 0.0) for v in set(s26.index) | set(sR.index))
    return a


def informe(nombre, d, est, grp, ton, mec, diatraea):
    ED = "C(edad_b, Treatment('12-13'))"
    mec_f = " + ".join(f"mec_{y}" for y in range(2011, 2027)) if mec == "por año" else "mec"
    f = f"y ~ C({grp}) + C(anio) + C(Mes) + {ED} + C(soca_b) + C(var) + inc + Corte + {mec_f} + llu03_100 + d32"
    mo = fitc(f, est, grp)
    ed = {k: mo.params.get(f"{ED}[T.{k}]", 0.0) for k in EDADES}
    print(f"\n==================== {nombre} (n={int(mo.nobs)}, R²={mo.rsquared:.2f})")
    for k in ["llu03_100", "d32", "inc", "Corte"] + (["mec_2024", "mec_2025", "mec_2026"] if mec == "por año" else ["mec"]):
        print(f"  {k:<10} {100 * (np.exp(mo.params[k]) - 1):+6.2f}%  (p={mo.pvalues[k]:.3f})")
    filas = []
    for lab, meses in PERIODOS:
        q = d[(d.anio == 2026) & d.Mes.isin(meses)].copy()
        r = d[d.anio.between(2020, 2025) & d.Mes.isin(meses)].copy()
        if not len(q):
            continue
        T, HA = q[ton].sum(), q.Area.sum()
        gap = wm(q.dropna(subset=["rel"]), "rel") - wm(r.dropna(subset=["rel"]), "rel")
        a = aportes(mo, q, r, ed, mec, d)
        a["Diatraea"] = diatraea(q, r)
        row = {"periodo": lab, "cortes": len(q), "ha": HA, "t": T, "TCH": T / HA, "resultado": 100 * gap}
        row.update({k: 100 * v for k, v in a.items()})
        row["sin explicar"] = 100 * gap - sum(100 * v for v in a.values() if pd.notna(v))
        row["t_resultado"] = -gap * T
        row.update({"t_" + k: -v * T for k, v in a.items() if pd.notna(v)})
        row["t_sin explicar"] = row["t_resultado"] - sum(row["t_" + k] for k in a if "t_" + k in row)
        row.update({"lluvia 0-3 m 2026": wm(q, "llu03"), "lluvia 0-3 m ref": wm(r, "llu03"),
                    "días >32 2026": wm(q, "d32"), "días >32 ref": wm(r, "d32"),
                    "% mecánica 2026": 100 * wm(q, "mec"), "% mecánica ref": 100 * wm(r, "mec"),
                    "% incendio 2026": 100 * wm(q, "inc"), "% incendio ref": 100 * wm(r, "inc")})
        filas.append(row)
    t = pd.DataFrame(filas).set_index("periodo")
    cols_ = ["cortes", "ha", "TCH", "resultado", "cosecha mecánica", "calor (días >32 °C)", "arranque (lluvia 0-3 m)",
             "Diatraea", "variedad", "edad", "incendio", "soca", "sin explicar"]
    print(t[cols_].round(1).to_string())
    print(t[[c for c in t.columns if c.startswith("t_")]].round(-2).to_string())
    print(t[["lluvia 0-3 m 2026", "lluvia 0-3 m ref", "días >32 2026", "días >32 ref", "% mecánica 2026",
             "% mecánica ref", "% incendio 2026", "% incendio ref"]].round(1).to_string())
    if mec == "por año":
        print("  brecha mecánica por año:", {y: round(100 * (np.exp(mo.params[f"mec_{y}"]) - 1), 1) for y in range(2020, 2027)})
    return mo, t


def preparar(c, var_col, top_n=60):
    c["y"] = np.log(c.tch_calc)
    c["anio"] = c.anio.astype(int)
    c["Mes"] = c.Mes.astype(int)
    c["edad_b"] = pd.cut(c.Edad, [0, 9, 10, 11, 12, 13, 14, 15, 40], labels=EDADES).astype(str)
    c["soca_b"] = pd.cut(c.Corte.fillna(3), [0, 1, 2, 3, 4, 5, 7, 99]).astype(str)
    c["Corte"] = c.Corte.fillna(3)
    top = c[var_col].value_counts()
    c["var"] = c[var_col].where(c[var_col].map(top) >= top_n, "otra")
    c["llu03_100"] = c.llu03 / 100
    return c


# ================================================================ RIOPAILA
e = pd.read_csv(REPO / "clima_estaciones.csv", sep=";")
e["f"] = pd.to_datetime(e.fc_obsrvcion)
EMA = {}
for est, cod in (("La Paila", "PAI"), ("Bugalagrande", "BLG")):
    d = e[e.nm_ema == est].set_index("f").sort_index()
    d = d[~d.index.duplicated()].reindex(dias)
    falta = d.index >= "2026-01-01"
    d.loc[falta & d.tmp_mxma.isna(), "tmp_mxma"] = Mt[cod].reindex(d.index)[falta & d.tmp_mxma.isna()]
    d.loc[falta & d.prcptcion.isna(), "prcptcion"] = Ml[cod].reindex(d.index)[falta & d.prcptcion.isna()]
    EMA[est] = d[["prcptcion", "tmp_mxma"]].rename(columns={"prcptcion": "llu", "tmp_mxma": "tmax"})
es = pd.read_csv(REPO / "estaciones_suertes.csv", sep=";", encoding="latin-1", skiprows=1)
est_de = {str(s).strip(): ("Bugalagrande" if "buga" in str(x).lower() else "La Paila")
          for s, x in zip(es.iloc[:, 1], es.iloc[:, 2])}

c = Z[Z.planta == "riopaila"].copy()
c["ste"] = c["Ubic Técnica"]
c["anio"] = c["AÑO"]
c["cosecha"] = pd.to_datetime(dict(year=c["AÑO"].astype(int), month=c.Mes.astype(int), day=15))
c["ini"] = c.cosecha - pd.to_timedelta(c.Edad * 30.44, unit="D")
c["est"] = c.ste.map(est_de).fillna("La Paila")
c["llu03"] = [rain(EMA[s].llu, a) for s, a in zip(c.est, c.ini)]
c["d32"] = [hot(EMA[s].tmax, f) for s, f in zip(c.est, c.cosecha)]
c["mec"] = (c["Tipo Cosecha"] == "Mecanico").astype(int)
c["inc"] = (c["Causa Cosecha"] == "Incendio").astype(int)
for y in range(2011, 2027):
    c[f"mec_{y}"] = c.mec * (c.anio == y)
c = preparar(c, "Variedad Principal")
c = c[c.anio >= 2011]
c["rel"] = relativo(c, "ste", "y", "anio")
c.to_pickle(SAL / "rio_factores_sep.pkl")
print("\nRiopaila 2026 por mes (cortes, % mecánica):",
      c[c.anio == 2026].groupby("Mes").apply(lambda g: (len(g), round(100 * wm(g, "mec")))).to_dict())
est_r = c.dropna(subset=["llu03", "d32"])
mo_r, t_r = informe("RIOPAILA", c[c.anio >= 2020], est_r, "ste", "TC-Total", "por año", lambda q, r: 0.0)
t_r.to_pickle(SAL / "tabla_riopaila_sep.pkl")

# ================================================================ CASTILLA
w = openpyxl.load_workbook(CAS / "Historico por suerte (2020-2026) Castilla.xlsx", read_only=True, data_only=True)
it = w["Base"].iter_rows(values_only=True)
enc = [str(x).strip() if x is not None else f"c{j}" for j, x in enumerate(next(it))]
h = pd.DataFrame([r for r in it if any(v is not None for v in r)], columns=dg._dedup(enc))
w.close()
for k in ["MES", "AÑO", "Corte", "Area", "TONS REAL", "TCH", "Edad"]:
    h[k] = pd.to_numeric(h[k], errors="coerce")
for k in ["TIPO CORTE", "Tenencia", "Estacion", "VARIEDAD", "Incendio Y/O Programa"]:
    h[k] = h[k].astype(str).str.strip().str.upper()
h["ut"] = h["U.T."].astype(str).str.strip()
h["fuente"] = "histórico por suerte"

# septiembre de tierras propias desde el G&P (ciclo completo; si empezó en agosto, reemplaza la fila de agosto)
vm = cargar_var_map()
g9 = Z[(Z.planta == "castilla") & (Z["AÑO"] == 2026) & (Z.Mes == ULTIMO_MES)].copy()
estacion = h.sort_values(["AÑO", "MES"]).groupby("ut").Estacion.last()
sep = pd.DataFrame({
    "ut": g9["Ubic Técnica"].values, "MES": ULTIMO_MES, "AÑO": 2026, "Corte": g9.Corte.values, "Area": g9.Area.values,
    "TONS REAL": g9["TC-Total"].values, "TCH": g9.tch_calc.values, "Edad": g9.Edad.values,
    "TIPO CORTE": g9["Tipo Cosecha"].str.upper().values, "Tenencia": "1.AGRICOLA",
    "Incendio Y/O Programa": g9["Causa Cosecha"].str.upper().values,
    "VARIEDAD": [str(map_variedad(v, vm) or v).upper() for v in g9["Variedad Principal"]],
    "fuente": "G&P septiembre"})
sep["Estacion"] = sep.ut.map(estacion).fillna("")
ago_dup = (h["AÑO"] == 2026) & (h.MES == ULTIMO_MES - 1) & h.Tenencia.str.contains("AGRIC") & \
    h.set_index(["ut", "Corte"]).index.isin(sep.set_index(["ut", "Corte"]).index)
print(f"\nCastilla: {len(sep)} cortes de septiembre (G&P, tierras propias); {int(ago_dup.sum())} venían de agosto y se reemplazan")
# control: el histórico por suerte y el G&P coinciden en ene-ago (tierras propias)
g8 = Z[(Z.planta == "castilla") & (Z["AÑO"] == 2026) & (Z.Mes <= 8)]
ha = h[(h["AÑO"] == 2026) & h.Tenencia.str.contains("AGRIC")]
mm = ha.merge(g8[["Ubic Técnica", "Corte", "tch_calc"]], left_on=["ut", "Corte"], right_on=["Ubic Técnica", "Corte"])
print(f"  control ene-ago: {len(mm)} de {len(ha)} cortes del histórico cruzan con el G&P; TCH igual (±1) en "
      f"{100 * ((mm.TCH - mm.tch_calc).abs() <= 1).mean():.0f} %")
h = pd.concat([h[~ago_dup], sep], ignore_index=True)
h = h[h.TCH.between(20, 300) & (h.Area > 0) & h.MES.between(1, 12) & h.Edad.between(8, 24)].copy()
h["anio"], h["Mes"] = h["AÑO"].astype(int), h.MES.astype(int)
h["cosecha"] = pd.to_datetime(dict(year=h.anio, month=h.Mes, day=15))
h["ini"] = h.cosecha - pd.to_timedelta(h.Edad * 30.44, unit="D")
h["tch_calc"] = h.TCH
h["mec"] = (h["TIPO CORTE"] == "MECANICO").astype(int)
h["inc"] = h["Incendio Y/O Programa"].str.contains("INC").astype(int)
h["ten"] = h.Tenencia.str.split(".").str[1].str.strip()

# lluvia: pluviómetro de la suerte (mediana regional si no tiene)
p = pd.read_csv(CAS / "Recursos Hidricos" / "Precipitacion diaria por pluviometro.csv", sep=";", decimal=",",
                encoding="utf-8-sig", low_memory=False)
p["Fecha"] = pd.Timestamp("1899-12-30") + pd.to_timedelta(pd.to_numeric(p.Fecha, errors="coerce"), unit="D")
p["mm"] = pd.to_numeric(p.Precipitacion.astype(str).str.replace(",", "."), errors="coerce")
p["PLUV"] = p.PLUV.astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
L = p.dropna(subset=["Fecha", "mm"]).pivot_table(index="Fecha", columns="PLUV", values="mm", aggfunc="mean").reindex(dias)
L["REGION"] = L.median(axis=1)
sp = pd.read_excel(CAS / "Recursos Hidricos" / "Pluviometro suerte.xlsx", sheet_name="Suerte-Pluviometro")
sp.columns = ["ut", "pluv", "est_nombre"][: len(sp.columns)]
sp["ut"] = sp.ut.astype(str).str.strip()
sp["pluv"] = sp.pluv.astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
h = h.merge(sp[["ut", "pluv"]].drop_duplicates("ut"), on="ut", how="left")
h["pluv"] = np.where(h.pluv.isin(L.columns), h.pluv, "REGION")
print(f"  lluvia: pluviómetros hasta {L.drop(columns='REGION').dropna(how='all').index.max().date()}; "
      f"{100 * (h.pluv != 'REGION').mean():.0f} % de los cortes con pluviómetro propio")
h["llu03"] = [rain(L[pl], a) if a >= pd.Timestamp("2017-01-01") else np.nan for pl, a in zip(h.pluv, h.ini)]

# calor: estación de Cenicaña de cada suerte (promedio regional si no tiene)
cs = pd.read_csv(CAS / "Consultas Climatologícas Cenicaña(Datos Originales).csv", sep=";", encoding="latin-1", dtype=str)
cs["fecha"] = pd.to_datetime(cs.Fecha, format="%d/%m/%Y", errors="coerce")
cs["tmax"] = pd.to_numeric(cs["Temperatura Maxima"].astype(str).str.strip().str.replace(",", "."), errors="coerce")
NOM = {"candelaria": "Candelaria", "pradera": "Pradera", "el tiple": "El Tiple", "cenicana": "Cenicana", "florida": "Florida",
       "bocas del palo": "Bocas del Palo", "ortigal": "Ortigal", "miranda": "Miranda", "el naranjo": "El Naranjo",
       "sder de quilichao": "Santander de Quilichao"}
cs["est"] = cs.Estacion.str.strip().str.lower().map(NOM).fillna(cs.Estacion.str.strip())
cs.loc[~cs.tmax.between(15, 42), "tmax"] = np.nan
TM = cs.pivot_table(index="fecha", columns="est", values="tmax", aggfunc="mean").reindex(dias)
for nombre_, cod in (("Pradera", "PRA"), ("Florida", "FLO"), ("Candelaria", "CAN"), ("El Tiple", "TIP"),
                     ("Cenicana", "CEN"), ("Miranda", "MIR")):
    if nombre_ in TM and cod in Mt:
        falta = TM[nombre_].isna() & (TM.index >= "2026-01-01")
        TM.loc[falta, nombre_] = Mt[cod].reindex(TM.index)[falta]
TM["REGION"] = TM.drop(columns=["Corinto"], errors="ignore").mean(axis=1)
COD = {"MIR": "Miranda", "BDP": "Bocas del Palo", "ORT": "Ortigal", "CEN": "Cenicana", "PRA": "Pradera",
       "CAN": "Candelaria", "SDQ": "Santander de Quilichao", "NAR": "El Naranjo", "TIP": "El Tiple", "FLO": "Florida"}
h["est_c"] = h.Estacion.map(COD).fillna("REGION")


def hot_c(col, f):
    for k in (col, "REGION"):
        if k in TM:
            v = hot(TM[k], f)
            if pd.notna(v):
                return v
    return np.nan


h["d32"] = [hot_c(k, f) for k, f in zip(h.est_c, h.cosecha)]

# barrenador: evaluación de 150 días antes a 20 días después de la cosecha
w = openpyxl.load_workbook(CAS / "Historico Barrenadores y salivazo.xlsm", read_only=True, data_only=True)
rows = list(w["Barrenadores"].iter_rows(values_only=True))
w.close()
b = pd.DataFrame(rows[1:], columns=[" ".join(str(x).split()) for x in rows[0]]).dropna(how="all")
b["fecha"] = pd.to_datetime(b.Fecha, errors="coerce")
b["II"] = pd.to_numeric(b.II_Diatraea, errors="coerce")
b["ut"] = b["Sector suerte"].astype(str).str.strip()
b = b[b.Tipo.astype(str).str.strip() == "Evaluada"].dropna(subset=["II", "fecha"])
h["k"] = np.arange(len(h))
mb = b[["ut", "fecha", "II"]].merge(h[["ut", "cosecha", "k"]], on="ut")
mb = mb[(mb.fecha >= mb.cosecha - pd.Timedelta(days=150)) & (mb.fecha <= mb.cosecha + pd.Timedelta(days=20))]
mb = mb.sort_values("fecha").drop_duplicates("k", keep="last")
h = h.merge(mb[["k", "II"]], on="k", how="left")

h = preparar(h, "VARIEDAD")
h["rel"] = relativo(h, "ut", "y", "anio")
h.to_pickle(SAL / "cas_factores_sep.pkl")
A = h[h.ten == "AGRICOLA"].copy()
print("  tierras propias 2026 por mes (cortes, % mecánica, % incendio, % con Diatraea):",
      A[A.anio == 2026].groupby("Mes").apply(lambda g: (len(g), round(100 * wm(g, "mec")), round(100 * wm(g, "inc")),
                                                        round(100 * g.Area[g.II.notna()].sum() / g.Area.sum()))).to_dict())
mA = A.dropna(subset=["II", "rel"]).copy()
mA["am"] = mA.cosecha.dt.strftime("%Y-%m")
bI = fitc("rel ~ II + C(am) + Edad + Corte + mec", mA, "ut").params["II"]
print(f"  Diatraea: {100 * bI:+.2f} % de TCH por punto de %I.I.")


def dia_c(q, r):
    e26, eR = mA[(mA.anio == 2026) & mA.Mes.isin(q.Mes.unique())], mA[mA.anio.between(2020, 2025) & mA.Mes.isin(q.Mes.unique())]
    return bI * (wm(e26, "II") - wm(eR, "II")) if len(e26) >= 5 and len(eR) >= 5 else np.nan


est_c = h.dropna(subset=["llu03", "d32"])
mo_c, t_c = informe("CASTILLA tierras propias (estimado con todas las tenencias)", A, est_c, "ut", "TONS REAL", "uno", dia_c)
t_c.to_pickle(SAL / "tabla_castilla_sep.pkl")
