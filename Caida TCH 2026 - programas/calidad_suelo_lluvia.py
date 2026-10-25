# -*- coding: utf-8 -*-
"""Base de maduración Riopaila: columnas BI-BL (Oport. Rot, Calidad Rot., Oport. Fert, Calidad Fert) de las filas vacías.

Oport. Rot = control oficial de roturación, 1.ª labor: TARDIO → TARDIO; IDEAL o NORMAL → IDEAL-NORMAL; REALIZADO con 0 días
entre programación y roturación (la fórmula del control deja la celda en blanco) → IDEAL-NORMAL. Oport. Fert no se llena: no
hay control de fertilización y los días a la fertilización de SimpleAgri no separan IDEAL-NORMAL de TARDIO.

Calidad (HUMEDO / FRIABLE), regla de Cañicultura: si llovió en el pluviómetro de la suerte cerca de la fecha de la labor → HUMEDO; si no → FRIABLE.
Pluviómetro de cada suerte = columna PLUVIOMETRO del control de roturación (Suerte_Pluv). Lluvia diaria = «Ponderado de
Precipitaciones año 2026» (2026 completo y jun-sep 2025); el resto de 2025 (oct-dic, hojas vacías en el ponderado) sale del
registro diario de los mismos pluviómetros (repo Precipitacion-riopaila), que es idéntico al ponderado donde ambos tienen dato.
Fecha de la roturación: control oficial (1.ª labor); si el corte no está en el control, SimpleAgri (fecha de confirmación − 10
días). Fecha de la fertilización: SimpleAgri (1.ª fertilización del ciclo, confirmación − desfase calibrado).
La ventana y el umbral se calibran contra las filas ene-ago que ya traen la calificación.

Uso: python calidad_suelo_lluvia.py <base.csv> <CONTROL DE ROTURACION.xlsm> <Ponderado.xlsx> <precipitacion.csv.csv>
                                    <carpeta planillas SimpleAgri> <salida.csv>
"""
import pickle
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
pd.set_option("display.width", 220)
BASE, CTRL, POND, REPO, SA, SAL = sys.argv[1:7]

# ---------------------------------------------------------------- lluvia diaria por pluviómetro
# hojas del ponderado por año (identificadas comparando día a día con el registro y con «Precip. Acum 2026»)
HOJAS = {2026: ["Precip.Enero", "Precip. Febrero", "Precip. Marzo", "Precip.Abril ", "Precip Mayo", "Precip.Junio ",
                "Precip. Julio", "Precip. Agosto", "Precip.Septiembre ", "Precip.Octu 2026"],
         2025: ["Precip Jun", "Precip Jul", "Precip Agosto", "Precip Septiembre"]}
MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "oct": 10}
wb = openpyxl.load_workbook(POND, read_only=True, data_only=True)
filas = []
for anio, hojas in HOJAS.items():
    for nom in hojas:
        mes = next(v for k, v in MESES.items() if k in nom.lower())
        F = list(wb[nom].iter_rows(values_only=True))
        h = next(i for i, r in enumerate(F[:10]) if r and "PLUV No" in [str(x).strip() for x in r if x is not None])
        enc = [str(x).strip() if x is not None else "" for x in F[h]]
        jp = enc.index("PLUV No")
        cols = {}
        for j in range(jp + 1, len(enc)):
            if re.fullmatch(r"\d+(\.0)?", enc[j]):
                dia = int(float(enc[j]))
                if dia in cols.values():
                    break
                cols[j] = dia
        for r in F[h + 1:]:
            pl = str(r[jp]).strip().replace(".0", "") if r and r[jp] is not None else ""
            if re.fullmatch(r"\d+", pl):
                for j, dia in cols.items():
                    # la plantilla trae 31 columnas en todos los meses: el 31 de septiembre no existe
                    if j < len(r) and isinstance(r[j], (int, float)) and dia <= pd.Timestamp(anio, mes, 1).days_in_month:
                        filas.append((pl, pd.Timestamp(anio, mes, dia), float(r[j])))
P = pd.DataFrame(filas, columns=["PLUV", "Fecha", "mm"])
R = pd.read_csv(REPO, sep=";", encoding="latin-1", usecols=[1, 2, 3], header=0, dtype=str)
R.columns = ["PLUV", "Fecha", "mm"]
R["Fecha"] = pd.to_datetime(R.Fecha.str.strip(), format="%d/%m/%Y", errors="coerce")
R["mm"] = pd.to_numeric(R.mm.str.replace(",", "."), errors="coerce")
R["PLUV"] = R.PLUV.str.strip()
R = R.dropna()
# control de calidad: donde se cruzan (jun-sep 2025, ene-jul 2026) deben ser iguales
x = P.merge(R, on=["PLUV", "Fecha"], suffixes=("", "_r"))
x = x[x.Fecha < "2026-08-01"]
print(f"lluvia: ponderado {P.Fecha.min().date()} a {P.Fecha.max().date()} ({P.PLUV.nunique()} pluviómetros); "
      f"igual al registro en {100 * (abs(x.mm - x.mm_r) < 0.5).mean():.1f} % de {len(x):,} días-pluviómetro comunes")
R = R[(R.Fecha < "2026-01-01") & ~((R.Fecha >= "2025-06-01") & (R.Fecha < "2025-10-01"))]
L = pd.concat([P, R]).groupby(["Fecha", "PLUV"]).mm.mean().unstack().sort_index()
L = L.reindex(pd.date_range(L.index.min(), L.index.max()))
print(f"serie diaria final: {L.index.min().date()} a {L.index.max().date()}")


def lluvia(pl, f, w):
    """mm en el pluviómetro desde w días antes hasta el día de la labor (NaN si no hay dato)."""
    if pd.isna(f) or pl not in L.columns:
        return np.nan
    s = L[pl].loc[f - pd.Timedelta(days=w): f]
    return s.sum() if s.notna().any() else np.nan


# ---------------------------------------------------------------- control de roturación
F = list(openpyxl.load_workbook(CTRL, read_only=True, data_only=True)["B.D. ROTURACIÓN"].iter_rows(values_only=True))
h = next(i for i, r in enumerate(F) if r and str(r[0]).strip() == "SECTOR-SUERTE")
enc = [re.sub(r"\s+", " ", str(c)).strip() if c is not None else f"c{j}" for j, c in enumerate(F[h])]
C = pd.DataFrame([r[: len(enc)] for r in F[h + 1:]], columns=enc).loc[:, lambda d: ~d.columns.duplicated()]
C = C[C["SECTOR-SUERTE"].notna()].copy()
C["ste"] = C["SECTOR-SUERTE"].astype(str).str.strip()
C["corte_f"] = pd.to_datetime(C["FECHA DE CORTE"], errors="coerce")
C["f1"] = pd.to_datetime(C["FECHA DE ROTURACION PRIMERA"], errors="coerce")
C["pluv"] = C.PLUVIOMETRO.astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
C.loc[~C.pluv.str.fullmatch(r"\d+"), "pluv"] = ""
C["suelo1"] = C["CONDICION SUELO PRIMERA"].astype(str).str.strip().str.upper()
pl_ste = C[C.pluv != ""].groupby("ste").pluv.agg(lambda s: s.mode().iloc[0])
pl_sec = C[C.pluv != ""].assign(sec=C.ste.str[:4]).groupby("sec").pluv.agg(lambda s: s.mode().iloc[0])

# ---------------------------------------------------------------- SimpleAgri: roturación y fertilización confirmadas
GR = {"rot": {"540", "541", "542", "543", "544"}, "fert": {"580", "581", "582", "583", "591", "593"}}
sa = []
for f in sorted(p for p in Path(SA).glob("*.pkl") if p.stem >= "2024-01"):
    for r in pickle.load(open(f, "rb")):
        a = str(r.get("Actividad") or "").strip()
        g = next((k for k, v in GR.items() if a in v), None)
        if g and str(r.get("Estado Planilla", "")).startswith("CONFIRMADO"):
            sa.append((str(r.get("Terreno") or "").strip(), g, r.get("Fecha Confirmacion")))
S = pd.DataFrame(sa, columns=["ste", "g", "f"])
S["f"] = pd.to_datetime(S.f, errors="coerce").dt.normalize()
S = S.dropna().drop_duplicates().sort_values("f")
SG = {k: g.f.values for k, g in S.groupby(["ste", "g"])}


def primera_sa(ste, g, d0, hasta):
    v = SG.get((ste, g))
    if v is None or pd.isna(d0):
        return pd.NaT
    v = v[(v >= np.datetime64(d0 - pd.Timedelta(days=15))) & (v <= np.datetime64(d0 + pd.Timedelta(days=hasta)))]
    return pd.Timestamp(v[0]) if len(v) else pd.NaT


# ---------------------------------------------------------------- base de maduración
d = pd.read_csv(BASE, sep=";", encoding="latin-1", dtype=str, keep_default_na=False)
d["ste"] = d.suerte.str.strip()
d["mes"] = pd.to_numeric(d.MES, errors="coerce")
d["edad"] = pd.to_numeric(d.Edad.str.replace(",", "."), errors="coerce")
d["ini"] = pd.to_datetime(dict(year=2026, month=d.mes, day=15)) - pd.to_timedelta(d.edad * 30.44, unit="D")


def en_control(r):
    g = C[(C.ste == r.ste) & C.corte_f.notna()]
    if not len(g) or pd.isna(r.ini):
        return None
    dd = (g.corte_f - r.ini).dt.days.abs()
    return dd.idxmin() if dd.min() <= 75 else None


d["k"] = [en_control(r) for r in d.itertuples()]
d["corte"] = [C.at[k, "corte_f"] if pd.notna(k) else r for k, r in zip(d.k, d.ini)]
d["f1_ctrl"] = [C.at[k, "f1"] if pd.notna(k) else pd.NaT for k in d.k]
d["suelo1"] = [C.at[k, "suelo1"] if pd.notna(k) else "" for k in d.k]
C["op1"] = C["OPORTUNIDAD 1RA LABOR"].astype(str).str.strip().str.upper()
C["dias1"] = pd.to_numeric(C["DIAS ENTRE PROGRAMACION Y ROTURACION"], errors="coerce")
C["est1"] = C["ESTADO PRIMERA LABORE"].astype(str).str.strip().str.upper()
d["op1"] = [C.at[k, "op1"] if pd.notna(k) else "" for k in d.k]
d["dias1"] = [C.at[k, "dias1"] if pd.notna(k) else np.nan for k in d.k]
d["est1"] = [C.at[k, "est1"] if pd.notna(k) else "" for k in d.k]
d["oport_rot"] = np.select([d.op1 == "TARDIO", d.op1.isin(["IDEAL", "NORMAL"]), (d.est1 == "REALIZADO") & (d.dias1 <= 0)],
                           ["TARDIO", "IDEAL-NORMAL", "IDEAL-NORMAL"], "")
d["oport_de"] = np.where(d.op1.isin(["IDEAL", "NORMAL", "TARDIO"]), "control: " + d.op1,
                         np.where(d.oport_rot != "", "control: realizada con 0 días", ""))
ya = d[(d["Oport. Rot"].str.strip() != "") & (d.oport_rot != "")]
print(f"Oport. Rot: la regla del control reproduce {100 * (ya.oport_rot == ya['Oport. Rot'].str.strip()).mean():.0f} % de {len(ya)} filas ya calificadas")
d["pluv"] = d.ste.map(pl_ste).fillna(d.ste.str[:4].map(pl_sec)).fillna("")
d["rot_sa"] = [primera_sa(s, "rot", c0, 150) for s, c0 in zip(d.ste, d.corte)]
d["fert_sa"] = [primera_sa(s, "fert", c0, 240) for s, c0 in zip(d.ste, d.corte)]
lag = (d.rot_sa - d.f1_ctrl).dt.days
print(f"\nfilas: {len(d)} · con su corte en el control: {d.k.notna().sum()} · con pluviómetro: {(d.pluv != '').sum()} · "
      f"roturación en SimpleAgri: {d.rot_sa.notna().sum()} · fertilización en SimpleAgri: {d.fert_sa.notna().sum()}")
print(f"desfase SimpleAgri − control en la roturación: mediana {lag.median():.0f} días (n={lag.notna().sum()})")
d["f_rot"] = d.f1_ctrl.fillna(d.rot_sa - pd.Timedelta(days=10))
d["fuente_rot"] = np.where(d.f1_ctrl.notna(), "control", np.where(d.rot_sa.notna(), "SimpleAgri", ""))

d["pluv_de"] = np.where(d.ste.isin(pl_ste.index), "suerte", np.where(d.pluv != "", "sector", ""))

# ---------------------------------------------------------------- regla
VENT, UMBRAL, DESF = 5, 1.0, 10  # ≥1 mm desde 5 días antes hasta el día de la labor; SimpleAgri confirma ~10 días después
d["f_fert"] = d.fert_sa - pd.Timedelta(days=DESF)
d["mm_rot"] = [lluvia(p, f, VENT) for p, f in zip(d.pluv, d.f_rot)]
d["mm_fert"] = [lluvia(p, f, VENT) for p, f in zip(d.pluv, d.f_fert)]


def calificar(mm, umbral=UMBRAL):
    mm = np.asarray(mm, dtype=float)
    return np.where(np.isnan(mm), "", np.where(mm >= umbral, "HUMEDO", "FRIABLE"))


d["cal_rot"], d["cal_fert"] = calificar(d.mm_rot), calificar(d.mm_fert)

# ---------------------------------------------------------------- validación contra lo ya calificado (ene-ago)
for col, est in (("Calidad Rot.", "cal_rot"), ("Calidad Fert", "cal_fert")):
    ya = d[(d[col].str.strip() != "") & (d[est] != "")]
    print(f"\n{col}: la regla reproduce {100 * (ya[est] == ya[col].str.strip()).mean():.0f} % de {len(ya)} filas ya calificadas")
    print(pd.crosstab(ya[col].str.strip(), ya[est], rownames=["base"], colnames=["lluvia"]).to_string())
ctl = d[(d["Calidad Rot."].str.strip() != "") & d.suelo1.isin(["FRIABLE", "HUMEDO"])]
print(f"(la condición del suelo anotada en el control coincide solo {100 * (ctl.suelo1 == ctl['Calidad Rot.'].str.strip()).mean():.0f} %)")
ya = d[(d["Calidad Rot."].str.strip() != "") & d.f_rot.notna()]
sens = {}
for w in (3, 5, 7, 8):
    mm = [lluvia(p, f, w) for p, f in zip(ya.pluv, ya.f_rot)]
    sens[w] = {u: round(100 * (calificar(mm, u) == ya["Calidad Rot."].str.strip().values).mean()) for u in (1, 5, 10)}
print("sensibilidad Calidad Rot. (% que coincide; filas = días antes, columnas = mm):")
print(pd.DataFrame(sens).T.to_string())

# ---------------------------------------------------------------- llenar las celdas vacías
o = pd.read_csv(BASE, sep=";", encoding="latin-1", dtype=str, keep_default_na=False)
llen = {}
for col, est in (("Oport. Rot", "oport_rot"), ("Calidad Rot.", "cal_rot"), ("Calidad Fert", "cal_fert")):
    m = (o[col].str.strip() == "") & (d[est] != "")
    o.loc[m, col] = d.loc[m, est]
    llen[col] = m
    print(f"\n{col}: {m.sum()} celdas llenadas {d.loc[m, est].value_counts().to_dict()} · siguen vacías {(o[col].str.strip() == '').sum()}")
o.to_csv(SAL, sep=";", encoding="latin-1", index=False)

sop = d[llen["Oport. Rot"] | llen["Calidad Rot."] | llen["Calidad Fert"]
        | (o[["Oport. Rot", "Calidad Rot.", "Calidad Fert"]].apply(lambda s: s.str.strip()) == "").any(axis=1)]
sop = pd.DataFrame({
    "Suerte": sop.ste, "Mes cosecha": sop.mes, "Corte": sop.Corte, "Fecha corte anterior": sop.corte.dt.date,
    "Oport. Rot": o.loc[sop.index, "Oport. Rot"], "Oport. Rot sale de": sop.oport_de,
    "Días programación → 1.ª roturación (control)": sop.dias1,
    "Pluviómetro": sop.pluv, "Pluviómetro tomado de": sop.pluv_de,
    "Fecha 1.ª roturación": sop.f_rot.dt.date, "Fuente fecha roturación": sop.fuente_rot.replace("SimpleAgri", "SimpleAgri − 10 días"),
    f"Lluvia {VENT} días antes y el día (mm) rot.": sop.mm_rot, "Calidad Rot.": o.loc[sop.index, "Calidad Rot."],
    "Fecha 1.ª fertilización (SimpleAgri − 10 días)": sop.f_fert.dt.date,
    f"Lluvia {VENT} días antes y el día (mm) fert.": sop.mm_fert, "Calidad Fert": o.loc[sop.index, "Calidad Fert"]})
xl = Path(SAL).with_name(Path(SAL).stem + " - soporte lluvia.xlsx")
with pd.ExcelWriter(xl, engine="openpyxl") as w:
    sop.to_excel(w, sheet_name="Septiembre", index=False)
    pd.DataFrame({"Regla": [
        "Oport. Rot: control de roturación, 1.ª labor. TARDIO → TARDIO; IDEAL o NORMAL → IDEAL-NORMAL; realizada con 0 días entre "
        "programación y roturación (el control la deja en blanco) → IDEAL-NORMAL.",
        f"Calidad: HUMEDO si en el pluviómetro de la suerte llovió ≥{UMBRAL:.0f} mm desde {VENT} días antes hasta el día de la labor; si no, FRIABLE.",
        "Pluviómetro de la suerte: columna PLUVIOMETRO del control de roturación (si la suerte no está, el del sector).",
        "Lluvia diaria: Ponderado de Precipitaciones año 2026 (2026 y jun-sep 2025); oct-dic 2025 del registro diario de los mismos pluviómetros.",
        "Fecha de roturación: control oficial (1.ª labor); si no está, SimpleAgri confirmada − 10 días (desfase medido contra el control).",
        "Fecha de fertilización: 1.ª fertilización confirmada en SimpleAgri − 10 días.",
        "Validación ene-ago: ver la salida del programa calidad_suelo_lluvia.py."]}).to_excel(w, sheet_name="Regla", index=False)
    for ws in w.book.worksheets:
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
print(f"\nguardado: {SAL}\nsoporte: {xl}")
