# -*- coding: utf-8 -*-
r"""Riopaila: ¿a qué edad del cultivo se riega? Frente a la ventana crítica de Cenicaña (0-4,5 meses formación y
macollamiento; 4,5-8,5 gran crecimiento, donde el estrés hídrico hace daño irreversible; > 8,5 maduración, estrés moderado
deseable). Prueba aparte: no cambia el modelo ni el informe.

Entradas:
  · «Balance Hídrico 2026 actualizado.xlsx», hoja «Consolidado de riegos»: cada riego con fecha, área regada, EDAD (meses,
    la del registro), estado de la suerte y lámina (MM), abr-2020 a jun-2026 (la misma hoja de balance_riego.py).
  · Clima mensual de Cenicaña (hoja Mensual): falta de agua del mes = máx(0,8 × evaporación − lluvia, 0) de la estación.
  · estaciones_suertes.csv del repo Precipitacion-riopaila y rio_integrado_dataset.pkl (cosechas con su inicio de ciclo).
1) Riego aplicado por etapa: hectáreas regadas y volumen (mm × área regada) por la EDAD del registro, por año.
2) Cosechas con el ciclo completo dentro del registro (inicio ≥ may-2020, cosecha ≤ may-2026): falta de agua y riego de cada
   etapa (mes del ciclo = días desde el inicio / 30,44); qué parte de la falta cubrió el riego en cada etapa.
3) TCH: regresión del log TCH como el modelo de la caída (suerte, año, mes, edad, etapa de la soca, variedad, incendio,
   corte, cosecha mecanizada, calor), ponderada por área, errores por suerte, con la falta sin cubrir y el riego de las
   tres etapas juntas.
4) Sacarosa % caña («Sacarosa C» del ZQM y del G&P 2026) y azúcar por ha (TCH × sacarosa) con la misma regresión; la
   maduración también partida en «8,5 meses a 2 antes del corte» y «2 últimos meses» (el agostamiento). Sin madurante en
   los controles (puede ir junto con la suspensión del riego).

5) Aporte a la caída de lo cosechado ene-sep 2026: efecto × (falta sin cubrir 2026 − la de 2021-25 en los mismos meses).
   Respaldo de la causa calor y sequía: no entra a la cascada (contaría el agua dos veces).

Uso: python riego_por_edad.py <datos intermedios> <Clima_Cenicana_...xlsx> <CONSOLIDADO DE RIEGOS 2026-.xlsx o Balance Hídrico>
                             <repo Precipitacion-riopaila> <salida.xlsx> [<carpeta web>]
Con <carpeta web> escribe data/riego_etapas.json (causa calor y sequía y pestaña «Para la gerencia»).
"""
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.formula.api as smf

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
pd.set_option("display.width", 220)
PK, CLIMA, BH, REPO, SAL = (Path(a) for a in sys.argv[1:6])
ETC = 0.8
CORTES, ETAPAS = [-1, 4.5, 8.5, 99], ["0-4,5 formación", "4,5-8,5 gran crecimiento", "> 8,5 maduración"]
wm = lambda g, k: float(np.average(g[k], weights=g.Area)) if len(g) else np.nan  # noqa: E731

# ---------------------------------------------------------------- riegos
# la hoja se llama «Consolidado de riegos» (Balance Hídrico) o «Consolidado de riegos » (CONSOLIDADO DE RIEGOS 2026-, con
# el encabezado en la fila 6; ese libro trae además una «Consolidado de riegos» corta, del año actual): la más larga
import openpyxl  # noqa: E402

_wb = openpyxl.load_workbook(BH, read_only=True)
hoja = max((h for h in _wb.sheetnames if h.strip() == "Consolidado de riegos"), key=lambda h: _wb[h].max_row)
_wb.close()
cab = pd.read_excel(BH, sheet_name=hoja, header=None, nrows=15, usecols="U")
fila = next(i for i, v in enumerate(cab.iloc[:, 0]) if str(v).strip().upper() == "SEC-SUERTE")
r = pd.read_excel(BH, sheet_name=hoja, header=fila, usecols="A,H,J,N,U,V,AM")
r.columns = ["fecha", "area_regada", "edad", "estado", "ste", "area_neta", "mm"]
r["fecha"] = pd.to_datetime(r.fecha, errors="coerce")
r["ste"] = r.ste.astype(str).str.strip()
for k in ["area_regada", "edad", "area_neta", "mm"]:
    r[k] = pd.to_numeric(r[k], errors="coerce")
r = r.dropna(subset=["fecha", "mm", "area_regada"])
r = r[r.mm.between(1, 400) & (r.area_regada > 0)]
r["vol"] = r.mm * r.area_regada / 1000                       # miles de mm·ha (= 10 000 m³)
r["mm_ste"] = r.mm * r.area_regada / r.area_neta
print(f"riegos sin área neta de la suerte (no se pueden llevar a mm de la suerte): {int((~np.isfinite(r.mm_ste)).sum())}")
r = r[np.isfinite(r.mm_ste)]
r["etapa"] = pd.cut(r.edad, CORTES, labels=ETAPAS)
r["anio"] = r.fecha.dt.year
print(f"riegos: {len(r)} de {r.fecha.min().date()} a {r.fecha.max().date()} · sin edad: {int(r.edad.isna().sum())}")

t1 = r.dropna(subset=["etapa"]).pivot_table(index="anio", columns="etapa", values="area_regada", aggfunc="sum", observed=False)
t1 = t1.div(t1.sum(axis=1), axis=0).mul(100)
t1["ha regadas"] = r.groupby("anio").area_regada.sum()
t1v = r.dropna(subset=["etapa"]).pivot_table(index="anio", columns="etapa", values="vol", aggfunc="sum", observed=False)
t1v = t1v.div(t1v.sum(axis=1), axis=0).mul(100)
t1mm = r.dropna(subset=["etapa"]).groupby(["anio", "etapa"], observed=False).apply(
    lambda g: np.average(g.mm, weights=g.area_regada) if len(g) else np.nan).unstack()
print("\n1) % de las hectáreas regadas por etapa (EDAD del registro):")
print(t1.round(1).to_string())
print("   % del volumen:")
print(t1v.round(1).to_string())
print("   lámina media por riego (mm):")
print(t1mm.round(0).to_string())
t_mes = r[r.edad.between(0, 18)].assign(m=lambda x: x.edad.astype(int)).groupby(["m"]).area_regada.sum()
t_mes = (100 * t_mes / t_mes.sum()).rename("% ha regadas").to_frame()
t_est = r.groupby("estado").area_regada.sum().sort_values(ascending=False)
t_est = (100 * t_est / t_est.sum()).rename("% ha regadas").to_frame()
print("   por estado de la suerte (registro):")
print(t_est.round(1).head(12).to_string())

# ---------------------------------------------------------------- falta de agua y riego de cada etapa, por cosecha
m = pd.read_excel(CLIMA, sheet_name="Mensual")
m = m[m["Estación"].isin(["La Paila", "Bugalagrande"])][["Estación", "Fecha", "Precipitación (mm)", "Evaporación (mm)"]]
m.columns = ["est", "f", "p", "e"]
m["falta"] = (ETC * m.e - m.p).clip(lower=0)
B = {k: g.set_index(pd.to_datetime(g.f).dt.to_period("M")).falta for k, g in m.groupby("est")}
e = pd.read_csv(REPO / "estaciones_suertes.csv", sep=";", skiprows=1, encoding="latin-1").iloc[:, [1, 2]].dropna()
EST = dict(zip(e.iloc[:, 0].astype(str).str.strip(), e.iloc[:, 1].astype(str).str.strip()))
RG = {k: g[["fecha", "mm_ste"]] for k, g in r.groupby("ste")}

d = pd.read_pickle(PK / "rio_integrado_dataset.pkl")
d = d[(d.ini >= "2020-05-01") & (d.cosecha <= r.fecha.max() - pd.Timedelta(days=15)) & d.d32.notna()].copy()
d["estB"] = d.ste.map(EST).fillna(d.est)
FK, RK = ["f_ini", "f_gc", "f_mad"], ["r_ini", "r_gc", "r_mad"]
# la maduración partida: de 8,5 meses a 2 meses antes del corte, y los 2 últimos meses (el agostamiento)
FK2, RK2 = ["f_mad1", "f_ult"], ["r_mad1", "r_ult"]


def etapas(x):
    g = B.get(x.estB)
    out = dict.fromkeys(FK + RK + FK2 + RK2, np.nan)
    if g is None:
        return pd.Series(out)
    # falta por mes calendario, repartida a la etapa por la edad a mitad de mes
    meses = pd.period_range(x.ini.to_period("M"), x.cosecha.to_period("M"), freq="M")
    ed = np.array([((p.to_timestamp() + pd.Timedelta(days=14)) - x.ini).days / 30.44 for p in meses])
    fv = g.reindex(meses).values
    if np.isnan(fv).sum() > 1:
        return pd.Series(out)
    fv = np.nan_to_num(fv)
    rs = RG.get(x.ste)
    if rs is not None:
        rr = rs[(rs.fecha >= x.ini) & (rs.fecha < x.cosecha)]
        er = ((rr.fecha - x.ini).dt.days / 30.44).values
        mmr = rr.mm_ste.values
    else:
        er, mmr = np.array([]), np.array([])
    for (a, b), fk, rk in zip(((-1, 4.5), (4.5, 8.5), (8.5, 99)), FK, RK):
        out[fk] = fv[(ed > a) & (ed <= b)].sum()
        out[rk] = mmr[(er > a) & (er <= b)].sum()
    u = max((x.cosecha - x.ini).days / 30.44 - 2, 8.5)
    for (a, b), fk, rk in zip(((8.5, u), (u, 99)), FK2, RK2):
        out[fk] = fv[(ed > a) & (ed <= b)].sum()
        out[rk] = mmr[(er > a) & (er <= b)].sum()
    return pd.Series(out)


d = pd.concat([d, d.apply(etapas, axis=1)], axis=1).dropna(subset=FK)
print(f"\n2) cosechas con el ciclo completo en el registro: {len(d)} ({d.Area.sum():,.0f} ha), cosechadas "
      f"{d.cosecha.min():%b-%Y} a {d.cosecha.max():%b-%Y}")
filas = []
for et, fk, rk in zip(ETAPAS, FK, RK):
    filas.append({"etapa": et, "falta de agua mm": wm(d, fk), "riego mm": wm(d, rk),
                  "% de la falta total": 100 * wm(d, fk) / sum(wm(d, k) for k in FK),
                  "% del riego total": 100 * wm(d, rk) / sum(wm(d, k) for k in RK),
                  "riego / falta %": 100 * wm(d, rk) / wm(d, fk),
                  "% área con falta > 100 mm": 100 * wm(d.assign(z=(d[fk] > 100).astype(float)), "z"),
                  "% área regada (algún riego)": 100 * wm(d.assign(z=(d[rk] > 0).astype(float)), "z")})
t2 = pd.DataFrame(filas)
print(t2.round(1).to_string(index=False))
por_anio = d.groupby("anio").apply(lambda g: pd.Series({**{f"falta {e_}": wm(g, k) for e_, k in zip(ETAPAS, FK)},
                                                       **{f"riego {e_}": wm(g, k) for e_, k in zip(ETAPAS, RK)}, "ha": g.Area.sum()}))
print(por_anio.round(0).to_string())

# ---------------------------------------------------------------- 3) TCH
for fk, rk in zip(FK, RK):
    d[fk + "_sc"] = (d[fk] - d[rk]).clip(lower=0) / 100
    d[rk + "_100"] = d[rk] / 100
BASE = "y ~ C(ste) + C(anio) + C(Mes) + C(edad_b) + C(soca_b) + C(var) + inc + Corte + mec + d32"
ref = d[d.anio.between(2021, 2025)]
TN = float(np.average(ref.tch_calc, weights=ref.Area))
res = []
for nom, xs in (("falta de agua (sin restar el riego)", [k + "_f" for k in FK]),
                ("falta sin cubrir y riego", [k + "_sc" for k in FK] + [k + "_100" for k in RK])):
    for fk in FK:
        d[fk + "_f"] = d[fk] / 100
    mo = smf.wls(BASE + " + " + " + ".join(xs), data=d, weights=d.Area)
    ft = mo.fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(d.loc[mo.data.row_labels, "ste"])[0]})
    for k in xs:
        et = ETAPAS[(FK + RK).index(k.rsplit("_", 1)[0]) % 3]
        pct = 100 * (np.exp(ft.params[k]) - 1)
        res.append({"modelo": nom, "variable": ("falta sin cubrir" if k.endswith("_sc") else "riego" if k.endswith("_100")
                                                else "falta de agua"), "etapa": et, "% TCH por 100 mm": pct,
                    "TCH por 100 mm": pct / 100 * TN, "p": ft.pvalues[k], "n": int(ft.nobs)})
t3 = pd.DataFrame(res)
print(f"\n3) TCH (normal 2021-25 {TN:.1f}):")
print(t3.round(3).to_string(index=False))

# ---------------------------------------------------------------- 4) sacarosa y azúcar por ha (lo que busca el agostamiento)
# sacarosa % caña («Sacarosa C» del ZQM) de la misma fila que usa el modelo (la de mayor TC-Total), como madurante_analisis.py
sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
import openpyxl  # noqa: E402
from tchmotor import dataset_gee as dg  # noqa: E402

z = dg._hoja(dg.ZQM_HIST, 0)
z.columns = [str(k).strip() for k in z.columns]
w = openpyxl.load_workbook(r"C:\Users\Agr349\Documents\G&P Suertes_26.xlsx", read_only=True, data_only=True)
it = w["ZQM101"].iter_rows(values_only=True)
for fila in it:
    if any(str(v).strip() == "Ubic Técnica" for v in fila if v is not None):
        hdr = dg._dedup([str(v).strip() if v is not None else f"c{j}" for j, v in enumerate(fila)])
        break
g26 = pd.DataFrame([f_ for f_ in it if any(v is not None for v in f_[:6])], columns=hdr)
w.close()
g26["AÑO"] = 2026
COL = ["AÑO", "Mes", "Ubic Técnica", "Corte", "TC-Total", "Sacarosa C"]
s = pd.concat([z[COL], g26[COL]], ignore_index=True)
for k in ["AÑO", "Mes", "Corte", "TC-Total", "Sacarosa C"]:
    s[k] = pd.to_numeric(s[k], errors="coerce")
s["ste"] = s["Ubic Técnica"].astype(str).str.strip()
s = s.dropna(subset=["TC-Total", "Sacarosa C"]).sort_values("TC-Total").drop_duplicates(["ste", "AÑO", "Mes", "Corte"], keep="last")
s = s.rename(columns={"AÑO": "anio", "Sacarosa C": "sac"})[["ste", "anio", "Mes", "Corte", "sac"]]
d = d.merge(s, on=["ste", "anio", "Mes", "Corte"], how="left").set_axis(d.index)
d.loc[~d.sac.between(7, 17), "sac"] = np.nan
d["y_az"] = d.y + np.log(d.sac / 100)                       # log de toneladas de azúcar (sacarosa) por ha
ds = d.dropna(subset=["sac"])
SAC = float(np.average(ds[ds.anio.between(2021, 2025)].sac, weights=ds[ds.anio.between(2021, 2025)].Area))
print(f"\n4) sacarosa en {len(ds)} de {len(d)} cosechas (normal 2021-25 {SAC:.2f} % caña)")
for k in FK2:
    d[k + "_sc"] = (d[k] - d["r" + k[1:]]).clip(lower=0) / 100
    d["r" + k[1:] + "_100"] = d["r" + k[1:]] / 100
ds = d.dropna(subset=["sac"])
NOMBRE = {"f_ini": ETAPAS[0], "f_gc": ETAPAS[1], "f_mad": ETAPAS[2], "f_mad1": "8,5 meses a 2 antes del corte",
          "f_ult": "2 últimos meses (agostamiento)"}
res4 = []
for parte, fks in (("tres etapas", FK), ("maduración partida", ["f_ini", "f_gc"] + FK2)):
    xs = [k + "_sc" for k in fks] + ["r" + k[1:] + "_100" for k in fks]
    for var, nom, esc in (("y", "TCH %", None), ("sac", "sacarosa (puntos % caña)", 1), ("y_az", "azúcar por ha %", None)):
        mo = smf.wls(BASE.replace("y ~", f"{var} ~") + " + " + " + ".join(xs), data=ds, weights=ds.Area)
        ft = mo.fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(ds.loc[mo.data.row_labels, "ste"])[0]})
        for fk in fks:
            k = fk + "_sc"
            v = ft.params[k] if esc else 100 * (np.exp(ft.params[k]) - 1)
            res4.append({"partición": parte, "resultado": nom, "etapa": NOMBRE[fk], "por 100 mm de falta sin cubrir": v,
                         "p": ft.pvalues[k], "n": int(ft.nobs)})
t4 = pd.DataFrame(res4)
print(t4.pivot_table(index=["partición", "etapa"], columns="resultado", values=["por 100 mm de falta sin cubrir", "p"],
                     sort=False).round(3).to_string())
print("falta y riego de los 2 últimos meses (mm, ponderado):", round(wm(d, "f_ult")), round(wm(d, "r_ult")),
      "· de 8,5 a 2 antes:", round(wm(d, "f_mad1")), round(wm(d, "r_mad1")))

# ---------------------------------------------------------------- 5) ¿aporta a la caída de 2026? (lo cosechado ene-sep)
# aporte = efecto × (falta sin cubrir 2026 − la de 2021-25 en los mismos meses), con la maduración partida
FKP = ["f_ini", "f_gc"] + FK2
ap = []
q5 = d[(d.anio == 2026) & (d.Mes <= 9)]
r5 = d[d.anio.between(2021, 2025) & (d.Mes <= 9)]
TN5 = float(np.average(np.exp(q5.y), weights=q5.Area))
for var, nom in (("y", "TCH"), ("y_az", "azúcar por ha")):
    dd = d.dropna(subset=["sac"]) if var == "y_az" else d
    xs = [k + "_sc" for k in FKP] + ["r" + k[1:] + "_100" for k in FKP]
    mo = smf.wls(BASE.replace("y ~", f"{var} ~") + " + " + " + ".join(xs), data=dd, weights=dd.Area)
    ft = mo.fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(dd.loc[mo.data.row_labels, "ste"])[0]})
    for fk in FKP:
        k = fk + "_sc"
        cambio = wm(q5, k) - wm(r5, k)
        ap.append({"resultado": nom, "etapa": NOMBRE[fk], "sin cubrir 2021-25 mm": 100 * wm(r5, k), "sin cubrir 2026 mm": 100 * wm(q5, k),
                   "aporte %": 100 * (np.exp(ft.params[k] * cambio) - 1),
                   "aporte TCH": ft.params[k] * cambio * TN5 if var == "y" else np.nan})
t5 = pd.DataFrame(ap)
print(f"\n5) aporte a 2026 (cosechas ene-sep 2026: {len(q5)}, {q5.Area.sum():,.0f} ha, TCH {TN5:.1f}):")
print(t5.round(2).to_string(index=False))

# ---------------------------------------------------------------- informe: data/riego_etapas.json (si se da la carpeta web)
if len(sys.argv) > 6:
    import json  # noqa: E402

    r1 = lambda v, n=1: None if v is None or pd.isna(v) else round(float(v), n)  # noqa: E731
    ef = t4[t4["partición"] == "maduración partida"].pivot_table(index="etapa", columns="resultado",
                                                                  values=["por 100 mm de falta sin cubrir", "p"], sort=False)
    v_ = "por 100 mm de falta sin cubrir"
    out = {
        "riegos": int(len(r)), "desde": r.fecha.min().date().isoformat(), "hasta": r.fecha.max().date().isoformat(),
        "cosechas": int(len(d)), "ha": r1(d.Area.sum(), 0), "cos_desde": f"{d.cosecha.min():%Y-%m}", "cos_hasta": f"{d.cosecha.max():%Y-%m}",
        "etapas": [[x["etapa"], r1(x["falta de agua mm"], 0), r1(x["riego mm"], 0), r1(x["% de la falta total"], 0),
                    r1(x["% del riego total"], 0), r1(x["riego / falta %"], 0)] for _, x in t2.iterrows()],
        "ult": [r1(wm(d, "f_ult"), 0), r1(wm(d, "r_ult"), 0)],
        "efectos": [[e_, r1(ef.loc[e_, (v_, "TCH %")], 1), r1(ef.loc[e_, ("p", "TCH %")], 3),
                     r1(ef.loc[e_, (v_, "sacarosa (puntos % caña)")], 2), r1(ef.loc[e_, ("p", "sacarosa (puntos % caña)")], 3),
                     r1(ef.loc[e_, (v_, "azúcar por ha %")], 1), r1(ef.loc[e_, ("p", "azúcar por ha %")], 3)] for e_ in ef.index],
        "aporte_n": [int(len(q5)), r1(q5.Area.sum(), 0), r1(TN5)],
        "aporte": [[e_, r1(g.iloc[0]["sin cubrir 2021-25 mm"], 0), r1(g.iloc[0]["sin cubrir 2026 mm"], 0),
                    r1(g[g.resultado == "TCH"]["aporte TCH"].iloc[0], 2), r1(g[g.resultado == "azúcar por ha"]["aporte %"].iloc[0], 2),
                    r1(g[g.resultado == "TCH"]["aporte %"].iloc[0], 2)]
                   for e_, g in t5.groupby("etapa", sort=False)],
        "anios": [[int(a)] + [r1(x[f"riego {e_}"], 0) for e_ in ETAPAS] for a, x in por_anio.iterrows()],
    }
    dest = Path(sys.argv[6]) / "data" / "riego_etapas.json"
    json.dump(out, open(dest, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    print("escrito:", dest)

with pd.ExcelWriter(SAL, engine="openpyxl") as wr:
    pd.DataFrame({"nota": [
        "Riego por edad del cultivo frente a la ventana crítica de Cenicaña (0-4,5 / 4,5-8,5 / > 8,5 meses).",
        "Fuente del riego: Balance Hídrico 2026 actualizado, hoja Consolidado de riegos (EDAD del registro).",
        "Falta de agua = máx(0,8 × evaporación − lluvia, 0) mensual de la estación de la suerte (Cenicaña).",
        "Hojas 2-3: cosechas con todo el ciclo dentro del registro de riego (inicio ≥ may-2020, cosecha ≤ may-2026).",
        "Desde 2024 cada riego aparece con menos mm (posible cambio de registro): mirar también las hectáreas regadas.",
        "No cambia el modelo, el informe ni el soporte."]}).to_excel(wr, sheet_name="Nota", index=False)
    t1.round(2).to_excel(wr, sheet_name="Riego por etapa (% ha)")
    t1v.round(2).to_excel(wr, sheet_name="Riego por etapa (% vol)")
    t1mm.round(1).to_excel(wr, sheet_name="Lámina por riego")
    t_mes.round(2).to_excel(wr, sheet_name="Riego por mes de edad")
    t_est.round(2).to_excel(wr, sheet_name="Riego por estado")
    t2.to_excel(wr, sheet_name="Falta y riego por etapa", index=False)
    por_anio.round(1).to_excel(wr, sheet_name="Falta y riego por año")
    t3.to_excel(wr, sheet_name="TCH por etapa", index=False)
    t4.to_excel(wr, sheet_name="TCH, sacarosa y azúcar", index=False)
    t5.to_excel(wr, sheet_name="Aporte 2026 (ene-sep)", index=False)
print("guardado:", SAL)
