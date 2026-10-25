# -*- coding: utf-8 -*-
"""Matriz ponderada Riopaila 2026: cuántas t/ha le quita (o le da) cada variable a cada suerte, con su efecto MEDIDO.

Paso 1 — variables que mide el modelo integrado (caida_integrada_rio.py, cosechas 2011-2026, cada suerte frente a sí
misma): calor, lluvia del arranque, sequía al final del ciclo, riego, cosecha mecánica, labores, renovación, variedades,
edad, número de corte e incendio. Se usan sus aportes suerte por suerte (caida_por_suerte.pkl); cuadran con la lámina.
Paso 2 — calificaciones que solo tiene la matriz (roya, carbón, madurante, permanencia, calidad de labores, nivel freático,
CAPEX, ganado, tiempo térmico, lluvia del mes antes de cosecha): se mide su efecto sobre lo que el paso 1 no explicó, en
las suertes de 2026, con la condición de que peor nota = menos TCH (si sale al revés o en cero, pesa 0).
Toneladas: aporte en log × TCH real × (e^L − 1)/L (L = caída en log de la suerte), así los aportes suman exacto la caída
de cada suerte (TCH histórico 2020-25 − TCH real). La calificación de variedad de la matriz no se usa: sale del mismo TCH.

Uso: python matriz_ponderada.py <caida_por_suerte.pkl> <rio matriz.xlsx> <hoja> <Caida_TCH_..._soporte.xlsx> <salida.xlsx>
"""
import sys
import warnings

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import lsq_linear

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
PKL, MAT, HOJA, SOP, SAL = sys.argv[1:6]

x = pd.read_pickle(PKL)
m = pd.read_excel(MAT, sheet_name=HOJA)
m = m[m.suerte.notna()].reset_index(drop=True)
m["ste"], m["Mes"], m["Cor"] = m.suerte.astype(str).str.strip(), pd.to_numeric(m.MES), pd.to_numeric(m.Cor, errors="coerce")
INFO = ["ste", "Mes", "Corte", "Area", "TC-Total", "rel", "var", "mec", "d32", "rot", "no_renov"]
x = x.rename(columns={c: "A·" + c for c in x.columns if c not in INFO}).rename(columns={"Corte": "Cor"})  # aportes con prefijo: la matriz tiene «Incendio»
# cruce suerte + mes (+ corte si la suerte se cosechó dos veces el mismo mes)
x["n"] = x.groupby(["ste", "Mes"]).cumcount()
m["n"] = m.groupby(["ste", "Mes"]).cumcount()
d = m.merge(x.drop(columns=["Area"]), on=["ste", "Mes", "n"], how="left", suffixes=("", "_x"))
print(f"matriz: {len(m)} filas · cruzan con el modelo: {d['var'].notna().sum()} ({d.rel.notna().sum()} con historia) · no cruzan: "
      f"{d.loc[d['var'].isna(), ['ste', 'Mes', 'Ing']].values.tolist()}")

# ---------------------------------------------------------------- paso 1: nombres como en la matriz
PASO1 = {  # columna del modelo → (nombre en la matriz ponderada, grupo, calificación de la matriz que reemplaza)
    "Calor (días >32 °C)": ("Días Tmax >32 °C", "No controlable", "Calificacion Class dias T_Max mas 32"),
    "Lluvia del arranque (0-3 meses)": ("Precipitación 0 a 3 meses", "No controlable", "Calificacion Precip 0 a 3"),
    "Sequía al final del ciclo": ("Balance hídrico: sequía al final", "No controlable", "Calificacion Balance Hidrico"),
    "Riego aplicado": ("Balance hídrico: riego aplicado", "Flujo de caja", "Calificacion Balance Hidrico"),
    "Renovación no hecha": ("No renovación", "Flujo de caja", "Calificación No Renovación"),
    "Roturación tardía": ("Oportunidad roturación", "Flujo de caja", "Calificación Oportunidad Roturación"),
    "Fertilización después de 120 días": ("Oportunidad fertilización", "Flujo de caja", "Calificación Oportunidad Fertilización"),
    "Malezas (herbicida después de 90 días)": ("Control de malezas (nueva)", "Flujo de caja", ""),
    "Cosecha mecánica": ("Cosecha mecánica (nueva)", "Agronómico y operativo", ""),
    "Variedades": ("Variedad", "Agronómico y operativo", "Calificación Variedad"),
    "Edad de cosecha": ("Edad de cosecha", "Agronómico y operativo", "Calificación Edad Cosecha"),
    "Número de corte": ("N° de corte", "Agronómico y operativo", "calificación N° Corte"),
    "Incendio": ("Incendio (nueva)", "Agronómico y operativo", ""),
}
# ---------------------------------------------------------------- paso 2: calificaciones que solo tiene la matriz
PASO2 = {
    "Calificación Roya": ("Roya", "Agronómico y operativo"),
    "Calificación Carbón": ("Carbón", "Agronómico y operativo"),
    "Calificación Dosis Maduración": ("Dosis madurante", "Agronómico y operativo"),
    "Calificación semanas": ("Semanas madurante", "Agronómico y operativo"),
    "Calificación TP": ("Tiempo de permanencia", "Agronómico y operativo"),
    "Calificación Calidad Roturación": ("Calidad roturación", "No controlable"),
    "Calificación Calidad fertilización": ("Calidad fertilización", "No controlable"),
    "Calificación Nivel freatico": ("Nivel freático", "No controlable"),
    "Calificacion CAPEX": ("CAPEX", "Flujo de caja"),
    "Calificación daño Ganado externo": ("Daño de ganado", "Agronómico y operativo"),
    "Calificacion Clas T Termico": ("Tiempo térmico", "No controlable"),
    "Calificacion Clas 1 mes antes de cosecha": ("Lluvia 1 mes antes de cosecha", "No controlable"),
}
col = {c.strip().lower()[:30]: c for c in m.columns}
P2 = {}
for k, v in PASO2.items():
    real = next((c for c in m.columns if c.strip().lower()[:30] == k.strip().lower()[:30]), None)
    if real is None:
        real = next((c for c in m.columns if c.strip().lower().startswith(k.strip().lower()[:24])), None)
    if real is not None and pd.to_numeric(d[real], errors="coerce").nunique() > 1:
        P2[real] = v
print("paso 2:", [v[0] for v in P2.values()])

# solo hay notas de 2026: cada suerte se mide frente a la nota promedio de 2026 (no frente a la nota 1), así estas
# variables reparten «sin explicar» entre suertes pero no se inventan una parte de la caída total frente a la historia
e = d.dropna(subset=["rel"]).copy()
MU = {c: np.average(pd.to_numeric(e[c], errors="coerce").fillna(1), weights=e.Area) for c in P2}
Xn = pd.DataFrame({c: pd.to_numeric(e[c], errors="coerce").fillna(1) - MU[c] for c in P2})
y, w = e["A·Sin explicar"].values, e.Area.values
X = np.column_stack([np.ones(len(e)), Xn.values])
sw = np.sqrt(w / w.mean())
lo = np.r_[-np.inf, np.full(len(P2), -np.inf)]
hi = np.r_[np.inf, np.zeros(len(P2))]  # peor nota solo puede restar
b = lsq_linear(X * sw[:, None], y * sw, bounds=(lo, hi)).x
libre = sm.WLS(y, X, weights=w).fit()
pesos2 = pd.DataFrame({"efecto libre % por punto": 100 * (np.exp(libre.params[1:]) - 1), "p": libre.pvalues[1:],
                       "peso usado % por punto": 100 * (np.exp(b[1:]) - 1),
                       "nota 5 frente a nota 1 (t/ha)": -np.expm1(4 * b[1:]) * np.average(e["TC-Total"] / e.Area, weights=e.Area)},
                      index=list(P2))
print(pesos2.round(3).to_string())

# ---------------------------------------------------------------- aportes por suerte en t/ha
tch = d["TC-Total"] / pd.to_numeric(d.Area)
L = -(d["A·caída"])
fac = np.where(L.abs() > 1e-9, tch * np.expm1(L) / L, tch)
fac = pd.Series(np.where(L.isna(), tch, fac), index=d.index)  # sin historia: primer orden
out = pd.DataFrame(index=d.index)
for k, (nom, _, _) in PASO1.items():
    out[nom] = -d["A·" + k] * fac
for c, (nom, _) in P2.items():
    nota = pd.to_numeric(d[c], errors="coerce").fillna(1) - MU[c]
    out[nom] = -(b[1 + list(P2).index(c)] * nota) * fac
resto = d["A·Sin explicar"] - sum(b[1 + i] * (pd.to_numeric(d[c], errors="coerce").fillna(1) - MU[c]) for i, c in enumerate(P2))
out["Sin explicar"] = -resto * fac
caida = (tch * np.expm1(L)).where(d.rel.notna())
out.loc[d.rel.isna(), "Sin explicar"] = np.nan
# la que más quitó, contando «sin explicar»: si gana esa, la matriz no explica la caída de esa suerte
principal = out.idxmax(axis=1).where(out.max(axis=1) > 0.5, "").replace("Sin explicar", "No lo explica ninguna variable")

area = pd.to_numeric(d.Area)
base = pd.DataFrame({"Suerte": d.suerte, "Jefe de zona": d.get("Jefe de zona"), "Técnico": d.get("Tecnico"), "Zona": d.get("Zona"),
                     "Mes": d.Mes, "Corte": d.Cor, "Variedad": d.get("Variedad"), "Tipo de cosecha": d.get("TIPO COS"),
                     "Área (ha)": area, "TCH real": tch, "TCH histórico 2020-25": tch + caida,
                     "Caída (t/ha)": caida, "Factor que más quitó": principal})
mat = pd.concat([base, out.round(2)], axis=1)
mat.loc[d["var"].isna(), "Factor que más quitó"] = "sin dato en el modelo (otro ingenio o no cruza)"
mat.loc[d.rel.isna() & d["var"].notna(), "Caída (t/ha)"] = np.nan

# ---------------------------------------------------------------- resumen y pesos
ok = d.rel.notna()
tot_t = (caida * area)[ok].sum()
filas = []
for k, (nom, gru, reemplaza) in PASO1.items():
    filas.append({"variable": nom, "grupo": gru, "reemplaza a": reemplaza, "fuente del peso": "modelo 2011-2026 (7.433 cosechas)"})
for c, (nom, gru) in P2.items():
    filas.append({"variable": nom, "grupo": gru, "reemplaza a": c.strip(), "fuente del peso": "matriz 2026 (lo que el modelo no explicó)",
                  "efecto libre % por punto": pesos2.loc[c, "efecto libre % por punto"], "p": pesos2.loc[c, "p"],
                  "peso usado % por punto": pesos2.loc[c, "peso usado % por punto"],
                  "nota 5 frente a nota 1 (t/ha)": pesos2.loc[c, "nota 5 frente a nota 1 (t/ha)"]})
filas.append({"variable": "Sin explicar", "grupo": "Sin explicar", "reemplaza a": "", "fuente del peso": "lo que ninguna variable explica"})
res = pd.DataFrame(filas)
res["toneladas"] = [(out[v] * area)[ok].sum() for v in res.variable]
res["t/ha"] = res.toneladas / area[ok].sum()
res["% de la caída"] = 100 * res.toneladas / tot_t
coef = pd.read_excel(SOP, sheet_name="Coeficientes").set_index("variable")
EF = {"Días Tmax >32 °C": ("d32", "por día >32 °C"), "Precipitación 0 a 3 meses": ("llu03_100", "por 100 mm"),
      "Balance hídrico: sequía al final": ("def100", "por 100 mm de déficit"), "Balance hídrico: riego aplicado": ("def100", "por 100 mm de déficit"),
      "No renovación": ("no_renov", "suerte no renovada"), "Oportunidad roturación": ("rot_tard", "roturación tardía"),
      "Oportunidad fertilización": ("fert_120", "fertilización >120 días"), "Control de malezas (nueva)": ("mal_extra30", "por mes de atraso >90 días"),
      "Cosecha mecánica (nueva)": ("mec_2026", "mecánica frente a manual (2026)"), "N° de corte": ("Corte", "por corte"),
      "Incendio (nueva)": ("inc", "cosecha por incendio")}
res["efecto medido"] = [f"{coef.loc[EF[v][0], 'efecto %']:+.1f} % {EF[v][1]} (p={coef.loc[EF[v][0], 'p']:.2f})" if v in EF else
                        ("por variedad (modelo)" if v == "Variedad" else "por rango de edad (modelo)" if v == "Edad de cosecha" else
                         (f"{r_:+.1f} % por punto de nota, frente a la suerte promedio de 2026" if pd.notna(r_ := res.loc[res.variable == v, 'peso usado % por punto'].iloc[0]) else ""))
                        for v in res.variable]
res = res.sort_values("toneladas", ascending=False)
grp = res.groupby("grupo")[["toneladas", "t/ha", "% de la caída"]].sum().sort_values("toneladas", ascending=False)
EXPL = 1 - out.loc[ok, "Sin explicar"].var() / caida[ok].var()
print(f"suerte por suerte, las variables explican {100 * EXPL:.0f} % de las diferencias de caída entre suertes")
print(f"\ncaída exacta, suertes con historia: {tot_t:,.0f} t en {area[ok].sum():,.0f} ha ({tot_t / area[ok].sum():.1f} t/ha)")
print(res[["variable", "grupo", "toneladas", "t/ha", "% de la caída", "efecto medido"]].round(1).to_string(index=False))
print(grp.round(1).to_string())

LEER = [
    "Qué es: para cada suerte cosechada en 2026, cuántas t/ha le quitó (positivo) o le dio (negativo) cada variable frente a lo que esa misma suerte daba en 2020-25.",
    "Las columnas de variables suman exacto la «Caída (t/ha)» de la suerte; lo que ninguna explica queda en «Sin explicar».",
    "Pesos: cada variable pesa lo que mide su efecto, no 1/3/5 igual para todas. Ver hoja «Pesos».",
    "Variables con dato medido (calor, lluvia, agua, cosecha, labores, renovación, variedad, edad, corte, incendio): efecto del modelo con 7.433 cosechas 2011-2026, cada suerte frente a sí misma.",
    "Calificaciones que solo tiene la matriz (roya, carbón, madurante, permanencia, calidad de labores, nivel freático, CAPEX, ganado, tiempo térmico, lluvia del mes antes): efecto medido en 2026 sobre lo que el modelo no explicó; si una peor nota no resta TCH, pesa 0.",
    "Como esas notas solo existen para 2026, cada suerte se compara con la suerte promedio de 2026: dicen qué suertes castigan más que las otras, pero en total suman cero, porque sin notas de años anteriores no se puede saber cuánto de la caída frente a la historia explican. Su tamaño está en «Pesos»: nota 5 frente a nota 1.",
    "La calificación de variedad de la matriz no se usa porque sale del mismo TCH (real frente al benchmark de la zona); en su lugar va el efecto de cada variedad del modelo.",
    "Se agregan cosecha mecánica, control de malezas e incendio, que no estaban en la matriz.",
    f"Suertes sin historia 2020-25 con ese código ({int((~ok & d['var'].notna()).sum())}): se muestran los aportes, sin caída ni «sin explicar». Filas de otros ingenios o que no cruzan: en blanco.",
    "La suma de esta matriz es la caída suerte a suerte; puede diferir un poco de la lámina, que mide la caída del conjunto.",
    f"Alcance: en el conjunto las variables explican el {100 * (1 - res.loc[res.variable == 'Sin explicar', 'toneladas'].iloc[0] / tot_t):.0f} % de la caída; "
    f"suerte por suerte explican solo el {100 * EXPL:.0f} % de las diferencias entre suertes. Sirve para explicar la caída del ingenio y "
    "las tendencias; en una suerte puntual, «Sin explicar» grande quiere decir que hay que ir a mirarla en campo.",
]
with pd.ExcelWriter(SAL, engine="openpyxl") as wr:
    mat.to_excel(wr, sheet_name="Matriz ponderada", index=False)
    res[["variable", "grupo", "toneladas", "t/ha", "% de la caída", "efecto medido", "fuente del peso", "reemplaza a",
         "nota 5 frente a nota 1 (t/ha)", "efecto libre % por punto", "p"]].to_excel(wr, sheet_name="Pesos", index=False)
    grp.reset_index().to_excel(wr, sheet_name="Por grupo", index=False)
    pd.DataFrame({"Cómo leerla": LEER}).to_excel(wr, sheet_name="Cómo leerla", index=False)

wb = load_workbook(SAL)
ws = wb["Matriz ponderada"]
hdr = PatternFill("solid", fgColor="1E3328")
for c in ws[1]:
    c.font, c.fill = Font(bold=True, color="F6F7F2"), hdr
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 48
for j, c in enumerate(ws[1], 1):
    L_ = c.column_letter
    ws.column_dimensions[L_].width = 26 if c.value in ("Jefe de zona", "Técnico", "Factor que más quitó") else 12
    if j > len(base.columns) - 1 or c.value in ("TCH real", "TCH histórico 2020-25", "Caída (t/ha)", "Área (ha)"):
        for cell in ws[L_][1:]:
            cell.number_format = "0.0"
primera = ws.cell(1, len(base.columns) + 1).column_letter
ultima = ws.cell(1, ws.max_column).column_letter
ws.conditional_formatting.add(f"{primera}2:{ultima}{ws.max_row}",
                              ColorScaleRule(start_type="num", start_value=-5, start_color="4E79A7", mid_type="num", mid_value=0,
                                             mid_color="FBFBF8", end_type="num", end_value=10, end_color="C4661D"))
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
for n_ in ("Pesos", "Por grupo"):
    s = wb[n_]
    for c in s[1]:
        c.font, c.fill = Font(bold=True, color="F6F7F2"), hdr
    for c in s[1]:
        s.column_dimensions[c.column_letter].width = 34 if c.value in ("variable", "efecto medido", "fuente del peso", "reemplaza a", "grupo") else 14
    for row in s.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.0"
wb["Cómo leerla"].column_dimensions["A"].width = 160
wb.save(SAL)
print("guardado:", SAL)
