# -*- coding: utf-8 -*-
"""Riopaila: madurante aplicado a cada cosecha y su efecto en el TCH, con la base de 12 cc/t de trinexapac.

Fuente de aplicaciones (solo lectura): Base_Reporte Aplicaciones Agrícola 2019-2026 (hoja Madurante; 2018 a sep-2026):
suerte, fecha, producto y dosis (L/ha) de cada aplicación. Productos agrupados:
  trinexapac (Bonus, Tronuspac) · glifosato (Glifosato, Roundup) · fusilade · promotor (Madukast/Maducast, Gibcas,
  Nutricast) · alternativo (fosfato de potasio, boro, Rendimiento, Núcleo 12, Suggar Max y otros).
A cada cosecha de rio_integrado_dataset.pkl se le asignan las aplicaciones de su suerte entre 150 días antes y 20 días
después del día 15 de su mes de cosecha (el ZQM da el mes). Dosis de trinexapac en cc/t = L/ha × 1000 / TCH esperado,
con TCH esperado = promedio de los 3 cortes anteriores de la suerte (lo que se sabía al aplicar; con el TCH real la dosis
por tonelada sube justo donde la caña rindió poco y el análisis saldría sesgado). En 2026 se muestra también con el PPTO.
Efecto: el modelo de caida_integrada_rio.py (efectos fijos de suerte, año y mes) + las variables de madurante, en las
cosechas desde jul-2018 (las anteriores van como «sin dato»).

Uso: <tch-motor>\\.venv-gee\\Scripts\\python.exe madurante_analisis.py <datos intermedios> <Base_Reporte Aplicaciones.xlsx>
       <Maestro Suertes 2026.xlsx> <salida.xlsx>
"""
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
pd.set_option("display.width", 250)
pd.set_option("display.max_columns", 30)
PK, APL, MAESTRO, SAL = Path(sys.argv[1]), sys.argv[2], sys.argv[3], sys.argv[4]
wm = lambda g, k: np.average(g[k], weights=g.Area) if len(g) else np.nan  # noqa: E731
pct = lambda b: 100 * np.expm1(b)  # noqa: E731

# ---------------------------------------------------------------- aplicaciones
a = pd.read_excel(APL, sheet_name="Madurante", header=3)
a.columns = [str(c).strip() for c in a.columns]
a = a.rename(columns={"SEC-STE": "ste", "FECHA DE APLICACIÓN": "fecha", "PRODUCTO*": "producto", "DOSIS": "dosis", "Aplicación": "aplic"})
a["ste"] = a.ste.astype(str).str.strip()
a["fecha"] = pd.to_datetime(a.fecha, errors="coerce")
a["dosis"] = pd.to_numeric(a.dosis.astype(str).str.extract(r"([\d.,]+)")[0].str.replace(",", "."), errors="coerce")
p = a.producto.astype(str).str.lower().str.strip()


def grupo(s):
    if re.search(r"bonus|tron+us", s):
        return "trinexapac"
    if re.search(r"glifo|roundup", s):
        return "glifosato"
    if "fusilade" in s:
        return "fusilade"
    if re.search(r"madu|gibca|nutrica", s):
        return "promotor"
    return "alternativo"


a["grupo"] = p.map(grupo)
a = a.dropna(subset=["fecha"])
print(f"{len(a)} aplicaciones {a.fecha.min().date()} a {a.fecha.max().date()} · {a.ste.nunique()} suertes")
print(pd.crosstab(a.fecha.dt.year, a.grupo))
print("dosis por grupo (L/ha):\n", a.groupby("grupo").dosis.describe().round(2))
print("productos más usados por grupo:")
for g_, d in a.groupby("grupo"):
    print(f"  {g_}: {d.producto.astype(str).str.strip().str.capitalize().value_counts().head(5).to_dict()}")

# ---------------------------------------------------------------- cosechas
c = pd.read_pickle(PK / "rio_integrado_dataset.pkl").sort_values(["ste", "cosecha"])
c["tch_esp"] = c.groupby("ste").tch_calc.transform(lambda s: s.shift().rolling(3, min_periods=1).mean())
c["tch_esp"] = c.tch_esp.fillna(c.groupby("ste").tch_calc.transform("mean"))
c = c.sort_index()
A = {s: d for s, d in a.groupby("ste")}
fil = []
for i, r in c[c.cosecha >= "2018-07-01"].iterrows():
    d = A.get(r.ste)
    v = {"i": i, "mad_dato": 1}
    if d is not None:
        d = d[(d.fecha > r.cosecha - pd.Timedelta(days=150)) & (d.fecha <= r.cosecha + pd.Timedelta(days=20))]
    if d is None or d.empty:
        v.update(tri_L=0.0, n_apl=0)
    else:
        t = d[d.grupo == "trinexapac"]
        v.update(tri_L=t.dosis.sum(), n_apl=len(d), tri_n=len(t),
                 tri_sem=((r.cosecha - t.fecha.min()).days / 7) if len(t) else np.nan,
                 **{g_: int((d.grupo == g_).any()) for g_ in ["trinexapac", "glifosato", "fusilade", "promotor", "alternativo"]},
                 productos=" + ".join(sorted(d.producto.astype(str).str.strip().str.capitalize().unique())))
    fil.append(v)
M = pd.DataFrame(fil).set_index("i")
c = c.join(M)
for g_ in ["trinexapac", "glifosato", "fusilade", "promotor", "alternativo"]:
    c[g_] = c[g_].fillna(0).astype(int)
c["mad_na"] = c.mad_dato.isna().astype(int)
c["con_mad"] = ((c.n_apl > 0) & (c.mad_na == 0)).astype(int)
c["tri_cc"] = np.where(c.trinexapac == 1, c.tri_L * 1000 / c.tch_esp, 0.0)
c["tri_gt12"] = (c.tri_cc > 12).astype(int)
c["tri_le12"] = ((c.trinexapac == 1) & (c.tri_cc <= 12)).astype(int)
cob = c[c.mad_na == 0].groupby("anio").agg(cosechas=("con_mad", "size"), con_madurante=("con_mad", "mean"),
                                           trinexapac=("trinexapac", "mean"), mas_de_12cc=("tri_gt12", "mean"),
                                           glifosato=("glifosato", "mean"), fusilade=("fusilade", "mean"))
print("\ncobertura por año de cosecha (parte de las cosechas):\n", cob.round(2))

# ---------------------------------------------------------------- efecto en el TCH (modelo integrado + madurante)
ED = "C(edad_b, Treatment('12-13'))"
MEC = " + ".join(f"mec_{y}" for y in range(2011, 2027))
NUEVAS = ["def100", "def_na", "rot_tard", "rot_na", "mal_extra30", "mal_na", "fert_120", "fert_na", "lab_na", "no_renov",
          "pm_new", "pm_seg", "pm_nd", "p_na", "nf_exc15", "nf_na"]
BASE = f"y ~ C(ste) + C(anio) + C(Mes) + {ED} + C(soca_b) + C(var) + inc + Corte + {MEC} + llu03_100 + d32 + " + " + ".join(NUEVAS)
est = c.dropna(subset=["llu03", "d32"]).copy()


def fit(f):
    mo = smf.wls(f, data=est, weights=est.Area)
    return mo.fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(est.loc[mo.data.row_labels, "ste"])[0]})


filas = []
for nom, t in [("con o sin madurante", "con_mad + mad_na"),
               ("por producto", "trinexapac + glifosato + fusilade + promotor + alternativo + mad_na"),
               ("trinexapac hasta 12 cc/t y más de 12", "tri_le12 + tri_gt12 + glifosato + fusilade + promotor + alternativo + mad_na"),
               ("trinexapac por cada 5 cc/t", "trinexapac + tri_cc5 + glifosato + fusilade + promotor + alternativo + mad_na")]:
    est["tri_cc5"] = est.tri_cc / 5
    mo = fit(BASE + " + " + t)
    for k in t.split(" + "):
        if k != "mad_na":
            filas.append({"modelo": nom, "variable": k, "efecto_pct": pct(mo.params[k]), "p": mo.pvalues[k],
                          "ic_bajo": pct(mo.conf_int().loc[k, 0]), "ic_alto": pct(mo.conf_int().loc[k, 1]),
                          "cosechas": int(est.loc[mo.model.data.row_labels, k].gt(0).sum())})
ef = pd.DataFrame(filas)
print(f"\nefecto en el TCH (n={len(est)}; cosechas con dato de madurante desde jul-2018):")
print(ef.round(3).to_string(index=False))

# ---------------------------------------------------------------- efecto en la sacarosa (lo que busca el madurante)
# sacarosa % caña («Sacarosa C» del ZQM) de la misma fila que usa el modelo (la de mayor TC-Total de la cosecha): ZQM histórico + G&P 2026
sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
from tchmotor import dataset_gee as dg  # noqa: E402

z = dg._hoja(dg.ZQM_HIST, 0)
z.columns = [str(k).strip() for k in z.columns]
w = openpyxl.load_workbook(r"C:\Users\Agr349\Documents\G&P Suertes_26.xlsx", read_only=True, data_only=True)
it = w["ZQM101"].iter_rows(values_only=True)
for r in it:
    if any(str(v).strip() == "Ubic Técnica" for v in r if v is not None):
        hdr = dg._dedup([str(v).strip() if v is not None else f"c{j}" for j, v in enumerate(r)])
        break
g26 = pd.DataFrame([r for r in it if any(v is not None for v in r[:6])], columns=hdr)
w.close()
g26["AÑO"] = 2026
s = pd.concat([z[["AÑO", "Mes", "Ubic Técnica", "Corte", "TC-Total", "Sacarosa C"]],
               g26[["AÑO", "Mes", "Ubic Técnica", "Corte", "TC-Total", "Sacarosa C"]]], ignore_index=True)
for k in ["AÑO", "Mes", "Corte", "TC-Total", "Sacarosa C"]:
    s[k] = pd.to_numeric(s[k], errors="coerce")
s["ste"] = s["Ubic Técnica"].astype(str).str.strip()
s = s.dropna(subset=["TC-Total", "Sacarosa C"]).sort_values("TC-Total").drop_duplicates(["ste", "AÑO", "Mes", "Corte"], keep="last")
s = s.rename(columns={"AÑO": "anio", "Sacarosa C": "sac"})[["ste", "anio", "Mes", "Corte", "sac"]]
c = c.merge(s, on=["ste", "anio", "Mes", "Corte"], how="left").set_axis(c.index)
c.loc[~c.sac.between(7, 17), "sac"] = np.nan
es = c.dropna(subset=["sac", "llu03", "d32"]).copy()
print(f"\nsacarosa % caña en {c.sac.notna().sum()} de {len(c)} cosechas; promedio por año:",
      c.groupby("anio").sac.mean().round(2).to_dict())
BS = f"sac ~ C(ste) + C(anio) + C(Mes) + {ED} + C(soca_b) + C(var) + inc + mec + llu03_100 + d32"
fs = []
for nom, t in [("por producto", "trinexapac + glifosato + fusilade + promotor + alternativo + mad_na"),
               ("trinexapac hasta 12 cc/t y más de 12", "tri_le12 + tri_gt12 + glifosato + fusilade + promotor + alternativo + mad_na")]:
    mo = smf.wls(BS + " + " + t, data=es, weights=es.Area)
    mo = mo.fit(cov_type="cluster", cov_kwds={"groups": pd.factorize(es.loc[mo.data.row_labels, "ste"])[0]})
    for k in t.split(" + "):
        if k != "mad_na":
            fs.append({"modelo": nom, "variable": k, "puntos_sacarosa": mo.params[k], "p": mo.pvalues[k],
                       "ic_bajo": mo.conf_int().loc[k, 0], "ic_alto": mo.conf_int().loc[k, 1]})
efs = pd.DataFrame(fs)
print("efecto en la sacarosa (puntos de % caña, misma suerte):")
print(efs.round(3).to_string(index=False))

# ---------------------------------------------------------------- 2026: quiénes superaron 12 cc/t
mae = list(openpyxl.load_workbook(MAESTRO, read_only=True, data_only=True)["MAESTRO"].iter_rows(min_row=5, values_only=True))
mae = pd.DataFrame(mae[1:], columns=[str(x).strip() if x is not None else f"c{j}" for j, x in enumerate(mae[0])])
mae = mae[mae.SUERTE.notna()].assign(ste=lambda d: d.SUERTE.astype(str).str.strip()).drop_duplicates("ste").set_index("ste")
q = c[c.anio == 2026].copy()
q["ppto"] = q.ste.map(pd.to_numeric(mae["TCH PPTO"], errors="coerce")).where(lambda v: v > 20)
q["hist"] = q.ste.map(c[c.anio.between(2015, 2025)].groupby("ste").apply(lambda g: wm(g, "tch_calc")))
q["sac_hist"] = q.ste.map(c[c.anio.between(2015, 2025)].dropna(subset=["sac"]).groupby("ste").apply(lambda g: wm(g, "sac")))
q["tri_cc_ppto"] = np.where(q.trinexapac == 1, q.tri_L * 1000 / q.ppto, np.nan)
q["tri_cc_real"] = np.where(q.trinexapac == 1, q.tri_L * 1000 / q.tch_calc, np.nan)
q["clase"] = np.select([q.tri_gt12 == 1, q.tri_le12 == 1, q.con_mad == 1], ["trinexapac > 12 cc/t", "trinexapac ≤ 12 cc/t",
                                                                           "otro madurante"], "sin registro de madurante")
res = q.dropna(subset=["hist"]).groupby("clase").apply(lambda g: pd.Series({
    "cosechas": len(g), "ha": g.Area.sum(), "cc_t": wm(g, "tri_cc") if (g.trinexapac == 1).any() else np.nan,
    "tch_2026": wm(g, "tch_calc"), "tch_hist": wm(g, "hist"), "cambio_pct": 100 * (wm(g, "tch_calc") / wm(g, "hist") - 1),
    "sac_2026": wm(g.dropna(subset=["sac", "sac_hist"]), "sac"), "sac_hist": wm(g.dropna(subset=["sac", "sac_hist"]), "sac_hist")}))
print("\n2026 frente a la historia 2015-25 de cada suerte, por clase de maduración:\n", res.round(1))
print("\n2026: cosechas con trinexapac por encima de 12 cc/t (TCH esperado) — y con el PPTO y el real")
lst = q[q.trinexapac == 1].sort_values("tri_cc", ascending=False)[
    ["ste", "Mes", "Area", "tri_L", "tri_n", "tri_sem", "tch_esp", "ppto", "tch_calc", "tri_cc", "tri_cc_ppto", "tri_cc_real", "productos"]]
print(f"  trinexapac en {len(lst)} cosechas; > 12 cc/t: esperado {int((lst.tri_cc > 12).sum())}, PPTO {int((lst.tri_cc_ppto > 12).sum())}, "
      f"real {int((lst.tri_cc_real > 12).sum())}")
print(lst.head(40).round(1).to_string(index=False))
print("\nproductos aplicados a las cosechas de 2026:\n", q[q.con_mad == 1].productos.value_counts().head(15))
with pd.ExcelWriter(SAL, engine="openpyxl") as w:
    ef.to_excel(w, sheet_name="Efecto en el TCH", index=False)
    efs.to_excel(w, sheet_name="Efecto en la sacarosa", index=False)
    res.reset_index().to_excel(w, sheet_name="2026 por clase", index=False)
    lst.to_excel(w, sheet_name="2026 trinexapac cc por t", index=False)
    cob.reset_index().to_excel(w, sheet_name="Cobertura", index=False)
    q[["ste", "Mes", "Area", "tch_calc", "hist", "ppto", "clase", "tri_L", "tri_cc", "tri_cc_ppto", "tri_cc_real", "tri_sem",
       "productos"]].to_excel(w, sheet_name="Cosechas 2026", index=False)
c[["ste", "anio", "Mes", "Corte", "con_mad", "mad_na", "trinexapac", "glifosato", "fusilade", "promotor", "alternativo",
   "tri_L", "tri_cc", "tri_gt12", "tri_le12", "tri_sem", "productos"]].to_pickle(PK / "madurante_cosechas.pkl")
print("guardado:", SAL)
