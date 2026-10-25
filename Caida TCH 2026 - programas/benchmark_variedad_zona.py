# -*- coding: utf-8 -*-
"""Benchmark de TCH por variedad × zona agroecológica (histórico G&P / ZQM 2016-2025) y calificación de las cosechas 2026.

Benchmark = mediana del TCH de los cortes de esa variedad en esa zona (un ciclo por corte, TC-Total ÷ área). Con menos de
MIN_N cortes se usa la variedad en toda la planta; si tampoco, la zona; si no, la planta. Zona agroecológica del maestro
(en el ZQM 2023-26 viene dañada). Calificación de cada cosecha 2026 (ene-sep G&P + octubre de la app):
1 = TCH más de 5 % sobre el benchmark (óptimo) · 3 = entre −5 % y +5 % (medio) · 5 = más de 5 % por debajo (bajo).

Uso: python benchmark_variedad_zona.py <salida.xlsx>
"""
import sys
import warnings

import numpy as np
import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, r"C:\Users\Agr349\tch-motor")
sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
import presupuesto_cosecha as pc  # noqa: E402
from tchmotor import dataset_gee as dg, produccion as pr  # noqa: E402
from tchmotor.agronomia import cargar_var_map, map_variedad  # noqa: E402

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
SAL = sys.argv[1]
MIN_N, DESDE, HASTA = 5, 2016, 2025
GP = r"C:\Users\Agr349\Documents\G&P Suertes_26.xlsx"
COLS = ["AÑO", "Mes", "Ubic Técnica", "Empresa", "Corte", "Variedad Principal", "Area", "Edad", "TC-Total"]


def leer_gp():
    w = openpyxl.load_workbook(GP, read_only=True, data_only=True)
    filas, hdr, vac = [], None, 0
    for r in w["ZQM101"].iter_rows(values_only=True):
        if hdr is None:
            if any(str(c).strip() == "Ubic Técnica" for c in r if c is not None):
                hdr = dg._dedup([str(x).strip() if x is not None else f"c{j}" for j, x in enumerate(r)])
            continue
        if all(v is None for v in r[:6]):
            vac += 1
            if vac > 50:
                break
            continue
        vac = 0
        filas.append(r)
    g = pd.DataFrame(filas, columns=hdr)
    g["AÑO"] = 2026
    return g


hist, gp = dg._hoja(dg.ZQM_HIST, 0), leer_gp()
for d_ in (hist, gp):
    for k in COLS:
        if k not in d_:
            d_[k] = np.nan
Z = pd.concat([hist[COLS], gp[COLS]], ignore_index=True)
Z = Z[pd.to_numeric(Z.Mes, errors="coerce").isin(range(1, 13))]
Z["Tipo Cosecha"] = ""
c = dg.cortes_zqm(Z)
vm = cargar_var_map()
c["var"] = [str(map_variedad(v, vm) or v).strip().upper() for v in c.variedad]
c["fuente"] = "G&P"
c.loc[c["var"].isin(["#N/A", "NAN", "NONE", ""]), "var"] = "SIN DATO"
mp, _ = pr.cargar_maestro()
rec, _ = pr.cosechas_recientes(pr.leer_cosechas_recientes(), mp, dg.cortes_zqm(dg.leer_zqm()))
rec = rec[(rec.anio == 2026) & (rec.mes >= 10)].copy()
rec["var"] = [str(map_variedad(v, vm) or v).strip().upper() for v in rec.variedad]
rec["fuente"] = "app de Cosecha (oct)"
rec.loc[rec["var"].isin(["#N/A", "NAN", "NONE", ""]), "var"] = "SIN DATO"
c = pd.concat([c, rec[[k for k in c.columns if k in rec.columns]]], ignore_index=True).drop_duplicates(["sec_ste", "anio", "mes"])
zona = {}
for planta in ("riopaila", "castilla"):
    ma, _ = pc.leer_maestro(planta)
    zona.update({s: (v["zagro"] or "").strip().upper() for s, v in ma.items()})
c["zona"] = c.sec_ste.map(zona).fillna("")
c.loc[c.zona.isin(["", "NAN", "CAÑA", "ND", "OCAS"]), "zona"] = "SIN ZONA"
print("cortes con zona:", c.groupby("planta").apply(lambda g: f"{100 * (g.zona != 'SIN ZONA').mean():.0f} %").to_dict())

# ---------------------------------------------------------------- benchmarks
h = c[c.anio.between(DESDE, HASTA)]
filas = []
for planta, g in h.groupby("planta"):
    niveles = [("variedad × zona", ["var", "zona"]), ("variedad (toda la planta)", ["var"]), ("zona (todas las variedades)", ["zona"])]
    for nivel, keys in niveles:
        for k, x in g.groupby(keys):
            k = k if isinstance(k, tuple) else (k,)
            var = k[0] if "var" in keys else "TODAS"
            zn = k[-1] if "zona" in keys else "TODAS"
            filas.append(dict(planta=planta, nivel=nivel, var=var, zona=zn, n=len(x), ha=x.ha.sum(), bench=x.tch.median(),
                              p25=x.tch.quantile(.25), p75=x.tch.quantile(.75), desde=int(x.anio.min()), hasta=int(x.anio.max())))
    filas.append(dict(planta=planta, nivel="planta", var="TODAS", zona="TODAS", n=len(g), ha=g.ha.sum(), bench=g.tch.median(),
                      p25=g.tch.quantile(.25), p75=g.tch.quantile(.75), desde=int(g.anio.min()), hasta=int(g.anio.max())))
B = pd.DataFrame(filas)
B["key"] = B.planta + "|" + B["var"] + "|" + B.zona
OK = set(B.key[B.n >= MIN_N])

# ---------------------------------------------------------------- cosechas 2026
q = c[c.anio == 2026].copy()


def clave(r):
    for k in (f"{r.planta}|{r['var']}|{r.zona}", f"{r.planta}|{r['var']}|TODAS", f"{r.planta}|TODAS|{r.zona}", f"{r.planta}|TODAS|TODAS"):
        if k in OK:
            return k
    return f"{r.planta}|TODAS|TODAS"


q["key"] = [clave(r) for _, r in q.iterrows()]
q = q.merge(B[["key", "nivel", "bench"]], on="key", how="left")
q["dev"] = q.tch / q.bench - 1
q["calif"] = np.where(q.dev > 0.05, 1, np.where(q.dev < -0.05, 5, 3))
for planta, g in q.groupby("planta"):
    print(f"\n{planta}: {len(g)} cosechas 2026 · {g.ha.sum():,.0f} ha · nivel del benchmark {g.nivel.value_counts().to_dict()}")
    print("   calificación (ha):", g.groupby("calif").ha.sum().round(0).to_dict(), "· cortes:", g.calif.value_counts().sort_index().to_dict())

# ---------------------------------------------------------------- libro
F = "Arial"
FB, FNB, FH = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, bold=True, color="FFFFFF")
FIN, FNOTE, FTIT = Font(name=F, size=10, color="0000FF"), Font(name=F, size=9, italic=True, color="595959"), Font(name=F, size=14, bold=True)
HEAD, YEL = PatternFill("solid", fgColor="1E3328"), PatternFill("solid", fgColor="FFFF00")
FILL = {1: PatternFill("solid", fgColor="D9EAD3"), 3: PatternFill("solid", fgColor="FFF2CC"), 5: PatternFill("solid", fgColor="F4CCCC")}
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
wb = Workbook()


def enc(ws, fila, cols):
    for i, (n, w) in enumerate(cols, 1):
        x = ws.cell(fila, i, n)
        x.font, x.alignment, x.fill = FH, CEN, HEAD
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[fila].height = 38


# Benchmark
bw = wb.active
bw.title = "Benchmark"
bw["A1"], bw["A1"].font = f"Benchmark de TCH por variedad × zona agroecológica · histórico G&P {DESDE}-{HASTA}", FTIT
bw["A2"] = "Banda del rango medio (±)"
bw["A2"].font = FNB
bw["B2"] = 0.05
bw["B2"].font, bw["B2"].fill, bw["B2"].number_format = FIN, YEL, "0%"
bw["C2"] = f"Benchmark = mediana del TCH de los cortes. Se usa el nivel más fino con al menos {MIN_N} cortes. Óptimo > benchmark × (1 + banda); bajo < benchmark × (1 − banda)."
bw["C2"].font = FNOTE
cols = [("Clave", 28), ("Planta", 10), ("Nivel", 24), ("Variedad", 14), ("Zona agroecológica", 12), ("Cortes", 8), ("Área (ha)", 10),
        ("Benchmark TCH (mediana)", 12), ("P25", 8), ("P75", 8), ("Óptimo: más de", 11), ("Bajo: menos de", 11), ("Años", 10)]
enc(bw, 4, cols)
B = B.sort_values(["planta", "nivel", "var", "zona"]).reset_index(drop=True)
for i, r in enumerate(B.itertuples(), 5):
    vals = [r.key, r.planta.capitalize(), r.nivel, r.var, r.zona, int(r.n), float(r.ha), float(r.bench), float(r.p25), float(r.p75)]
    for j, v in enumerate(vals, 1):
        x = bw.cell(i, j, v)
        x.font = FIN
        x.number_format = {6: "0", 7: "#,##0", 8: "0.0", 9: "0.0", 10: "0.0"}.get(j, "General")
    bw.cell(i, 11, f"=H{i}*(1+$B$2)").number_format = "0.0"
    bw.cell(i, 12, f"=H{i}*(1-$B$2)").number_format = "0.0"
    bw.cell(i, 13, f"{r.desde}-{r.hasta}").font = FB
NB = 4 + len(B)
bw.freeze_panes = "B5"
bw.auto_filter.ref = f"A4:M{NB}"
RK, RB = f"Benchmark!$A$5:$A${NB}", f"Benchmark!$H$5:$H${NB}"

# Calificación 2026
cw = wb.create_sheet("Calificación 2026", 0)
cw["A1"], cw["A1"].font = "Calificación de las cosechas 2026 frente al benchmark de su variedad en su zona agroecológica", FTIT
cw["A2"] = "1 = más de 5 % sobre el benchmark (óptimo) · 3 = entre −5 % y +5 % (medio) · 5 = más de 5 % por debajo (bajo). La banda se cambia en Benchmark!B2."
cw["A2"].font = FNOTE
cols = [("Planta", 10), ("Suerte", 10), ("Mes", 6), ("Corte", 6), ("Variedad", 14), ("Zona agroecológica", 12), ("Área (ha)", 9), ("TCH 2026", 9),
        ("Benchmark usado", 28), ("Nivel del benchmark", 22), ("Benchmark TCH", 10), ("Frente al benchmark", 11), ("Calificación", 10), ("Fuente", 18)]
enc(cw, 4, cols)
q = q.sort_values(["planta", "mes", "sec_ste"]).reset_index(drop=True)
for i, r in enumerate(q.itertuples(), 5):
    vals = [r.planta.capitalize(), r.sec_ste, int(r.mes), int(r.corte) if pd.notna(r.corte) else None, r.var, r.zona, float(r.ha), float(r.tch), r.key, r.nivel]
    for j, v in enumerate(vals, 1):
        x = cw.cell(i, j, v)
        x.font = FIN
        x.number_format = {7: "0.00", 8: "0.0"}.get(j, "General")
    cw.cell(i, 11, f"=INDEX({RB},MATCH(I{i},{RK},0))").number_format = "0.0"
    cw.cell(i, 12, f"=H{i}/K{i}-1").number_format = "+0.0%;-0.0%"
    cw.cell(i, 13, f"=IF(L{i}>Benchmark!$B$2,1,IF(L{i}<-Benchmark!$B$2,5,3))").number_format = "0"
    cw.cell(i, 13).fill = FILL[int(r.calif)]
    cw.cell(i, 14, r.fuente).font = FB
NC = 4 + len(q)
cw.freeze_panes = "C5"
cw.auto_filter.ref = f"A4:N{NC}"

# Resumen
rs = wb.create_sheet("Resumen", 0)
rs["A1"], rs["A1"].font = "Resumen: cosechas 2026 por calificación varietal", FTIT
enc(rs, 3, [("Planta", 12), ("Calificación", 14), ("Cortes", 10), ("Área (ha)", 12), ("% del área", 10), ("TCH medio", 10)])
CP, CC, CA, CT = (f"'Calificación 2026'!${k}$5:${k}${NC}" for k in ("A", "M", "G", "H"))
f = 4
for planta in ("Riopaila", "Castilla"):
    for cal, nom in ((1, "1 · óptimo"), (3, "3 · medio"), (5, "5 · bajo")):
        rs.cell(f, 1, planta).font = FB
        rs.cell(f, 2, nom).font = FB
        rs.cell(f, 3, f'=COUNTIFS({CP},A{f},{CC},{cal})')
        rs.cell(f, 4, f'=SUMIFS({CA},{CP},A{f},{CC},{cal})').number_format = "#,##0"
        rs.cell(f, 5, f'=IF(SUMIFS({CA},{CP},A{f})=0,"",D{f}/SUMIFS({CA},{CP},A{f}))').number_format = "0%"
        rs.cell(f, 6, f'=IF(D{f}=0,"",SUMPRODUCT(({CP}=A{f})*({CC}={cal})*{CA}*{CT})/D{f})').number_format = "0.0"
        rs.cell(f, 2).fill = FILL[cal]
        f += 1
rs.cell(f + 1, 1, f"Histórico G&P/ZQM {DESDE}-{HASTA} para el benchmark; 2026 = G&P ene-sep + app de Cosecha octubre. Zona agroecológica del maestro.").font = FNOTE
wb.save(SAL)
print("guardado:", SAL)
