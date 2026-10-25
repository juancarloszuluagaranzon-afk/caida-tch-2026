# -*- coding: utf-8 -*-
"""Riopaila: riego y lluvia de la caña que se corta en 2027 (lo que ya vivió hasta hoy) frente a lo normal.

Cohorte 2027: suertes del maestro (uso CAÑA, no RENOVACION) cuyo ciclo actual empezó entre el 1-nov-2025 y hoy
(inicio = max(siembra, último corte)). Por etapa ya cumplida (meses 0-4 y 4-8): lluvia y evaporación de su EMA
(La Paila / Bugalagrande: repo hasta 8-sep-2026 + Meteoportal hasta 18-oct), riego del consolidado (hasta 11-oct),
déficit que quedó después del riego, y lluvia de los primeros 91 días (arranque).
Normal = ciclos de las cosechas 2022-25 que empezaron el mismo mes del año (rio_riego.pkl de riego_2026.py).
Efectos (mismos modelos de riego_2026.py y factores_sep.py, cosechas Riopaila):
  déficit después del riego: meses 0-4 +0,70 %/100 mm, 4-8 −0,45 %/100 mm, 8 a cosecha −1,25 %/100 mm
  lluvia del arranque (0-3 meses): −1,20 %/100 mm
Escenario El Niño para el final del ciclo (aún no ocurre): déficit final sin cubrir como en 2026 (+117 mm sobre lo normal).

Uso: python riego_2027.py <CONSOLIDADO DE RIEGOS.xlsx> <rio_riego.pkl> <carpeta repo Precipitacion-riopaila> <ha 2027> <TCH 2027>
"""
import sys
import warnings
from datetime import date
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, r"C:\Users\Agr349\tch-motor\scripts")
import presupuesto_cosecha as pc  # noqa: E402

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
XLS, RPKL, REPO = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
HA27, TCH27 = float(sys.argv[4]), float(sys.argv[5])
HOY = pd.Timestamp(date.today())
B_DEF = {"e0_4": 0.0070, "e4_8": -0.0045, "fin": -0.0125}
B_ARR = -0.0120
DIAS_MES = 30.44
wm = lambda g, k: np.average(g[k].dropna(), weights=g.area[g[k].notna()]) if g[k].notna().any() else np.nan  # noqa: E731

# clima diario por estación
e = pd.read_csv(REPO / "clima_estaciones.csv", sep=";")
e["f"] = pd.to_datetime(e.fc_obsrvcion)
wb = openpyxl.load_workbook(r"C:\Users\Agr349\Documents\RIOPAILA AGRICOLA\Clima Cenicaña\Clima_Cenicana_Riopaila_Castilla_1993-2026.xlsx",
                            read_only=True, data_only=True)
it = wb["Diario 2026"].iter_rows(values_only=True)
enc = next(it)
Mt = pd.DataFrame(list(it), columns=enc)
Mt["f"] = pd.to_datetime(Mt.Fecha)
CL = {}
for est, cod in (("La Paila", "PAI"), ("Bugalagrande", "BLG")):
    d = e[e.nm_ema == est].set_index("f").sort_index()[["prcptcion", "evp_clclada"]]
    m = Mt[Mt["Código"] == cod].set_index("f")[["Lluvia (mm)", "Evaporación (mm)"]].rename(
        columns={"Lluvia (mm)": "prcptcion", "Evaporación (mm)": "evp_clclada"})
    CL[est] = d.combine_first(m)
ULT_CLIMA = min(s.prcptcion.dropna().index.max() for s in CL.values())
es = pd.read_csv(REPO / "estaciones_suertes.csv", sep=";", encoding="latin-1", skiprows=1)
est_de = {str(s).strip(): ("Bugalagrande" if "buga" in str(x).lower() else "La Paila") for s, x in zip(es.iloc[:, 1], es.iloc[:, 2])}

# riegos
filas = list(openpyxl.load_workbook(XLS, read_only=True, data_only=True)["Consolidado de riegos "].iter_rows(values_only=True))
h = next(i for i, r in enumerate(filas) if r and str(r[0]).strip() == "Fecha")
encr = [str(c).strip() if c is not None else f"c{j}" for j, c in enumerate(filas[h])]
R = pd.DataFrame(filas[h + 1:], columns=encr)
R = R[R.Fecha.notna()]
R = R[~R.duplicated(subset=[c for c in encr if not c.startswith("c")])].copy()
R["f"] = pd.to_datetime(R.Fecha, errors="coerce")
R["ste"] = R["SEC-SUERTE"].astype(str).str.strip()
R["a"] = pd.to_numeric(R["ÁREA REGADA"], errors="coerce")
R["mm"] = pd.to_numeric(R["MM"], errors="coerce")
R["mm"] = R.mm.where(R.mm.between(20, 300), R.mm.where(R.mm.between(20, 300)).median())
R = R[R.f.notna() & (R.a > 0)]
R["vol"] = R.a * R.mm
ULT_RIEGO = R.f.max()
G = {s: g for s, g in R.groupby("ste")}

# cohorte 2027
ma, _ = pc.leer_maestro("riopaila")
m = pd.DataFrame([dict(ste=s, area=v["area"] or 0, var=v["var"], ini=max([d for d in (v["siembra"], v["ult"]) if d] or [None]) if (v["siembra"] or v["ult"]) else None)
                  for s, v in ma.items()])
m["ini"] = pd.to_datetime(m.ini)
co = m[(m["var"].str.upper() != "RENOVACION") & (m.ini >= "2025-11-01") & (m.ini <= HOY) & (m.area > 0)].copy()
co["est"] = co.ste.map(est_de).fillna("La Paila")
lim = min(ULT_CLIMA, ULT_RIEGO)
print(f"cohorte 2027: {len(co)} suertes · {co.area.sum():,.0f} ha · ciclo iniciado nov-2025 a hoy; datos hasta {lim.date()}")
TR = {"e0_4": (0, 4), "e4_8": (4, 8)}


def etapa(r, k):
    a, b = TR[k]
    d0, d1 = r.ini + pd.Timedelta(days=DIAS_MES * a), r.ini + pd.Timedelta(days=DIAS_MES * b)
    if d1 > lim:
        return np.nan, np.nan, np.nan
    s = CL[r.est].loc[d0:d1 - pd.Timedelta(days=1)]
    n = (d1 - d0).days
    llu, et = s.prcptcion.mean() * n, s.evp_clclada.mean() * n
    g = G.get(r.ste)
    mm = 0.0 if g is None else g[(g.f >= d0) & (g.f < d1)].vol.sum() / r.area
    return llu, et, mm


for k in TR:
    v = [etapa(r, k) for r in co.itertuples()]
    co[f"llu_{k}"], co[f"et_{k}"], co[f"mm_{k}"] = zip(*v)
    co[f"def_{k}"] = co[f"et_{k}"] - co[f"llu_{k}"] - co[f"mm_{k}"]
co["llu03"] = [CL[r.est].prcptcion.loc[r.ini:r.ini + pd.Timedelta(days=90)].mean() * 91
               if r.ini + pd.Timedelta(days=90) <= lim else np.nan for r in co.itertuples()]

# normal: ciclos de las cosechas 2022-25 que empezaron el mismo mes del año
rr = pd.read_pickle(RPKL)
rr = rr[rr.anio.between(2022, 2025)].copy()
rr["mes_ini"] = rr.ini.dt.month
norm = rr.groupby("mes_ini").apply(lambda g: pd.Series({"def_e0_4": np.average(g.def_e0_4.dropna(), weights=g.Area[g.def_e0_4.notna()]),
                                                        "def_e4_8": np.average(g.def_e4_8.dropna(), weights=g.Area[g.def_e4_8.notna()]),
                                                        "llu03": np.average(g.llu03.dropna(), weights=g.Area[g.llu03.notna()])}))
co["mes_ini"] = co.ini.dt.month
for k in ("def_e0_4", "def_e4_8", "llu03"):
    co[f"n_{k}"] = co.mes_ini.map(norm[k])

ton27 = HA27 * TCH27
print(f"\nproducción 2027 de referencia: {HA27:,.0f} ha × TCH {TCH27} = {ton27:,.0f} t")
res = {}
for k, col in (("e0_4", "def_e0_4"), ("e4_8", "def_e4_8")):
    x = co.dropna(subset=[col])
    if not len(x):
        print(f"  etapa {k}: ninguna suerte la ha completado todavía")
        continue
    dlt = wm(x, col) - wm(x, f"n_{col}")
    ef = B_DEF[k] * dlt / 100
    t = -ef * (x.area.sum() / HA27) * ton27
    res[f"riego_{k}"] = (x.area.sum(), t)
    print(f"  etapa {k}: {len(x)} suertes · {x.area.sum():,.0f} ha · déficit después del riego {wm(x, col):.0f} mm vs normal "
          f"{wm(x, 'n_' + col):.0f} → efecto {100 * ef:+.1f} % en esas suertes · {-t:+,.0f} t (− = pérdida)")
    print(f"      lluvia {wm(x, 'llu_' + k):.0f} mm · evaporación {wm(x, 'et_' + k):.0f} mm · riego {wm(x, 'mm_' + k):.0f} mm")
x = co.dropna(subset=["llu03"])
dl = wm(x, "llu03") - wm(x, "n_llu03")
ef = B_ARR * dl / 100
res["arranque"] = (x.area.sum(), -ef * (x.area.sum() / HA27) * ton27)
print(f"  arranque (lluvia 0-3 meses): {len(x)} suertes · {x.area.sum():,.0f} ha · {wm(x, 'llu03'):.0f} mm vs normal {wm(x, 'n_llu03'):.0f} "
      f"→ efecto {100 * ef:+.1f} % · {ef * (x.area.sum() / HA27) * ton27:+,.0f} t")
esc = B_DEF["fin"] * 117 / 100
print(f"\nescenario El Niño (final del ciclo como 2026, +117 mm sin cubrir): {100 * esc:+.1f} % → {esc * ton27:+,.0f} t en 2027")
co.to_pickle("cohorte_2027_riego.pkl")
pd.to_pickle({"res": res, "esc_fin": esc * ton27, "ton27": ton27}, "riego_2027.pkl")
